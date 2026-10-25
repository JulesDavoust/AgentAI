"""Module bibliotheque/excel_com.py, testé avec l'enregistreur (Excel n'est pas disponible ici)."""

import sys

import pytest

from vba2py.bibliotheque import excel_com
from vba2py.bibliotheque.excel_com import EnregistreurCOM, preparer_feuille, session_excel, xl


def creer_visuel_excel(wb):
    """Traduction ligne à ligne de la partie TCD + graphique de CreerVisuel (marco_25_10.bas),
    telle que l'agent doit l'écrire dans connecteurs/excel.py."""
    ws = wb.Worksheets("Feuille1")
    while ws.ChartObjects().Count > 0:
        ws.ChartObjects(1).Delete()
    while ws.PivotTables().Count > 0:
        ws.PivotTables(1).TableRange2.Clear()
    pc = wb.PivotCaches().Create(SourceType=xl.xlDatabase, SourceData="tblDonnees")
    pt = pc.CreatePivotTable(TableDestination=ws.Range("A5"), TableName="TCD_Volume")
    pt.PivotFields("Date").Orientation = xl.xlRowField
    pt.PivotFields("Region").Orientation = xl.xlColumnField
    champ = pt.AddDataField(pt.PivotFields("Volume_GWh"), "Volume (GWh)", xl.xlSum)
    champ.NumberFormat = "#,##0.0"
    pt.PivotFields("Date").NumberFormat = "dd/mm"
    co = ws.ChartObjects().Add(Left=ws.Range("J5").Left, Top=ws.Range("J5").Top, Width=520, Height=300)
    co.Chart.SetSourceData(pt.TableRange1)
    co.Chart.ChartType = xl.xlLine
    co.Chart.HasTitle = True
    co.Chart.ChartTitle.Text = "Volumes distribues par region (GWh)"


def test_enregistreur_reproduit_la_sequence_de_la_macro():
    wb = EnregistreurCOM("wb")
    creer_visuel_excel(wb)
    assert wb.instructions()[:7] == [
        "wb.Worksheets('Feuille1')",
        "wb.Worksheets('Feuille1').ChartObjects()",
        "wb.Worksheets('Feuille1').PivotTables()",
        "wb.PivotCaches()",
        "wb.PivotCaches().Create(SourceType=1, SourceData='tblDonnees')",
        "wb.Worksheets('Feuille1').Range('A5')",
        "wb.PivotCaches().Create(SourceType=1, SourceData='tblDonnees').CreatePivotTable("
        "TableDestination=<wb.Worksheets('Feuille1').Range('A5')>, TableName='TCD_Volume')",
    ]
    assert wb.affectations("PivotFields('Date').Orientation") == [xl.xlRowField]
    assert wb.affectations("PivotFields('Region').Orientation") == [xl.xlColumnField]
    [(args, _)] = wb.appels(".AddDataField")
    assert args[1:] == ("Volume (GWh)", -4157)
    assert wb.affectations("ChartTitle.Text") == ["Volumes distribues par region (GWh)"]
    assert wb.affectations("Chart.ChartType") == [4]


def test_enregistreur_reponses_configurables():
    wb = EnregistreurCOM("wb", reponses={"wb.Worksheets('Feuille1').Range('J5').Left": 300.0})
    assert wb.Worksheets("Feuille1").Range("J5").Left == 300.0
    assert wb.Worksheets("Feuille1").ChartObjects().Count == 0  # défaut
    assert list(wb.Worksheets("Feuille1").ChartObjects()) == []


def test_session_ouvre_enregistre_ferme(tmp_path):
    chemin = tmp_path / "copie.xlsm"
    chemin.write_bytes(b"")
    excel = EnregistreurCOM("excel")
    with session_excel(chemin, excel=excel) as wb:
        wb.Worksheets("Feuille1").Range("A1").Value = 1
    lignes = excel.instructions()
    assert lignes[0] == f"excel.Workbooks.Open({str(chemin.resolve())!r})"
    assert lignes[-2].endswith(".Save()") and lignes[-1].endswith(".Close(False)")
    assert not excel.appels("excel.Quit")  # instance fournie : pas fermée ici


def test_session_chemin_relatif_et_fermeture_sur_erreur(tmp_path, monkeypatch):
    (tmp_path / "copie.xlsm").write_bytes(b"")
    monkeypatch.chdir(tmp_path)
    excel = EnregistreurCOM("excel")
    monkeypatch.setattr(excel_com, "ouvrir_excel", lambda: excel)
    with pytest.raises(RuntimeError):
        with session_excel("copie.xlsm") as wb:
            raise RuntimeError("échec")
    lignes = excel.instructions()
    assert str((tmp_path / "copie.xlsm").resolve()) in lignes[0]
    assert not any(l.endswith(".Save()") for l in lignes)  # rien d'enregistré
    assert lignes[-2].endswith(".Close(False)") and lignes[-1] == "excel.Quit()"


class com_error(Exception):  # même nom que pywintypes.com_error
    pass


ERREUR_03EC = com_error(-2147352567, "Une exception s'est produite.", (0, None, None, None, 0, -2146827284), None)


class FauxCache:
    _oleobj_ = object()

    def CreatePivotTable(self, TableDestination, TableName):
        raise ERREUR_03EC


class FauxClasseur:
    """Objet « réel » minimal : se comporte comme un objet COM (attribut _oleobj_)."""

    _oleobj_ = object()

    def __init__(self):
        self.recu = {}

    def Worksheets(self, nom):
        feuille = FauxClasseur()
        feuille.nom = nom
        return feuille

    def Range(self, adresse):
        plage = FauxClasseur()
        plage.adresse = adresse
        return plage

    def PivotCaches(self):
        return self

    def Create(self, SourceType, SourceData):
        self.recu["Create"] = (SourceType, SourceData)
        return FauxCache()

    def Save(self):
        self.recu["Save"] = True

    def Close(self, enregistrer):
        self.recu["Close"] = enregistrer


class FauxExcel:
    def __init__(self):
        self.classeur = FauxClasseur()
        self.Workbooks = self

    def Open(self, chemin):
        return self.classeur


def test_traceur_indique_l_instruction_refusee(tmp_path):
    from vba2py.bibliotheque.excel_com import ErreurExcel

    (tmp_path / "c.xlsx").write_bytes(b"")
    excel = FauxExcel()
    with pytest.raises(ErreurExcel) as info:
        with session_excel(tmp_path / "c.xlsx", excel=excel) as wb:
            ws = wb.Worksheets("Feuille1")
            ws.Range("A1").Value = "Titre"  # affectation transmise à l'objet réel
            pc = wb.PivotCaches().Create(SourceType=xl.xlDatabase, SourceData="tblDonnees")
            pc.CreatePivotTable(TableDestination=ws.Range("A5"), TableName="TCD_Volume")
    message = str(info.value)
    assert message.splitlines()[0] == (
        "Excel a refusé : excel.Workbooks.Open(%r).PivotCaches().Create(SourceType=1, SourceData='tblDonnees')"
        ".CreatePivotTable(TableDestination=excel.Workbooks.Open(%r).Worksheets('Feuille1').Range('A5'), "
        "TableName='TCD_Volume')" % ((str((tmp_path / "c.xlsx").resolve()),) * 2)
    )
    assert "0x800A03EC" in message
    assert "SourceData='tblDonnees'" in message  # instructions précédentes
    assert info.value.code == -2146827284
    assert excel.classeur.recu["Create"] == (1, "tblDonnees")  # arguments déballés
    assert "Save" not in excel.classeur.recu and excel.classeur.recu["Close"] is False


def test_traceur_transmet_valeurs_et_erreurs_python(tmp_path):
    (tmp_path / "c.xlsx").write_bytes(b"")
    excel = FauxExcel()
    with session_excel(tmp_path / "c.xlsx", excel=excel) as wb:
        plage = wb.Worksheets("Feuille1").Range("B2")
        assert plage.adresse == "B2"  # valeur simple : renvoyée telle quelle
        with pytest.raises(AttributeError):
            wb.ProprieteInexistante  # erreur Python : non transformée
    assert excel.classeur.recu == {"Save": True, "Close": False}


def test_session_fichier_absent(tmp_path):
    with pytest.raises(FileNotFoundError):
        with session_excel(tmp_path / "absent.xlsx", excel=EnregistreurCOM()):
            pass


def test_constantes():
    assert (xl.xlRowField, xl.xlSum, xl.xlLine, xl.xlTypePDF) == (1, -4157, 4, 0)
    with pytest.raises(AttributeError, match="valeur numérique"):
        xl.xlConstanteInexistante


def test_preparer_feuille():
    from openpyxl import Workbook
    from openpyxl.chart import LineChart

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Titre"
    ws.add_chart(LineChart(), "J5")
    ws._pivots.append(object())
    preparer_feuille(ws)
    assert ws["A1"].value is None and ws._charts == [] and ws._pivots == []


@pytest.mark.skipif(sys.platform != "win32", reason="Excel réel : Windows uniquement")
def test_excel_reel(tmp_path):
    pytest.importorskip("win32com")
    import datetime as dt

    from openpyxl import Workbook, load_workbook
    from openpyxl.worksheet.table import Table

    wb = Workbook()
    wb.active.title = "Feuille1"
    d = wb.create_sheet("Feuille2")
    d.append(["Date", "Region", "Volume_GWh"])
    for jour in (1, 2):
        for r, v in (("Nord", 10.0), ("Sud", 20.0)):
            d.append([dt.datetime(2026, 10, jour), r, v * jour])
    d.add_table(Table(displayName="tblDonnees", ref=f"A1:C{d.max_row}"))
    wb.save(tmp_path / "reel.xlsx")
    with session_excel(tmp_path / "reel.xlsx") as wbx:
        creer_visuel_excel(wbx)
    ws = load_workbook(tmp_path / "reel.xlsx")["Feuille1"]
    assert [p.name for p in ws._pivots] == ["TCD_Volume"]
    assert ws._charts and ws._charts[0].pivotSource is not None
    assert ws["B7"].value == 10  # valeurs calculées par Excel
