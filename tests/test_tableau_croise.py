"""Module bibliotheque/tableau_croise.py : création et conservation d'un tableau croisé sans Excel."""

import csv
import datetime as dt
import shutil
import subprocess

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from vba2py.bibliotheque.tableau_croise import (
    FormeNonPriseEnCharge,
    assurer_graphique_croise,
    assurer_tableau_croise,
)

REGIONS = ["Nord", "Sud", "Est"]
# Volumes choisis pour des totaux faciles à vérifier à la main
VOLUMES = {("Nord", 1): 10.0, ("Sud", 1): 20.0, ("Est", 1): 30.0, ("Nord", 2): 1.5, ("Sud", 2): 2.5, ("Est", 2): 3.0}


def classeur(lignes_en_plus=()):
    wb = Workbook()
    wb.active.title = "Feuille1"
    d = wb.create_sheet("Feuille2")
    d.append(["Date", "Region", "Volume_GWh", "Statut"])
    for jour in (1, 2):
        for r in REGIONS:
            d.append([dt.datetime(2026, 10, jour), r, VOLUMES[(r, jour)], "Historique"])
    for l in lignes_en_plus:
        d.append(l)
    d.add_table(Table(displayName="tblDonnees", ref=f"A1:D{d.max_row}"))
    return wb


def creer(wb, **kwargs):
    params = dict(feuille="Feuille1", nom="TCD_Volume", table_source="tblDonnees", ancre="A5",
                  ligne="Date", colonne="Region", valeur="Volume_GWh", libelle_valeur="Volume (GWh)")
    params.update(kwargs)
    return assurer_tableau_croise(wb, **params)


def valeurs(ws, plage):
    return [[c.value for c in r] for r in ws[plage]]


def test_creation_valeurs_et_disposition(tmp_path):
    wb = classeur()
    res = creer(wb, format_valeurs="#,##0.0", format_lignes="dd/mm")
    assert (res.action, res.plage) == ("cree", "A5:E9")
    wb.save(tmp_path / "c.xlsx")

    ws = load_workbook(tmp_path / "c.xlsx")["Feuille1"]
    # Colonnes triées comme Excel (Est, Nord, Sud), totaux : jour 1 = 60, jour 2 = 7, Est = 33...
    assert valeurs(ws, "A5:E9") == [
        ["Volume (GWh)", "Étiquettes de colonnes", None, None, None],
        ["Étiquettes de lignes", "Est", "Nord", "Sud", "Total général"],
        [dt.datetime(2026, 10, 1), 30, 10, 20, 60],
        [dt.datetime(2026, 10, 2), 3, 1.5, 2.5, 7],
        ["Total général", 33, 11.5, 22.5, 67],
    ]
    assert ws["A7"].number_format == "dd/mm" and ws["B7"].number_format == "#,##0.0"


def test_definition_relue(tmp_path):
    wb = classeur()
    creer(wb)
    wb.save(tmp_path / "c.xlsx")
    [pivot] = load_workbook(tmp_path / "c.xlsx")["Feuille1"]._pivots
    assert pivot.name == "TCD_Volume"
    assert pivot.location.ref == "A5:E9"
    assert pivot.cache.refreshOnLoad is True
    assert pivot.cache.cacheSource.worksheetSource.name == "tblDonnees"
    assert [f.name for f in pivot.cache.cacheFields] == ["Date", "Region", "Volume_GWh", "Statut"]
    assert len(pivot.cache.records.r) == 6
    assert [f.axis for f in pivot.pivotFields] == ["axisRow", "axisCol", None, None]
    assert pivot.dataFields[0].name == "Volume (GWh)" and pivot.dataFields[0].subtotal == "sum"


def test_sans_colonne_et_moyenne():
    wb = classeur()
    res = creer(wb, colonne=None, fonction="average", libelle_valeur="Moyenne")
    ws = wb["Feuille1"]
    assert res.plage == "A5:B8"
    assert valeurs(ws, "A5:B8") == [
        ["Étiquettes de lignes", "Moyenne"],
        [dt.datetime(2026, 10, 1), 20],
        [dt.datetime(2026, 10, 2), pytest.approx(7 / 3)],
        ["Total général", pytest.approx(67 / 6)],
    ]


def test_conservation_d_un_tableau_existant(tmp_path):
    wb = classeur()
    creer(wb)
    wb.save(tmp_path / "c.xlsx")

    wb = load_workbook(tmp_path / "c.xlsx")
    res = creer(wb, ancre="Z99")  # paramètres ignorés : le tableau existe déjà
    assert res.action == "conserve" and res.plage == "A5:E9"
    assert any("Activer la modification" in a for a in res.avertissements)
    wb.save(tmp_path / "c2.xlsx")
    assert len(load_workbook(tmp_path / "c2.xlsx")["Feuille1"]._pivots) == 1


def test_graphique_croise_cree_puis_conserve(tmp_path):
    wb = classeur()
    res = creer(wb)
    chart = assurer_graphique_croise(wb, res, ancre="J5", titre="Volumes", nom_fichier="c.xlsx")
    assert len(chart.series) == 3  # une série par région, sans « Total général »
    assert chart.series[0].val.numRef.f == "'Feuille1'!$B$7:$B$8"
    assert chart.series[0].tx.strRef.f == "'Feuille1'!$B$6"
    wb.save(tmp_path / "c.xlsx")

    wb = load_workbook(tmp_path / "c.xlsx")
    res = creer(wb)
    [relu] = wb["Feuille1"]._charts
    assert relu.pivotSource.name == "[c.xlsx]Feuille1!TCD_Volume"
    assert assurer_graphique_croise(wb, res, ancre="J5") is relu  # pas de doublon


def test_deux_tableaux_identifiants_distincts():
    wb = classeur()
    creer(wb)
    creer(wb, nom="TCD_2", ancre="A20", colonne=None)
    assert sorted(p.cacheId for p in wb["Feuille1"]._pivots) == [0, 1]


def test_erreurs():
    wb = classeur()
    with pytest.raises(KeyError, match="Inconnu"):
        creer(wb, valeur="Inconnu")
    with pytest.raises(KeyError, match="tblAbsente"):
        creer(wb, table_source="tblAbsente")
    with pytest.raises(FormeNonPriseEnCharge):
        creer(wb, fonction="stdDev")


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice absent")
def test_libreoffice_reconstruit_le_tableau(tmp_path):
    """Les cellules sont vidées : seul un tableau croisé valide permet à LibreOffice de les recalculer."""
    wb = classeur(lignes_en_plus=[[dt.datetime(2026, 10, 3), "Nord", 100.0, "Phase de test"]])
    creer(wb)
    for row in wb["Feuille1"]["A5:E10"]:
        for c in row:
            c.value = None
    wb.save(tmp_path / "vide.xlsx")
    subprocess.run(
        ["soffice", "--headless", "--calc", "--convert-to",
         "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,false,false,false,-1",
         "--outdir", str(tmp_path / "lo"), str(tmp_path / "vide.xlsx")],
        check=True, capture_output=True, timeout=180,
    )
    lignes = list(csv.reader(open(tmp_path / "lo" / "vide-Feuille1.csv", encoding="utf-8")))
    assert lignes[6][1:5] == ["30", "10", "20", "60"]
    assert lignes[8][0].startswith("2026-10-03") and lignes[8][2] == "100"
    assert lignes[9][1:5] == ["33", "111.5", "22.5", "167"]


def test_effacer_hors_tableaux_croises():
    from openpyxl.chart import BarChart
    from vba2py.bibliotheque.tableau_croise import effacer_hors_tableaux_croises

    wb = classeur()
    res = creer(wb)
    assurer_graphique_croise(wb, res, ancre="J5")
    ws = wb["Feuille1"]
    ws["A1"] = "Titre"
    ws["H2"] = "KPI"
    ws.add_chart(BarChart(), "J30")  # graphique ordinaire : supprimé
    effacer_hors_tableaux_croises(ws)
    assert ws["A1"].value is None and ws["H2"].value is None
    assert ws["B7"].value == 30  # cellule du tableau croisé conservée
    assert len(ws._charts) == 1 and ws._charts[0].pivotSource is not None
