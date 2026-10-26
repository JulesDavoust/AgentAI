"""Référence VBA : neutralisation du code d'origine et exécution isolée (simulée ici, réelle sous Windows)."""

import json
from pathlib import Path

import pytest

from vba2py.bibliotheque.excel_com import EnregistreurCOM, excel_disponible
from vba2py.extract import SOURCE_DIRNAME
from vba2py.reference_vba import executer, executer_dans_excel, module_harnais, neutraliser

VBA = '''Attribute VB_Name = "Module1"
Option Explicit
Private Const DOSSIER As String = "C:\\Users\\jdavoust\\natran\\VBATF_CSV\\"   ' A MODIFIER "C:\\pas\\touche"
Public Sub Lancer()
    Dim ol As Object, s As String
    s = Trim$(InputBox("Region ?", "Filtre", "Nord"))
    If MsgBox("Continuer ?", vbYesNo + vbQuestion) = vbNo Then Exit Sub
    MsgBox "Texte qui cite MsgBox et InputBox dans une chaîne", vbInformation
    Set ol = CreateObject("Outlook.Application")
    Set d = CreateObject("Scripting.Dictionary")
    ThisWorkbook.FollowHyperlink DOSSIER & "carte.html"
    Shell "notepad.exe", vbNormalFocus
    x = Application.InputBox("Nombre ?", Type:=1)
    ' MsgBox dans un commentaire : inchangé
End Sub
'''


def test_neutralisation():
    code, modifications = neutraliser(VBA, Path("/bac/chemins"))
    lignes = code.split("\r\n")
    assert '"/bac/chemins\\C\\Users\\jdavoust\\natran\\VBATF_CSV\\"' in lignes[2]
    assert lignes[2].endswith(' \' A MODIFIER "C:\\pas\\touche"')  # commentaire intact
    assert 'Trim$(vba2py_InputBox("Region ?", "Filtre", "Nord"))' in code
    assert 'If vba2py_MsgBox("Continuer ?", vbYesNo + vbQuestion) = vbNo' in code
    assert 'vba2py_MsgBox "Texte qui cite MsgBox et InputBox dans une chaîne", vbInformation' in code
    assert 'Set ol = vba2py_CreateObject("Outlook.Application")' in code
    assert 'Set d = vba2py_CreateObject("Scripting.Dictionary")' in code  # filtré à l'exécution
    assert 'vba2py_FollowHyperlink DOSSIER & "carte.html"' in code
    assert 'vba2py_Shell "notepad.exe", vbNormalFocus' in code
    assert 'x = vba2py_InputBox("Nombre ?", Type:=1)' in code
    assert "    ' MsgBox dans un commentaire : inchangé" in lignes
    assert lignes[0] == 'Attribute VB_Name = "Module1"'
    assert any(m.startswith("chemin redirigé : C:\\Users\\jdavoust") for m in modifications)


def test_neutralisation_des_macros_reelles():
    """Les modules de l'exemple ne doivent plus contenir aucun appel externe ni dialogue direct."""
    import re

    for fichier in sorted(Path("examples/macros_natran_vf_vba").glob("*.bas")):
        code, _ = neutraliser(fichier.read_bytes().decode("cp1252"), Path("C:/bac"))
        sans_chaines = re.sub(r'"(?:[^"]|"")*"', '""', code)
        sans_commentaires = "\n".join(l.split("'")[0] for l in sans_chaines.splitlines())
        for interdit in (r"(?<![\w.])MsgBox\b", r"(?<![\w.])InputBox\b", r"(?<!vba2py_)CreateObject\(",
                         r"\.FollowHyperlink\b"):
            assert not re.search(interdit, sans_commentaires), (fichier.name, interdit)


def test_harnais():
    code = module_harnais(["Module1.Lancer", "CréationCSV.GenererCSVDuJourV3"])
    assert '        Case "Module1.Lancer"\r\n            Module1.Lancer\r\n' in code
    assert '        Case "CréationCSV.GenererCSVDuJourV3"\r\n            CréationCSV.GenererCSVDuJourV3\r\n' in code
    assert code.startswith('Attribute VB_Name = "vba2py_Harnais"')
    assert "Err.Raise vbObjectError + 513" in code  # Outlook & co. neutralisés
    with pytest.raises(ValueError):
        module_harnais(["Module1.Lancer : Kill \"x\""])


@pytest.fixture
def projet(tmp_path):
    d = tmp_path / "projet"
    (d / SOURCE_DIRNAME).mkdir(parents=True)
    (d / SOURCE_DIRNAME / "Module1.bas").write_text(VBA, encoding="utf-8")
    (d / "tests" / "fixtures").mkdir(parents=True)
    (d / "tests" / "fixtures" / "classeur.xlsx").write_bytes(b"PK")
    return d


def test_execution_et_reference_enregistree(projet):
    recus = {}

    def faux_lanceur(parametres, delai):
        recus.update(parametres)
        modules = [Path(m) for m in parametres["modules"]]
        assert [m.name for m in modules] == ["Module1.bas", "vba2py_Harnais.bas"]
        assert "vba2py_MsgBox" in modules[0].read_bytes().decode("cp1252")  # copie neutralisée en Windows-1252
        Path(parametres["journal"]).write_bytes("DEBUT\tModule1.Lancer\nMSGBOX\tContinuer ?\nFIN\tmacro terminée\n".encode("cp1252"))
        Path(parametres["classeur"]).write_bytes(b"xlsm")
        ecrit = Path(parametres["journal"]).parent / "chemins" / "C" / "Users" / "carte.html"
        ecrit.parent.mkdir(parents=True)
        ecrit.write_text("<html>")
        return "OK"

    resultat = executer(projet, "Module1.Lancer", "tests/fixtures/classeur.xlsx", ["oui"],
                        projet / "tests" / "references" / "lancer", ["Module1.Lancer"], lanceur=faux_lanceur)
    assert resultat.statut == "OK"
    assert recus["reponses"] == ["oui"] and recus["fixture"].endswith("classeur.xlsx")
    ref = projet / "tests" / "references" / "lancer"
    assert (ref / "classeur_apres.xlsm").read_bytes() == b"xlsm"
    assert (ref / "fichiers_ecrits" / "C" / "Users" / "carte.html").exists()
    journal = (ref / "journal.txt").read_text(encoding="utf-8")
    assert "# Modifications de neutralisation" in journal and "MSGBOX\tContinuer ?" in journal
    assert json.loads((ref / "resultat.json").read_text(encoding="utf-8"))["statut"] == "OK"


def test_userforms_refuses(projet):
    (projet / SOURCE_DIRNAME / "Saisie.frm").write_text("x")
    with pytest.raises(ValueError, match="UserForms"):
        executer(projet, "Module1.Lancer", "tests/fixtures/classeur.xlsx", [], projet / "r", ["Module1.Lancer"],
                 lanceur=lambda p, d: "OK")


def test_execution_dans_excel_sequence(tmp_path):
    excel = EnregistreurCOM("excel", reponses={})
    parametres = {"fixture": "C:/f.xlsx", "classeur": "C:/t/classeur_reference.xlsm", "modules": ["a.bas", "h.bas"],
                  "macro": "Module1.Lancer", "journal": "C:/t/journal.txt", "reponses": ["oui", "Nord"], "pid": "p"}
    executer_dans_excel(parametres, excel=excel)
    lignes = excel.instructions()
    assert "excel.AutomationSecurity = 1" in lignes and "excel.Visible = False" in lignes
    assert lignes[3] == "excel.Workbooks.Open('C:/f.xlsx')"
    assert lignes[4] == "excel.Workbooks.Open('C:/f.xlsx').SaveAs('C:/t/classeur_reference.xlsm', 52)"
    assert [l[l.index(".Import(") + 1:] for l in lignes if ".Import(" in l] == ["Import('a.bas')", "Import('h.bas')"]
    run = next(l for l in lignes if l.startswith("excel.Run("))
    assert run == ("excel.Run(\"'classeur_reference.xlsm'!vba2py_Harnais.vba2py_Lancer\", 'Module1.Lancer', "
                   "'C:/t/journal.txt', 'oui\\x1eNord')")
    assert lignes[-2].endswith(".Save()") and lignes[-1].endswith(".Close(False)")


def test_acces_au_projet_vba_refuse():
    class Classeur:
        def SaveAs(self, *a):
            pass

        @property
        def VBProject(self):
            raise RuntimeError("accès refusé")

        def Close(self, *a):
            pass

    class Excel:
        Visible = DisplayAlerts = AutomationSecurity = None

        class Workbooks:
            @staticmethod
            def Open(chemin):
                return Classeur()

    statut = executer_dans_excel({"fixture": "f", "classeur": "c.xlsm", "modules": [], "macro": "M.L",
                                  "journal": "j", "reponses": [], "pid": "p"}, excel=Excel())
    assert statut.startswith("ECHEC accès au projet VBA refusé")


@pytest.mark.skipif(not excel_disponible(), reason="Excel réel : Windows avec Excel uniquement")
def test_reference_sur_excel_reel(projet):
    """Exécution réelle : dialogue, Outlook et chemin absolu neutralisés, sans blocage."""
    from openpyxl import Workbook

    Workbook().save(projet / "tests" / "fixtures" / "classeur.xlsx")
    vba = '''Attribute VB_Name = "Module1"
Public Sub Lancer()
    Dim ol As Object, f As Integer
    If MsgBox("Continuer ?", vbYesNo) = vbNo Then Exit Sub
    ThisWorkbook.Worksheets(1).Range("A1").Value = InputBox("Valeur ?", "T", "defaut")
    On Error Resume Next
    Set ol = CreateObject("Outlook.Application")
    ThisWorkbook.Worksheets(1).Range("A2").Value = Err.Number
    On Error GoTo 0
    MkDir "C:\\vba2py_test_reference"
    f = FreeFile
    Open "C:\\vba2py_test_reference\\sortie.txt" For Output As #f
    Print #f, "bonjour"
    Close #f
End Sub
'''
    (projet / SOURCE_DIRNAME / "Module1.bas").write_text(vba, encoding="utf-8")
    resultat = executer(projet, "Module1.Lancer", "tests/fixtures/classeur.xlsx", ["oui", "saisi"],
                        projet / "tests" / "references" / "reel", ["Module1.Lancer"])
    assert resultat.statut == "OK", resultat.journal
    assert not Path("C:/vba2py_test_reference").exists()  # redirigé dans le bac à sable
    assert list((resultat.dossier / "fichiers_ecrits").rglob("sortie.txt"))
    from openpyxl import load_workbook
    ws = load_workbook(resultat.dossier / "classeur_apres.xlsm").worksheets[0]
    assert ws["A1"].value == "saisi" and ws["A2"].value != 0  # Outlook neutralisé : erreur interceptée
