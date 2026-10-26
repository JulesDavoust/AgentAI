"""Extraction du code VBA et de la structure des classeurs.

Deux types d'entrée sont acceptés :
- un classeur (.xlsm, .xlsb, .xls, .xlam) : les modules sont décodés avec oletools ;
- un dossier de modules exportés depuis l'éditeur VBA (.bas, .cls, .frm).

Aucune macro n'est exécutée : oletools fait une analyse statique du fichier.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

WORKBOOK_EXTENSIONS = {".xlsm", ".xlsb", ".xls", ".xlam"}
EXPORT_EXTENSIONS = {".bas", ".cls", ".frm"}
OPENPYXL_EXTENSIONS = {".xlsm", ".xlsx", ".xlam"}

SOURCE_DIRNAME = "_source_vba"


class ExtractionError(Exception):
    pass


@dataclass
class Module:
    name: str
    filename: str  # ex. "Module1.bas"
    code: str  # fins de ligne normalisées en "\n"
    form_controls: list[dict] | None = None

    @property
    def kind(self) -> str:
        return {".bas": "module", ".cls": "classe/document", ".frm": "userform"}[
            Path(self.filename).suffix.lower()
        ]

    @property
    def is_empty(self) -> bool:
        return not strip_vba_headers(self.code).strip()


@dataclass
class Project:
    """Un classeur (ou un dossier exporté) = un projet VBA à traduire."""

    source: Path
    relative_path: Path  # chemin de sortie relatif, reproduit l'arborescence d'entrée
    origin: str  # "classeur" ou "export"
    modules: list[Module] = field(default_factory=list)
    workbook_structure: dict | None = None
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- utilitaires


def normalize_newlines(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n")


def decode_export(raw: bytes) -> str:
    """Les exports VBA sont encodés dans la page de code Windows (cp1252 en français)."""
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


_ATTRIBUTE_RE = re.compile(r"^\s*Attribute\s+VB_", re.IGNORECASE)


def strip_vba_headers(code: str) -> str:
    """Retire l'en-tête d'export (VERSION/BEGIN...END) et les lignes `Attribute VB_*`.

    Ces lignes sont des métadonnées de l'éditeur VBA, pas du code exécuté.
    """
    lines = normalize_newlines(code).split("\n")
    if lines and lines[0].strip().upper().startswith("VERSION"):
        # L'en-tête se termine juste avant la première ligne "Attribute VB_Name"
        for i, line in enumerate(lines):
            if _ATTRIBUTE_RE.match(line):
                lines = lines[i:]
                break
    kept = [line.rstrip() for line in lines if not _ATTRIBUTE_RE.match(line)]
    while kept and not kept[-1]:
        kept.pop()
    while kept and not kept[0]:
        kept.pop(0)
    return "\n".join(kept)


# --------------------------------------------------------------------------- classeurs


def extract_from_workbook(path: Path) -> tuple[list[Module], list[str]]:
    from oletools.olevba import VBA_Parser

    warnings: list[str] = []
    try:
        parser = VBA_Parser(str(path))
    except Exception as exc:  # fichier corrompu ou format inconnu
        raise ExtractionError(f"{path.name} : lecture impossible ({exc})") from exc

    try:
        if not parser.detect_vba_macros():
            raise ExtractionError(f"{path.name} : aucune macro VBA trouvée")

        modules: dict[str, Module] = {}
        for _, _, vba_filename, vba_code in parser.extract_all_macros():
            if vba_filename in modules:
                continue
            modules[vba_filename] = Module(
                name=Path(vba_filename).stem,
                filename=vba_filename,
                code=normalize_newlines(vba_code or ""),
            )

        _attach_form_controls(parser, modules, warnings)

        try:
            if parser.detect_vba_stomping():
                warnings.append(
                    "VBA stomping détecté : le code source lisible peut différer du code "
                    "réellement exécuté par Excel. Préférez un export depuis l'éditeur VBA."
                )
        except Exception:
            pass
    finally:
        parser.close()

    return list(modules.values()), warnings


def _attach_form_controls(parser, modules: dict[str, Module], warnings: list[str]) -> None:
    """Récupère au mieux la liste des contrôles des UserForms (stockés en binaire)."""
    forms = [m for m in modules.values() if m.filename.lower().endswith(".frm")]
    if not forms:
        return
    try:
        controls: dict[str, list[dict]] = {}
        for _, storage, variable in parser.extract_form_strings_extended():
            form_name = storage.split("/")[0]
            controls.setdefault(form_name, []).append(
                {k: v for k, v in variable.items() if isinstance(v, (str, int, float, bool))}
            )
        for form in forms:
            form.form_controls = controls.get(form.name)
    except Exception as exc:
        warnings.append(f"Contrôles des UserForms non lus ({exc}).")
    if any(f.form_controls is None for f in forms):
        warnings.append(
            "Disposition des UserForms incomplète : pour une traduction fidèle des formulaires, "
            "exportez-les depuis l'éditeur VBA (le .frm exporté décrit tous les contrôles)."
        )


# --------------------------------------------------------------------------- dossiers exportés


def extract_from_export_dir(directory: Path) -> list[Module]:
    modules = []
    for file in sorted(directory.iterdir()):
        if file.is_file() and file.suffix.lower() in EXPORT_EXTENSIONS:
            modules.append(
                Module(
                    name=file.stem,
                    filename=file.name,
                    code=normalize_newlines(decode_export(file.read_bytes())),
                )
            )
    return modules


# --------------------------------------------------------------------------- structure du classeur


def read_workbook_structure(path: Path, max_rows: int = 20, max_formulas: int = 30) -> dict:
    """Décrit la forme du classeur (feuilles, en-têtes, types) sans copier ses valeurs."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=False, data_only=False, keep_vba=False)
    sheets = []
    formulas = []
    for ws in wb.worksheets:
        header_row = None
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10)):
            if any(c.value is not None for c in row):
                header_row = row
                break

        columns = []
        if header_row is not None:
            first_data_row = header_row[0].row + 1
            for cell in header_row:
                if cell.value is None:
                    continue
                types = set()
                for (value,) in ws.iter_rows(
                    min_row=first_data_row,
                    max_row=min(ws.max_row, first_data_row + max_rows - 1),
                    min_col=cell.column,
                    max_col=cell.column,
                    values_only=True,
                ):
                    if value is not None:
                        types.add(_value_type(value))
                columns.append(
                    {
                        "colonne": cell.column_letter,
                        "en_tete": str(cell.value),
                        "types_observes": sorted(types),
                    }
                )

        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    if len(formulas) < max_formulas:
                        formulas.append({"feuille": ws.title, "cellule": cell.coordinate, "formule": cell.value})

        sheets.append(
            {
                "nom": ws.title,
                "visibilite": {"visible": "visible", "hidden": "masquée", "veryHidden": "très masquée"}.get(
                    ws.sheet_state, ws.sheet_state
                ),
                "dimensions": ws.dimensions,
                "ligne_en_tete": header_row[0].row if header_row else None,
                "colonnes": columns,
                "tableaux": list(ws.tables.keys()),
            }
        )

    named_ranges = []
    for name, defined in wb.defined_names.items():
        named_ranges.append({"nom": name, "reference": defined.attr_text})

    return {
        "fichier": path.name,
        "feuilles": sheets,
        "boutons": lire_boutons(path),
        "plages_nommees": named_ranges,
        "formules_exemples": formulas,
        "note": "Structure uniquement : aucune valeur du classeur n'est copiée.",
    }


_MACRO_RE = re.compile(r"(?:\[\d+\]!|'[^']*'!|[^!'\"]*!)?([^!'\"]+)$")


def _nom_macro(reference: str) -> str:
    """« [0]!Module.Macro » ou « 'Classeur.xlsm'!Macro » → « Module.Macro » / « Macro »."""
    m = _MACRO_RE.search(reference.strip())
    return m.group(1) if m else reference.strip()


def _texte_xml(fragment: str) -> str:
    import html

    texte = re.sub(r"<[^>]+>", " ", fragment)
    return " ".join(html.unescape(texte).split())


def lire_boutons(path: Path) -> list[dict]:
    """Boutons et formes reliés à une macro (contrôles de formulaire, formes avec « Affecter une macro »).

    Information pour repérer les macros lancées par l'utilisateur : les boutons eux-mêmes ne sont
    pas recréés, les commandes Python les remplacent.
    """
    import posixpath
    import zipfile

    def chemin_cible(base: str, cible: str) -> str:
        if cible.startswith("/"):
            return cible[1:]
        return posixpath.normpath(posixpath.join(posixpath.dirname(base), cible))

    def relations(archive, partie: str) -> list[tuple[str, str]]:
        rels = posixpath.join(posixpath.dirname(partie), "_rels", posixpath.basename(partie) + ".rels")
        if rels not in archive.namelist():
            return []
        xml = archive.read(rels).decode("utf-8", "replace")
        return [
            (m.group("id"), chemin_cible(partie, m.group("cible")))
            for m in re.finditer(r'<Relationship\b(?=[^>]*\bId="(?P<id>[^"]+)")(?=[^>]*\bTarget="(?P<cible>[^"]+)")[^>]*>', xml)
        ]

    boutons: list[dict] = []
    try:
        with zipfile.ZipFile(path) as archive:
            classeur = archive.read("xl/workbook.xml").decode("utf-8", "replace")
            cibles = dict(relations(archive, "xl/workbook.xml"))
            for m in re.finditer(r"<sheet\b[^>]*>", classeur):
                balise = m.group(0)
                nom = re.search(r'name="([^"]*)"', balise)
                rid = re.search(r'r:id="([^"]*)"', balise)
                if not nom or not rid or rid.group(1) not in cibles:
                    continue
                feuille = _texte_xml(nom.group(1))
                for _, partie in relations(archive, cibles[rid.group(1)]):
                    if partie not in archive.namelist() or "drawing" not in partie.lower():
                        continue
                    xml = archive.read(partie).decode("utf-8", "replace")
                    if partie.endswith(".vml"):  # contrôles de formulaire (boutons)
                        for forme in re.findall(r"<v:shape\b.*?</v:shape>", xml, re.DOTALL):
                            macro = re.search(r"<x:FmlaMacro>(.*?)</x:FmlaMacro>", forme, re.DOTALL)
                            if macro:
                                texte = re.search(r"<v:textbox\b.*?</v:textbox>", forme, re.DOTALL)
                                boutons.append({"feuille": feuille, "texte": _texte_xml(texte.group(0)) if texte else "",
                                                "macro": _nom_macro(_texte_xml(macro.group(1)))})
                    else:  # formes DrawingML avec une macro affectée
                        for forme in re.findall(r"<xdr:sp\b.*?</xdr:sp>", xml, re.DOTALL):
                            macro = re.search(r'\bmacro="([^"]+)"', forme)
                            if macro:
                                textes = re.findall(r"<a:t>(.*?)</a:t>", forme, re.DOTALL)
                                boutons.append({"feuille": feuille, "texte": _texte_xml(" ".join(textes)),
                                                "macro": _nom_macro(_texte_xml(macro.group(1)))})
    except (zipfile.BadZipFile, KeyError):
        return []
    return boutons


def _value_type(value) -> str:
    import datetime

    if isinstance(value, bool):
        return "booleen"
    if isinstance(value, int):
        return "entier"
    if isinstance(value, float):
        return "decimal"
    if isinstance(value, (datetime.date, datetime.datetime, datetime.time)):
        return "date"
    if isinstance(value, str) and value.startswith("="):
        return "formule"
    return "texte"


def _safe_structure(path: Path, warnings: list[str]) -> dict | None:
    if path.suffix.lower() not in OPENPYXL_EXTENSIONS:
        warnings.append(f"Structure de {path.name} non lue (format {path.suffix} non pris en charge par openpyxl).")
        return None
    try:
        return read_workbook_structure(path)
    except Exception as exc:
        warnings.append(f"Structure de {path.name} non lue ({exc}).")
        return None


# --------------------------------------------------------------------------- découverte


def discover_projects(input_path: Path) -> list[Project]:
    """Trouve les projets VBA dans un fichier ou un dossier, en conservant l'arborescence.

    - Un dossier contenant des .bas/.cls/.frm est un projet « export ». Un classeur posé
      dans le même dossier sert seulement à lire la structure des feuilles.
    - Un classeur situé ailleurs est un projet « classeur » : sortie dans <dossier>/<nom>/.
    """
    input_path = input_path.resolve()
    if not input_path.exists():
        raise ExtractionError(f"{input_path} n'existe pas")

    if input_path.is_file():
        if input_path.suffix.lower() not in WORKBOOK_EXTENSIONS:
            raise ExtractionError(f"{input_path.name} : extension non prise en charge")
        return [_workbook_project(input_path, Path(_slug(input_path.stem)))]

    projects = []
    directories = [input_path] + sorted(
        p
        for p in input_path.rglob("*")
        if p.is_dir() and not any(part.startswith((".", "_")) for part in p.relative_to(input_path).parts)
    )
    for directory in directories:
        files = [f for f in sorted(directory.iterdir()) if f.is_file()]
        exported = [f for f in files if f.suffix.lower() in EXPORT_EXTENSIONS]
        workbooks = [f for f in files if f.suffix.lower() in WORKBOOK_EXTENSIONS | {".xlsx"}]
        rel_dir = directory.relative_to(input_path)

        if exported:
            project = Project(
                source=directory,
                relative_path=_slug_path(rel_dir) if rel_dir.parts else Path(_slug(directory.name)),
                origin="export",
                modules=extract_from_export_dir(directory),
            )
            if len(workbooks) == 1:
                project.workbook_structure = _safe_structure(workbooks[0], project.warnings)
            elif len(workbooks) > 1:
                project.warnings.append("Plusieurs classeurs dans le dossier exporté : structure non lue.")
            else:
                project.warnings.append(
                    "Aucun classeur à côté des modules exportés : la structure des feuilles est inconnue, "
                    "les données fictives seront déduites du code."
                )
            projects.append(project)
        else:
            for wb in workbooks:
                if wb.suffix.lower() == ".xlsx":
                    continue  # pas de macros possibles
                projects.append(_workbook_project(wb, _slug_path(rel_dir) / _slug(wb.stem)))

    if not projects:
        raise ExtractionError(f"Aucun classeur ni module VBA trouvé dans {input_path}")
    return projects


def _workbook_project(path: Path, relative_path: Path) -> Project:
    modules, warnings = extract_from_workbook(path)
    project = Project(source=path, relative_path=relative_path, origin="classeur", modules=modules, warnings=warnings)
    project.workbook_structure = _safe_structure(path, project.warnings)
    return project


def _slug(name: str) -> str:
    """Nom de dossier utilisable comme package Python : « Ventes 2024 » -> « ventes_2024 »."""
    import unicodedata

    ascii_name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^0-9a-zA-Z]+", "_", ascii_name).strip("_").lower()
    if not slug or slug[0].isdigit():
        slug = f"p_{slug}"
    return slug


def _slug_path(path: Path) -> Path:
    return Path(*[_slug(part) for part in path.parts]) if path.parts else Path()


# --------------------------------------------------------------------------- écriture


def write_project_sources(project: Project, out_root: Path) -> Path:
    """Écrit les modules et la structure dans <sortie>/<projet>/_source_vba/ (UTF-8)."""
    project_dir = out_root / project.relative_path
    source_dir = project_dir / SOURCE_DIRNAME
    source_dir.mkdir(parents=True, exist_ok=True)

    for module in project.modules:
        if module.is_empty:
            continue
        (source_dir / module.filename).write_text(module.code, encoding="utf-8")
        if module.form_controls:
            (source_dir / f"{module.name}.controls.json").write_text(
                json.dumps(module.form_controls, ensure_ascii=False, indent=2), encoding="utf-8"
            )

    if project.workbook_structure:
        (source_dir / "structure_classeur.json").write_text(
            json.dumps(project.workbook_structure, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
        )

    (source_dir / "EXTRACTION.md").write_text(extraction_summary(project), encoding="utf-8")
    return project_dir


def extraction_summary(project: Project) -> str:
    lines = [
        f"# Extraction : {project.source.name}",
        "",
        f"- Origine : {project.origin}",
        f"- Dossier de sortie : `{project.relative_path.as_posix()}`",
        "",
        "| Module | Type | Lignes de code | Statut |",
        "|---|---|---|---|",
    ]
    for m in project.modules:
        n = len(strip_vba_headers(m.code).splitlines())
        status = "vide (ignoré)" if m.is_empty else "extrait"
        lines.append(f"| {m.name} | {m.kind} | {n} | {status} |")
    structure = project.workbook_structure or {}
    if structure.get("boutons"):
        lines += ["", "## Macros reliées à des boutons", "", "| Feuille | Bouton | Macro |", "|---|---|---|"]
        lines += [f"| {b['feuille']} | {b['texte']} | {b['macro']} |" for b in structure["boutons"]]
    masquees = [f["nom"] for f in structure.get("feuilles", []) if f.get("visibilite", "visible") != "visible"]
    if masquees:
        lines += ["", f"Feuilles masquées : {', '.join(masquees)}"]
    if project.warnings:
        lines += ["", "## Avertissements", ""] + [f"- {w}" for w in project.warnings]
    return "\n".join(lines) + "\n"
