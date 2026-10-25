"""Tableau croisé dynamique et graphique croisé dynamique, sans Excel (openpyxl).

Module fourni par vba2py et copié tel quel dans chaque projet traduit : ne pas le modifier.

openpyxl ne propose pas de fonction pour créer un tableau croisé dynamique. Ce module écrit
directement sa définition, comme le ferait Excel :
- s'il existe déjà dans le classeur, il est **conservé** et marqué « à recalculer à
  l'ouverture » (Excel le met à jour à partir de la table source) ;
- sinon il est **créé** : définition complète, cache des données, valeurs déjà calculées dans
  les cellules (visibles même en Mode protégé), recalcul à l'ouverture.

Formes prises en charge à la création : un champ en lignes, zéro ou un champ en colonnes,
une valeur agrégée (somme, nombre, moyenne, min, max). Pour une autre forme, une erreur
`FormeNonPriseEnCharge` est levée : la signaler dans le rapport.

Correspondance VBA :
    PivotCaches.Create(xlDatabase, "tblDonnees").CreatePivotTable(Range("A5"), "TCD_Volume")
    .PivotFields("Date").Orientation = xlRowField
    .PivotFields("Region").Orientation = xlColumnField
    .AddDataField .PivotFields("Volume_GWh"), "Volume (GWh)", xlSum
→   assurer_tableau_croise(wb, feuille="Feuille1", nom="TCD_Volume", table_source="tblDonnees",
                           ancre="A5", ligne="Date", colonne="Region", valeur="Volume_GWh",
                           fonction="sum", libelle_valeur="Volume (GWh)")
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field

from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.pivot import PivotSource
from openpyxl.chart.series import SeriesLabel
from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, SharedItems, WorksheetSource
from openpyxl.pivot.fields import DateTimeField, Index, Missing, Number, Text
from openpyxl.pivot.record import Record, RecordList
from openpyxl.pivot.table import (
    DataField,
    FieldItem,
    Location,
    PivotField,
    PivotTableStyle,
    RowColField,
    RowColItem,
    TableDefinition,
)
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

FONCTIONS = {"sum": sum, "count": len, "average": lambda v: sum(v) / len(v), "min": min, "max": max}
LIBELLES_FONCTION = {"sum": "Somme", "count": "Nombre", "average": "Moyenne", "min": "Min", "max": "Max"}
TOTAL = "Total général"
ETIQUETTES_LIGNES = "Étiquettes de lignes"
ETIQUETTES_COLONNES = "Étiquettes de colonnes"


class FormeNonPriseEnCharge(ValueError):
    pass


@dataclass
class ResultatTableauCroise:
    feuille: str
    nom: str
    action: str  # "cree" ou "conserve"
    tableau: TableDefinition
    avertissements: list[str] = field(default_factory=list)

    @property
    def plage(self) -> str:
        return self.tableau.location.ref


# --------------------------------------------------------------------------- lecture


def trouver_tableau_croise(wb, nom: str):
    for ws in wb.worksheets:
        for pivot in ws._pivots:
            if pivot.name == nom:
                return ws, pivot
    return None, None


def lire_table(wb, nom_table: str) -> tuple[list[str], list[list]]:
    """En-têtes et lignes d'un tableau structuré (ListObject)."""
    for ws in wb.worksheets:
        if nom_table in ws.tables:
            min_col, min_row, max_col, max_row = range_boundaries(ws.tables[nom_table].ref)
            lignes = [list(r) for r in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True)]
            entetes = [str(h) for h in lignes[0]]
            donnees = [l for l in lignes[1:] if any(v is not None for v in l)]
            return entetes, donnees
    raise KeyError(f"Tableau « {nom_table} » introuvable dans le classeur")


# --------------------------------------------------------------------------- tableau croisé


def assurer_tableau_croise(
    wb,
    *,
    feuille: str,
    nom: str,
    table_source: str,
    ancre: str,
    ligne: str,
    valeur: str,
    colonne: str | None = None,
    fonction: str = "sum",
    libelle_valeur: str | None = None,
    format_valeurs: str | None = None,
    format_lignes: str | None = None,
    style: str = "PivotStyleLight16",
) -> ResultatTableauCroise:
    """Conserve le tableau croisé `nom` s'il existe, sinon le crée sur `feuille` en `ancre`."""
    ws_existant, existant = trouver_tableau_croise(wb, nom)
    if existant is not None:
        return _conserver(ws_existant, existant, table_source)
    if fonction not in FONCTIONS:
        raise FormeNonPriseEnCharge(f"fonction d'agrégation non prise en charge : {fonction}")
    ws = wb[feuille]
    entetes, donnees = lire_table(wb, table_source)
    for champ in [ligne, valeur] + ([colonne] if colonne else []):
        if champ not in entetes:
            raise KeyError(f"Champ « {champ} » absent du tableau « {table_source} » ({entetes})")
    libelle_valeur = libelle_valeur or f"{LIBELLES_FONCTION[fonction]} de {valeur}"
    pivot = _creer(wb, ws, nom, table_source, ancre, entetes, donnees, ligne, colonne, valeur,
                   fonction, libelle_valeur, format_valeurs, format_lignes, style)
    ws._pivots.append(pivot)
    return ResultatTableauCroise(feuille=ws.title, nom=nom, action="cree", tableau=pivot)


def _conserver(ws, pivot, table_source: str) -> ResultatTableauCroise:
    avertissements = []
    source = pivot.cache.cacheSource
    if source is not None and source.worksheetSource is not None:
        source.worksheetSource.name = table_source
        source.worksheetSource.ref = None
        source.worksheetSource.sheet = None
    else:
        avertissements.append("source du tableau croisé non modifiable (source externe) : laissée telle quelle")
    pivot.cache.refreshOnLoad = True
    avertissements.append(
        "tableau croisé existant conservé : ses valeurs sont recalculées par Excel à l'ouverture "
        "(en Mode protégé, cliquer sur « Activer la modification »)"
    )
    return ResultatTableauCroise(feuille=ws.title, nom=pivot.name, action="conserve", tableau=pivot, avertissements=avertissements)


def _cle_tri(v):
    # Ordre d'Excel : nombres et dates croissants, puis textes (insensible à la casse), vides à la fin
    if v is None:
        return (3, "")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return (0, v)
    if isinstance(v, (dt.date, dt.datetime)):
        return (1, dt.datetime.combine(v, dt.time()) if not isinstance(v, dt.datetime) else v)
    return (2, str(v).lower())


def _uniques(valeurs):
    vus, ordre = set(), []
    for v in valeurs:
        if v not in vus:
            vus.add(v)
            ordre.append(v)
    return ordre


def _item(v):
    if v is None:
        return Missing()
    if isinstance(v, bool):
        return Text(v=str(v))
    if isinstance(v, (int, float)):
        return Number(v=v)
    if isinstance(v, (dt.date, dt.datetime)):
        return DateTimeField(v=v if isinstance(v, dt.datetime) else dt.datetime.combine(v, dt.time()))
    return Text(v=str(v))


def _shared_items(valeurs, axe: bool) -> SharedItems:
    non_vides = [v for v in valeurs if v is not None]
    nombres = [v for v in non_vides if isinstance(v, (int, float)) and not isinstance(v, bool)]
    dates = [v for v in non_vides if isinstance(v, (dt.date, dt.datetime))]
    kwargs = {}
    if non_vides and len(dates) == len(non_vides):
        dates_dt = [d if isinstance(d, dt.datetime) else dt.datetime.combine(d, dt.time()) for d in dates]
        kwargs = dict(containsSemiMixedTypes=False, containsNonDate=False, containsDate=True, containsString=False,
                      minDate=min(dates_dt), maxDate=max(dates_dt))
    elif non_vides and len(nombres) == len(non_vides):
        kwargs = dict(containsSemiMixedTypes=False, containsString=False, containsNumber=True,
                      minValue=min(nombres), maxValue=max(nombres))
        if all(float(n).is_integer() for n in nombres):
            kwargs["containsInteger"] = True
    if len(non_vides) < len(valeurs):
        kwargs["containsBlank"] = True
    if not axe:
        return SharedItems(**kwargs)
    uniques = _uniques(valeurs)
    return SharedItems(_fields=[_item(v) for v in uniques], count=len(uniques), **kwargs)


def _creer(wb, ws, nom, table_source, ancre, entetes, donnees, ligne, colonne, valeur,
           fonction, libelle_valeur, format_valeurs, format_lignes, style) -> TableDefinition:
    i_ligne, i_valeur = entetes.index(ligne), entetes.index(valeur)
    i_colonne = entetes.index(colonne) if colonne else None
    axes = {i_ligne} | ({i_colonne} if colonne else set())

    # --- cache : une entrée par colonne de la table source, comme Excel
    colonnes = list(zip(*donnees)) if donnees else [() for _ in entetes]
    uniques = {i: _uniques(colonnes[i]) for i in axes}
    cache_fields = []
    for i, nom_champ in enumerate(entetes):
        est_date = bool(colonnes[i]) and all(isinstance(v, (dt.date, dt.datetime)) for v in colonnes[i] if v is not None)
        cache_fields.append(CacheField(name=nom_champ, numFmtId=14 if est_date else 0,
                                       sharedItems=_shared_items(list(colonnes[i]), axe=i in axes)))
    records = RecordList(r=[
        Record(_fields=[Index(v=uniques[i].index(v)) if i in axes else _item(v) for i, v in enumerate(l)])
        for l in donnees
    ])
    cache = CacheDefinition(
        refreshOnLoad=True,
        recordCount=len(donnees),
        createdVersion=8, refreshedVersion=8, minRefreshableVersion=3,
        cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(name=table_source)),
        cacheFields=cache_fields,
    )
    cache.records = records

    # --- calcul des valeurs (identique à ce qu'affichera Excel)
    lignes_tri = sorted(uniques[i_ligne], key=_cle_tri)
    colonnes_tri = sorted(uniques[i_colonne], key=_cle_tri) if colonne else []
    agreger = FONCTIONS[fonction]
    groupes: dict = {}
    for l in donnees:
        v = l[i_valeur]
        if fonction != "count" and not isinstance(v, (int, float)):
            continue
        cle_c = l[i_colonne] if colonne else None
        for cle in [(l[i_ligne], cle_c), (l[i_ligne], TOTAL), (TOTAL, cle_c), (TOTAL, TOTAL)]:
            groupes.setdefault(cle, []).append(v)

    def cellule(cle):
        vals = groupes.get(cle)
        return agreger(vals) if vals else None

    # --- disposition (forme compacte d'Excel)
    col0, row0 = coordinate_from_string(ancre)
    c0 = column_index_from_string(col0)
    if colonne:
        entete_lignes = 2
        largeur = 1 + len(colonnes_tri) + 1
        ws.cell(row=row0, column=c0, value=libelle_valeur)
        ws.cell(row=row0, column=c0 + 1, value=ETIQUETTES_COLONNES)
        ws.cell(row=row0 + 1, column=c0, value=ETIQUETTES_LIGNES)
        for j, c in enumerate(colonnes_tri + [TOTAL]):
            ws.cell(row=row0 + 1, column=c0 + 1 + j, value=c)
        cles_col = colonnes_tri + [TOTAL]
    else:
        entete_lignes = 1
        largeur = 2
        ws.cell(row=row0, column=c0, value=ETIQUETTES_LIGNES)
        ws.cell(row=row0, column=c0 + 1, value=libelle_valeur)
        cles_col = [TOTAL]
    for k, r in enumerate(lignes_tri + [TOTAL]):
        rr = row0 + entete_lignes + k
        etiquette = ws.cell(row=rr, column=c0, value=r)
        if format_lignes and r != TOTAL:
            etiquette.number_format = format_lignes
        for j, c in enumerate(cles_col):
            cle = (r, c) if colonne else (r, TOTAL)
            cel = ws.cell(row=rr, column=c0 + 1 + j, value=cellule(cle))
            if format_valeurs:
                cel.number_format = format_valeurs
    hauteur = entete_lignes + len(lignes_tri) + 1
    ref = f"{ancre}:{get_column_letter(c0 + largeur - 1)}{row0 + hauteur - 1}"

    # --- définition du tableau croisé
    def items_axe(i):
        tri = sorted(uniques[i], key=_cle_tri)
        return [FieldItem(x=uniques[i].index(v)) for v in tri] + [FieldItem(t="default")]

    pivot_fields = []
    for i in range(len(entetes)):
        if i == i_ligne:
            pivot_fields.append(PivotField(axis="axisRow", showAll=False, items=items_axe(i)))
        elif colonne and i == i_colonne:
            pivot_fields.append(PivotField(axis="axisCol", showAll=False, items=items_axe(i)))
        elif i == i_valeur:
            pivot_fields.append(PivotField(dataField=True, showAll=False))
        else:
            pivot_fields.append(PivotField(showAll=False))

    def items_lignes(n):
        return [RowColItem(x=[Index(v=k)]) for k in range(n)] + [RowColItem(t="grand", x=[Index(v=0)])]

    cache_ids = [p.cacheId for w in wb.worksheets for p in w._pivots]
    pivot = TableDefinition(
        name=nom,
        cacheId=max(cache_ids, default=-1) + 1,
        dataCaption="Valeurs",
        applyNumberFormats=False, applyBorderFormats=False, applyFontFormats=False,
        applyPatternFormats=False, applyAlignmentFormats=False, applyWidthHeightFormats=True,
        updatedVersion=8, minRefreshableVersion=3, createdVersion=8,
        useAutoFormatting=True, itemPrintTitles=True, indent=0, outline=True, outlineData=True,
        multipleFieldFilters=False,
        location=Location(ref=ref, firstHeaderRow=1, firstDataRow=2 if colonne else 1, firstDataCol=1),
        pivotFields=pivot_fields,
        rowFields=[RowColField(x=i_ligne)],
        rowItems=items_lignes(len(lignes_tri)),
        colFields=[RowColField(x=i_colonne)] if colonne else [],
        colItems=items_lignes(len(colonnes_tri)) if colonne else [RowColItem()],
        dataFields=[DataField(name=libelle_valeur, fld=i_valeur, subtotal=fonction, baseField=0, baseItem=0)],
        pivotTableStyleInfo=PivotTableStyle(name=style, showRowHeaders=True, showColHeaders=True,
                                            showRowStripes=False, showColStripes=False, showLastColumn=True),
    )
    pivot.cache = cache
    return pivot


# --------------------------------------------------------------------------- graphique croisé


def graphiques_croises(ws, nom_tableau: str) -> list:
    return [c for c in ws._charts if c.pivotSource is not None and c.pivotSource.name.endswith(f"!{nom_tableau}")]


def assurer_graphique_croise(
    wb,
    resultat: ResultatTableauCroise,
    *,
    ancre: str,
    titre: str | None = None,
    type_graphique: str = "ligne",
    largeur_pt: float = 520,
    hauteur_pt: float = 300,
    nom_fichier: str = "Classeur.xlsx",
):
    """Conserve le graphique croisé lié au tableau s'il existe, sinon en crée un (vrai graphique Excel).

    Correspondance VBA :
        ChartObjects.Add(Left:=Range("J5").Left, Top:=..., Width:=520, Height:=300)
        .Chart.SetSourceData pt.TableRange1 : .ChartType = xlLine : .ChartTitle.Text = "..."
    `nom_fichier` : nom du classeur, utilisé par Excel pour relier le graphique au tableau.
    """
    ws = wb[resultat.feuille]
    existants = graphiques_croises(ws, resultat.nom)
    if existants:
        return existants[0]

    pivot = resultat.tableau
    min_col, min_row, max_col, max_row = range_boundaries(pivot.location.ref)
    ligne_entete = min_row + pivot.location.firstDataRow - 1  # ligne des étiquettes de colonnes
    premiere_donnee = ligne_entete + 1
    derniere_donnee = max_row - 1  # sans la ligne « Total général »
    derniere_colonne = max_col - 1 if pivot.colFields else max_col  # sans la colonne « Total général »

    chart = {"ligne": LineChart, "barres": BarChart}[type_graphique]()
    if titre:
        chart.title = titre
    chart.width = largeur_pt * 2.54 / 72
    chart.height = hauteur_pt * 2.54 / 72
    categories = Reference(ws, min_col=min_col, min_row=premiere_donnee, max_row=derniere_donnee)
    for col in range(min_col + 1, derniere_colonne + 1):
        valeurs = Reference(ws, min_col=col, min_row=premiere_donnee, max_row=derniere_donnee)
        serie = Series(valeurs, title_from_data=False)
        serie.tx = SeriesLabel(strRef=StrRef(f"'{ws.title}'!${get_column_letter(col)}${ligne_entete}"))
        if type_graphique == "ligne":
            serie.marker.symbol = "none"
            serie.smooth = False
        chart.series.append(serie)
    chart.set_categories(categories)
    chart.pivotSource = PivotSource(name=f"[{nom_fichier}]{ws.title}!{resultat.nom}", fmtId=0)
    ws.add_chart(chart, ancre)
    return chart


# --------------------------------------------------------------------------- nettoyage


def effacer_hors_tableaux_croises(ws) -> None:
    """Équivalent de `ws.Cells.Clear` qui épargne les tableaux croisés et graphiques croisés.

    En VBA, la macro efface la feuille puis recrée tableau et graphique. Ici ils sont conservés
    (et recalculés par Excel) : seules les autres cellules sont vidées, ainsi que les images et
    graphiques qui ne sont pas liés à un tableau croisé.
    """
    zones = [range_boundaries(p.location.ref) for p in ws._pivots]

    def dans_un_tableau(cell):
        return any(c1 <= cell.column <= c2 and r1 <= cell.row <= r2 for c1, r1, c2, r2 in zones)

    for row in ws.iter_rows():
        for cell in row:
            if not dans_un_tableau(cell):
                cell.value = None
                cell.style = "Normal"
    ws._images = []
    ws._charts = [c for c in ws._charts if c.pivotSource is not None]
