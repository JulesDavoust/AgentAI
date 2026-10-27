from __future__ import annotations

import random
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

import marco_25_10 as m

MAINTENANT = datetime(2026, 10, 25, 11, 15, 30)
AUJOURDHUI = MAINTENANT.date()


def _copie(chemin: Path, tmp_path: Path):
    wb = load_workbook(chemin)
    return wb


# ---------------------------------------------------------------- simulation

def test_simuler_volume_bornes_par_region():
    rng = random.Random(1)
    for reg, base in {"Nord": 120, "Sud": 95, "Est": 80, "Ouest": 70, "Centre": 60, "Inconnue": 60}.items():
        for _ in range(50):
            v = m.simuler_volume(reg, rng)
            assert base * 0.85 - 0.1 <= v <= base * 1.15 + 0.1  # base x (0,85..1,15), arrondi 0,1
            assert v == round(v, 1)


def test_simuler_temperature_bornes():
    rng = random.Random(2)
    assert all(2 <= m.simuler_temperature(rng) <= 16 for _ in range(100))


# ---------------------------------------------------------------- structure

def test_ensure_structure_cree_historique(classeur_vide, tmp_path):
    wb = load_workbook(classeur_vide)
    cree = m.ensure_structure(wb, AUJOURDHUI, MAINTENANT, rng=random.Random(0))
    assert cree is True
    assert wb.sheetnames == ["Dashboard", "Feuille2", "Feuille1", "Log"]
    ws = wb["Feuille2"]
    assert [c.value for c in ws[1]] == ["Date", "Region", "Volume_GWh", "Temperature_C", "Statut"]
    assert ws.tables["tblDonnees"].ref == "A1:E71"  # 14 jours x 5 régions = 70 lignes + en-tête
    assert ws["A2"].value == date(2026, 10, 11)  # Date - 14
    assert ws["A71"].value == date(2026, 10, 24)  # Date - 1
    assert [ws.cell(row=r, column=2).value for r in range(2, 7)] == list(m.REGIONS)
    assert ws["E2"].value == "Historique"
    assert ws["A2"].number_format == "dd/mm/yyyy"
    assert ws.tables["tblDonnees"].tableStyleInfo.name == "TableStyleMedium2"
    log = wb["Log"]
    assert [c.value for c in log[1]] == ["Horodatage", "Niveau", "Message"]
    assert log["C2"].value == "Structure initiale creee"


def test_ensure_structure_ne_touche_pas_au_tableau_existant(classeur_complet):
    wb = load_workbook(classeur_complet)
    assert m.ensure_structure(wb, AUJOURDHUI, MAINTENANT) is False
    assert wb["Feuille2"].tables["tblDonnees"].ref == "A1:E16"
    assert wb["Feuille2"].max_row == 16


# ---------------------------------------------------------------- données du jour

def test_ajouter_donnees_du_jour_nominal(classeur_complet):
    wb = load_workbook(classeur_complet)
    nb = m.ajouter_donnees_du_jour(wb, AUJOURDHUI, MAINTENANT, rng=random.Random(0))
    assert nb == 5
    ws = wb["Feuille2"]
    assert ws.tables["tblDonnees"].ref == "A1:E21"  # 15 lignes + 5
    assert ws.tables["tblDonnees"].autoFilter.ref == "A1:E21"
    assert ws["A17"].value == AUJOURDHUI and ws["B17"].value == "Nord"
    assert ws["E21"].value == "Phase de test" and ws["B21"].value == "Centre"
    assert wb["Log"]["C3"].value == "5 ligne(s) ajoutee(s) pour le 25/10/2026"


def test_ajouter_donnees_du_jour_idempotent(classeur_complet):
    wb = load_workbook(classeur_complet)
    m.ajouter_donnees_du_jour(wb, AUJOURDHUI, MAINTENANT)
    assert m.ajouter_donnees_du_jour(wb, AUJOURDHUI, MAINTENANT) == 0
    assert wb["Feuille2"].tables["tblDonnees"].ref == "A1:E21"
    assert wb["Log"]["C4"].value == "0 ligne(s) ajoutee(s) pour le 25/10/2026"


def test_ajouter_donnees_du_jour_region_deja_presente(classeur_complet):
    wb = load_workbook(classeur_complet)
    ws = wb["Feuille2"]
    ws.append([datetime(2026, 10, 25, 0, 0), "Sud", 1, 1, "x"])
    ws.tables["tblDonnees"].ref = "A1:E17"
    assert m.ajouter_donnees_du_jour(wb, AUJOURDHUI, MAINTENANT) == 4


def test_ajouter_donnees_sans_tableau(classeur_vide):
    wb = load_workbook(classeur_vide)
    with pytest.raises(KeyError):
        m.ajouter_donnees_du_jour(wb, AUJOURDHUI, MAINTENANT)


def test_ligne_existe_casse_exacte(classeur_complet):
    wb = load_workbook(classeur_complet)
    ws = wb["Feuille2"]
    t = ws.tables["tblDonnees"]
    assert m.ligne_existe(ws, t, date(2026, 10, 21), "Nord")
    assert not m.ligne_existe(ws, t, date(2026, 10, 21), "nord")  # '=' VBA sensible à la casse
    assert not m.ligne_existe(ws, t, date(2026, 10, 25), "Nord")


# ---------------------------------------------------------------- log

def test_log_msg_ajoute_apres_derniere_ligne(classeur_complet):
    wb = load_workbook(classeur_complet)
    m.log_msg(wb, "ERREUR", "boom", MAINTENANT)
    log = wb["Log"]
    assert (log["A3"].value, log["B3"].value, log["C3"].value) == (MAINTENANT, "ERREUR", "boom")
    assert log["A3"].number_format == "dd/mm/yyyy hh:mm:ss"


# ---------------------------------------------------------------- filtre

@pytest.mark.parametrize("saisie,attendu", [
    ("", None), ("   ", None), ("Nord", "Nord"), ("  est ", "est"), ("OUEST", "OUEST"), ("Paris", None),
])
def test_choisir_filtre(saisie, attendu):
    assert m.choisir_filtre(saisie) == attendu


def test_est_region_valide():
    assert m.est_region_valide("centre") and not m.est_region_valide("Nord-Est")


# ---------------------------------------------------------------- sauvegarde / mail

def test_chemin_sauvegarde(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    chemin = m.chemin_sauvegarde("a/b/source.xlsm", "rel/sous", MAINTENANT)
    assert chemin.is_absolute() and chemin.parent.is_dir()
    assert chemin.name == "VBA_Natran_Test_20261025_111530_VBATF.xlsm"
    assert chemin.parent == (tmp_path / "rel" / "sous").resolve()


def test_construire_mail():
    objet, corps = m.construire_mail("C:\\x\\y\\fichier.xlsm", AUJOURDHUI)
    assert objet == "[PHASE DE TEST] Mise a jour suivi distribution - 25/10/2026"
    assert "Fichier : fichier.xlsm\r\n" in corps
    assert corps.startswith("Bonjour,\r\n\r\n") and corps.endswith("Robot VBATF")


def _faux_mail(appels):
    def faux(destinataire, cc, objet, corps, piece_jointe, envoyer):
        assert Path(piece_jointe).is_absolute() and Path(piece_jointe).exists()
        appels.append((destinataire, cc, objet, corps, piece_jointe, envoyer))
    return faux


@pytest.mark.parametrize("mode,envoyer", [("envoyer", True), ("afficher", False)])
def test_envoyer_mail_suivi_modes(tmp_path, mode, envoyer):
    pj = tmp_path / "f.xlsx"
    pj.write_bytes(b"x")
    appels = []
    assert m.envoyer_mail_suivi(pj, AUJOURDHUI, mode_mail=mode, envoyer_mail=_faux_mail(appels))
    assert len(appels) == 1
    assert appels[0][0] == m.MAIL_TO and appels[0][1] == "" and appels[0][5] is envoyer


def test_envoyer_mail_suivi_aucun(tmp_path):
    appels = []
    assert not m.envoyer_mail_suivi(tmp_path / "f", AUJOURDHUI, mode_mail="aucun", envoyer_mail=_faux_mail(appels))
    assert appels == []


def test_envoyer_mail_suivi_erreur_non_bloquante(tmp_path, caplog):
    pj = tmp_path / "f.xlsx"
    pj.write_bytes(b"x")

    def casse(*a):
        raise OSError("Outlook absent")

    assert m.envoyer_mail_suivi(pj, AUJOURDHUI, envoyer_mail=casse) is False
    assert "Mail impossible" in caplog.text


# ---------------------------------------------------------------- pipeline

def _etapes(appels):
    def faux(chemin, **kwargs):
        assert Path(chemin).is_absolute() and Path(chemin).exists()  # fichier enregistré avant Excel
        appels.append(("excel", chemin, kwargs))
    return faux


def test_pipeline_complet(classeur_complet, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)  # dossier de sortie relatif
    ordre = []
    etapes = _etapes(ordre)

    def mail(*args):
        assert ordre and ordre[0][0] == "excel"  # mail après les étapes Excel
        ordre.append(("mail", args))
        assert Path(args[4]).is_absolute() and Path(args[4]).exists()

    chemin = m.lancer_pipeline_complet(
        classeur_complet, "out", saisie="sud", mode_mail="afficher", horloge=lambda: MAINTENANT,
        rng=random.Random(0), envoyer_mail=mail, etapes_excel=etapes)

    assert chemin == (tmp_path / "out" / "VBA_Natran_Test_20261025_111530_VBATF.xlsx").resolve()
    _, ch, kw = ordre[0]
    assert ch == chemin
    assert kw == dict(sheet_visuel="Feuille1", sheet_data="Feuille2", table_name="tblDonnees",
                      saisie_filtre="sud", structure_creee=False, maintenant=MAINTENANT)
    assert ordre[1][0] == "mail" and ordre[1][1][0] == m.MAIL_TO and ordre[1][1][5] is False
    wb = load_workbook(chemin)
    assert wb["Feuille2"].tables["tblDonnees"].ref == "A1:E21"
    log = [r[2].value for r in wb["Log"].iter_rows(min_row=2)]
    assert log == ["Structure initiale creee", "5 ligne(s) ajoutee(s) pour le 25/10/2026",
                   "Filtre applique : sud", "Visuel regenere"]
    assert wb["Feuille1"]["A1"].value is None  # nettoyé avant recréation par Excel
    # le classeur source n'est pas modifié (SaveCopyAs)
    assert load_workbook(classeur_complet)["Feuille2"].max_row == 16


def test_pipeline_classeur_vide_et_sans_mail(classeur_vide, tmp_path):
    ordre = []
    chemin = m.lancer_pipeline_complet(
        classeur_vide, tmp_path / "s", mode_mail="aucun", horloge=lambda: MAINTENANT,
        envoyer_mail=lambda *a: pytest.fail("mail inattendu"), etapes_excel=_etapes(ordre))
    assert chemin is not None
    kw = ordre[0][2]
    assert kw["structure_creee"] is True and kw["saisie_filtre"] is None
    assert load_workbook(chemin)["Feuille2"].tables["tblDonnees"].ref == "A1:E76"  # 75 lignes (cf. structure)


def test_pipeline_erreur_retourne_none(tmp_path, caplog):
    assert m.lancer_pipeline_complet(tmp_path / "absent.xlsx", tmp_path, etapes_excel=lambda *a, **k: None) is None
    assert "ERREUR" in caplog.text


def test_pipeline_defaut_bloque_par_conftest(classeur_complet, tmp_path):
    """Sans fausse fonction, les connecteurs réels sont bloqués : l'erreur est journalisée."""
    assert m.lancer_pipeline_complet(classeur_complet, tmp_path, horloge=lambda: MAINTENANT) is None


def test_main_aide(capsys):
    with pytest.raises(SystemExit):
        m.main(["--help"])
    sortie = capsys.readouterr().out
    for option in ("--dossier-sortie", "--mode-mail", "--saisie", "--mail-to", "--table-name"):
        assert option in sortie
