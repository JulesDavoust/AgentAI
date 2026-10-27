"""Crée les classeurs fictifs de tests/fixtures/ (idempotent : les fichiers sont recréés)."""

from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableStyleInfo

FIXTURES = Path(__file__).resolve().parent / "fixtures"
REGIONS = ("Nord", "Sud", "Est", "Ouest", "Centre")
VOLUMES = {"Nord": 120.0, "Sud": 95.5, "Est": 80, "Ouest": 70.2, "Centre": 60}
JOURS = (date(2026, 10, 21), date(2026, 10, 22), date(2026, 10, 23))


def classeur_complet(chemin: Path) -> None:
    """Dashboard, Feuille2 (tblDonnees, 15 lignes), Feuille1 (visuel), Log."""
    wb = Workbook()
    wb.active.title = "Dashboard"
    ws = wb.create_sheet("Feuille2")
    ws.append(["Date", "Region", "Volume_GWh", "Temperature_C", "Statut"])
    for i, jour in enumerate(JOURS):
        for r in REGIONS:
            ws.append([jour, r, VOLUMES[r] + i, 5.5 + i, "Historique"])
            ws.cell(row=ws.max_row, column=1).number_format = "dd/mm/yyyy"
    ref = f"A1:E{ws.max_row}"
    table = Table(displayName="tblDonnees", ref=ref)
    table.autoFilter = AutoFilter(ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    visuel = wb.create_sheet("Feuille1")
    visuel["A1"] = "Suivi distribution gaz - mis a jour le 24/10/2026 09:00"
    visuel["A2"] = "Volume total (GWh)"
    log = wb.create_sheet("Log")
    log.append(["Horodatage", "Niveau", "Message"])
    log.append([datetime(2026, 10, 24, 9, 0, 0), "INFO", "Structure initiale creee"])
    wb.save(chemin)


def classeur_vide(chemin: Path) -> None:
    """Classeur sans aucune des feuilles de la macro."""
    wb = Workbook()
    wb.active.title = "Dashboard"
    wb.save(chemin)


def generer() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    classeur_complet(FIXTURES / "classeur_complet.xlsx")
    classeur_vide(FIXTURES / "classeur_vide.xlsx")


if __name__ == "__main__":
    generer()
    print(f"Fixtures créées dans {FIXTURES}", file=sys.stderr)
