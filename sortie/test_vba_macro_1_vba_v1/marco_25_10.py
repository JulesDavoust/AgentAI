"""Traduction du module VBA `marco_25_10` (suivi distribution gaz).

Pipeline : structure -> données du jour -> filtre -> visuel -> sauvegarde -> mail.
Les parties qu'openpyxl ne sait pas faire (filtre appliqué, TCD, graphique croisé, AutoFit) sont
dans `connecteurs/excel.py` ; le mail est dans `connecteurs/outlook.py`.
"""

from __future__ import annotations

import argparse
import logging
import random
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableStyleInfo

from bibliotheque.excel_com import preparer_feuille
from connecteurs import excel as connecteur_excel
from connecteurs import outlook as connecteur_outlook

logger = logging.getLogger(__name__)

# --- CONFIGURATION (constantes VBA -> valeurs par défaut de la ligne de commande) ------------
SHEET_VISUEL = "Feuille1"
SHEET_DATA = "Feuille2"
SHEET_LOG = "Log"
TABLE_NAME = "tblDonnees"
FILE_TITLE = "VBA_Natran_Test"
FILE_SUFFIX = "VBATF"
MAIL_TO = "jdavoust201@headmind.com"
MAIL_CC = ""
DOSSIER_SORTIE = "sortie"

REGIONS = ("Nord", "Sud", "Est", "Ouest", "Centre")  # VBA : Regions()
FORMAT_DATE = "dd/mm/yyyy"


# --------------------------------------------------------------------------- outils

def get_or_create_sheet(wb, nom: str):
    """VBA GetOrCreateSheet : la feuille `nom`, créée en dernière position si absente."""
    if nom in wb.sheetnames:
        return wb[nom]
    return wb.create_sheet(nom)


def log_msg(wb, niveau: str, msg: str, maintenant: datetime, sheet_log: str = SHEET_LOG) -> None:
    """VBA LogMsg : ajoute une ligne (horodatage, niveau, message) dans la feuille de log."""
    ws = get_or_create_sheet(wb, sheet_log)
    if ws["A1"].value in (None, ""):
        for col, titre in enumerate(("Horodatage", "Niveau", "Message"), start=1):
            ws.cell(row=1, column=col, value=titre)
    n = 1  # équivalent de Cells(Rows.Count, 1).End(xlUp).Row : dernière ligne non vide de A
    for row in range(ws.max_row, 0, -1):
        if ws.cell(row=row, column=1).value not in (None, ""):
            n = row
            break
    n += 1
    ws.cell(row=n, column=1, value=maintenant).number_format = "dd/mm/yyyy hh:mm:ss"
    ws.cell(row=n, column=2, value=niveau)
    ws.cell(row=n, column=3, value=msg)


def _en_date(valeur) -> date | None:
    """CLng(valeur) côté date : datetime/date -> date, sinon None."""
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    return None


# --------------------------------------------------------------------------- 1. structure / données

def simuler_volume(reg: str, rng: random.Random) -> float:
    """VBA SimulerVolume : volume de base de la région x (0,85 à 1,15), arrondi à 0,1."""
    base = {"Nord": 120, "Sud": 95, "Est": 80, "Ouest": 70}.get(reg, 60)
    return round(base * (0.85 + rng.random() * 0.3), 1)


def simuler_temperature(rng: random.Random) -> float:
    """VBA SimulerTemperature : température entre 2 et 16, arrondie à 0,1."""
    return round(2 + rng.random() * 14, 1)


def ensure_structure(wb, aujourdhui: date, maintenant: datetime, *, sheet_data: str = SHEET_DATA,
                     sheet_visuel: str = SHEET_VISUEL, sheet_log: str = SHEET_LOG,
                     table_name: str = TABLE_NAME, rng: random.Random | None = None) -> bool:
    """VBA EnsureStructure : crée les feuilles et, si le tableau manque, 14 jours d'historique.

    Retourne True si la structure initiale a été créée (l'AutoFit des colonnes se fait alors dans
    Excel).
    """
    rng = rng or random.Random()
    ws_d = get_or_create_sheet(wb, sheet_data)
    get_or_create_sheet(wb, sheet_visuel)

    if table_name in ws_d.tables:
        return False

    # wsD.Cells.Clear
    for row in ws_d.iter_rows():
        for cell in row:
            cell.value = None
            cell.style = "Normal"
    for col, titre in enumerate(("Date", "Region", "Volume_GWh", "Temperature_C", "Statut"), start=1):
        ws_d.cell(row=1, column=col, value=titre)
    ligne = 2
    for i in range(14, 0, -1):
        for r in REGIONS:
            ws_d.cell(row=ligne, column=1, value=aujourdhui - timedelta(days=i))
            ws_d.cell(row=ligne, column=2, value=r)
            ws_d.cell(row=ligne, column=3, value=simuler_volume(r, rng))
            ws_d.cell(row=ligne, column=4, value=simuler_temperature(rng))
            ws_d.cell(row=ligne, column=5, value="Historique")
            ligne += 1
    ref = f"A1:E{ligne - 1}"
    table = Table(displayName=table_name, ref=ref)
    table.autoFilter = AutoFilter(ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws_d.add_table(table)
    ws_d.column_dimensions["A"].number_format = FORMAT_DATE
    for row in range(1, ligne):
        ws_d.cell(row=row, column=1).number_format = FORMAT_DATE
    log_msg(wb, "INFO", "Structure initiale creee", maintenant, sheet_log)
    return True


def ligne_existe(ws, table: Table, d: date, reg: str) -> bool:
    """VBA LigneExiste : une ligne du tableau a-t-elle déjà cette date et cette région ?"""
    _, min_row, _, max_row = range_boundaries(table.ref)
    for row in range(min_row + 1, max_row + 1):
        if _en_date(ws.cell(row=row, column=1).value) == d and ws.cell(row=row, column=2).value == reg:
            return True
    return False


def ajouter_donnees_du_jour(wb, aujourdhui: date, maintenant: datetime, *,
                            sheet_data: str = SHEET_DATA, sheet_log: str = SHEET_LOG,
                            table_name: str = TABLE_NAME, rng: random.Random | None = None) -> int:
    """VBA AjouterDonneesDuJour : une ligne « Phase de test » par région absente pour aujourd'hui.

    Retourne le nombre de lignes ajoutées. Lève KeyError si la feuille ou le tableau manque
    (erreur 9 en VBA).
    """
    rng = rng or random.Random()
    ws = wb[sheet_data]
    table = ws.tables[table_name]
    nb = 0
    for r in REGIONS:
        if not ligne_existe(ws, table, aujourdhui, r):
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            new_row = max_row + 1  # ListRows.Add : juste sous la dernière ligne du tableau
            ws.cell(row=new_row, column=1, value=aujourdhui).number_format = FORMAT_DATE
            ws.cell(row=new_row, column=2, value=r)
            ws.cell(row=new_row, column=3, value=simuler_volume(r, rng))
            ws.cell(row=new_row, column=4, value=simuler_temperature(rng))
            ws.cell(row=new_row, column=5, value="Phase de test")
            nouvelle_ref = f"A{min_row}:E{new_row}"
            table.ref = nouvelle_ref
            if table.autoFilter is not None:
                table.autoFilter.ref = nouvelle_ref
            nb += 1
    log_msg(wb, "INFO", f"{nb} ligne(s) ajoutee(s) pour le {aujourdhui:%d/%m/%Y}", maintenant, sheet_log)
    return nb


# --------------------------------------------------------------------------- 2. filtre

def est_region_valide(s: str) -> bool:
    """VBA EstRegionValide : comparaison insensible à la casse (vbTextCompare)."""
    return any(r.casefold() == s.casefold() for r in REGIONS)


def choisir_filtre(saisie: str) -> str | None:
    """VBA AppliquerFiltre (partie saisie) : région à filtrer, ou None (vide ou inconnue).

    `Trim$` ne retire que les espaces. Une région inconnue donnait un MsgBox : avertissement log.
    """
    saisie = saisie.strip(" ")
    if saisie == "":
        return None
    if not est_region_valide(saisie):
        logger.warning("Region inconnue : %s", saisie)
        return None
    return saisie


# --------------------------------------------------------------------------- 4. sauvegarde

def chemin_sauvegarde(classeur: str | Path, dossier_sortie: str | Path, maintenant: datetime,
                      file_title: str = FILE_TITLE, file_suffix: str = FILE_SUFFIX) -> Path:
    """VBA SauvegarderVersion : crée le dossier et renvoie le chemin absolu titre_date_suffixe.

    L'extension est celle du classeur source (la macro enregistrait une copie `.xlsm` du classeur
    macro).
    """
    dossier = Path(dossier_sortie)
    dossier.mkdir(parents=True, exist_ok=True)  # If Dir(...) = "" Then MkDir
    nom = f"{file_title}_{maintenant:%Y%m%d_%H%M%S}_{file_suffix}{Path(classeur).suffix}"
    return (dossier / nom).resolve()


# --------------------------------------------------------------------------- 5. mail

def construire_mail(chemin_pj: str | Path, aujourdhui: date) -> tuple[str, str]:
    """VBA EnvoyerMail : (objet, corps) du mail."""
    objet = f"[PHASE DE TEST] Mise a jour suivi distribution - {aujourdhui:%d/%m/%Y}"
    nom = str(chemin_pj).replace("\\", "/").rsplit("/", 1)[-1]  # Mid$(chemin, InStrRev(chemin, "\") + 1)
    corps = ("Bonjour,\r\n\r\n"
             "Ceci est un envoi automatique en phase de test.\r\n"
             "Le graphique et les donnees ont ete mis a jour.\r\n"
             f"Fichier : {nom}\r\n\r\n"
             "Cordialement,\r\nRobot VBATF")
    return objet, corps


def envoyer_mail_suivi(chemin_pj: str | Path, aujourdhui: date, *, mail_to: str = MAIL_TO,
                       mail_cc: str = MAIL_CC, mode_mail: str = "envoyer",
                       envoyer_mail: Callable | None = None) -> bool:
    """VBA EnvoyerMail : envoie (`envoyer`), affiche (`afficher`) ou ignore (`aucun`) le mail.

    `On Error GoTo MailErr` : une erreur d'envoi est journalisée et n'interrompt pas le pipeline.
    Retourne True si le mail a été envoyé ou affiché.
    """
    if mode_mail == "aucun":
        return False
    envoyer_mail = envoyer_mail or connecteur_outlook.envoyer_mail
    objet, corps = construire_mail(chemin_pj, aujourdhui)
    envoyer = mode_mail == "envoyer"
    try:
        envoyer_mail(mail_to, mail_cc, objet, corps, str(Path(chemin_pj).resolve()), envoyer)
    except Exception as exc:  # noqa: BLE001 - comme MailErr : on journalise et on continue
        logger.error("Mail impossible : %s", exc)
        return False
    logger.info("%s%s", "Mail envoye a " if envoyer else "Brouillon cree pour ", mail_to)
    return True


# --------------------------------------------------------------------------- pipeline

def lancer_pipeline_complet(classeur: str | Path, dossier_sortie: str | Path = DOSSIER_SORTIE, *,
                            sheet_visuel: str = SHEET_VISUEL, sheet_data: str = SHEET_DATA,
                            sheet_log: str = SHEET_LOG, table_name: str = TABLE_NAME,
                            file_title: str = FILE_TITLE, file_suffix: str = FILE_SUFFIX,
                            mail_to: str = MAIL_TO, mail_cc: str = MAIL_CC,
                            mode_mail: str = "envoyer", saisie: str = "",
                            horloge: Callable[[], datetime] = datetime.now,
                            rng: random.Random | None = None,
                            envoyer_mail: Callable | None = None,
                            etapes_excel: Callable | None = None) -> Path | None:
    """VBA LancerPipelineComplet : exécute tout le pipeline et renvoie le chemin de la copie.

    Le classeur d'origine n'est jamais modifié (la macro utilisait `SaveCopyAs`) : toutes les
    modifications sont enregistrées dans la copie datée. Retourne None en cas d'erreur
    (`ErrHandler` : erreur journalisée, pas de propagation).
    """
    etapes_excel = etapes_excel or connecteur_excel.appliquer_etapes_excel
    classeur = Path(classeur)
    try:
        wb = load_workbook(classeur, keep_vba=classeur.suffix.lower() == ".xlsm")
        maintenant = horloge()
        aujourdhui = maintenant.date()
        noms = dict(sheet_data=sheet_data, sheet_log=sheet_log, table_name=table_name, rng=rng)

        structure_creee = ensure_structure(wb, aujourdhui, maintenant, sheet_visuel=sheet_visuel, **noms)
        ajouter_donnees_du_jour(wb, aujourdhui, maintenant, **noms)

        # Filtre (la saisie remplace l'InputBox) ; appliqué ensuite par Excel
        region = choisir_filtre(saisie)
        if region is not None:
            log_msg(wb, "INFO", f"Filtre applique : {region}", maintenant, sheet_log)

        # Visuel : nettoyage côté openpyxl, création côté Excel
        preparer_feuille(get_or_create_sheet(wb, sheet_visuel))
        log_msg(wb, "INFO", "Visuel regenere", horloge(), sheet_log)

        chemin = chemin_sauvegarde(classeur, dossier_sortie, horloge(), file_title, file_suffix)
        wb.save(chemin)
        etapes_excel(chemin, sheet_visuel=sheet_visuel, sheet_data=sheet_data, table_name=table_name,
                     saisie_filtre=region, structure_creee=structure_creee, maintenant=maintenant)
        logger.info("Sauvegarde : %s", chemin)

        envoyer_mail_suivi(chemin, aujourdhui, mail_to=mail_to, mail_cc=mail_cc,
                           mode_mail=mode_mail, envoyer_mail=envoyer_mail)
        logger.info("Pipeline termine : %s", chemin)
        return chemin
    except Exception as exc:  # noqa: BLE001 - ErrHandler : erreur journalisée
        logger.error("ERREUR %s - %s", type(exc).__name__, exc)
        return None


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Pipeline suivi distribution gaz (macro LancerPipelineComplet)")
    parser.add_argument("classeur", help="classeur à traiter (remplace ThisWorkbook)")
    parser.add_argument("--dossier-sortie", default=DOSSIER_SORTIE)
    parser.add_argument("--sheet-visuel", default=SHEET_VISUEL)
    parser.add_argument("--sheet-data", default=SHEET_DATA)
    parser.add_argument("--sheet-log", default=SHEET_LOG)
    parser.add_argument("--table-name", default=TABLE_NAME)
    parser.add_argument("--file-title", default=FILE_TITLE)
    parser.add_argument("--file-suffix", default=FILE_SUFFIX)
    parser.add_argument("--mail-to", default=MAIL_TO)
    parser.add_argument("--mail-cc", default=MAIL_CC)
    parser.add_argument("--mode-mail", choices=["envoyer", "afficher", "aucun"], default="envoyer")
    parser.add_argument("--saisie", default="", help="région à filtrer (vide = aucun filtre)")
    args = parser.parse_args(argv)

    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    chemin = lancer_pipeline_complet(
        args.classeur, args.dossier_sortie, sheet_visuel=args.sheet_visuel, sheet_data=args.sheet_data,
        sheet_log=args.sheet_log, table_name=args.table_name, file_title=args.file_title,
        file_suffix=args.file_suffix, mail_to=args.mail_to, mail_cc=args.mail_cc,
        mode_mail=args.mode_mail, saisie=args.saisie)
    return 0 if chemin else 1


if __name__ == "__main__":
    sys.exit(main())
