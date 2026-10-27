import random
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

import marco_25_10 as m
from connecteurs import outlook

MAINTENANT = datetime(2026, 10, 25, 11, 15, 30)
JOUR = MAINTENANT.date()


@pytest.fixture
def wb(classeur_complet):
    return load_workbook(classeur_complet)


class FauxMail:
    def __init__(self, erreur=None):
        self.appels = []
        self.erreur = erreur

    def __call__(self, destinataire, cc, objet, corps, pieces_jointes, mode):
        if self.erreur:
            raise self.erreur
        self.appels.append(
            dict(destinataire=destinataire, cc=cc, objet=objet, corps=corps, pj=list(pieces_jointes), mode=mode)
        )


# --- outils ---------------------------------------------------------------
def test_regions_et_dossier():
    assert m.regions() == ["Nord", "Sud", "Est", "Ouest", "Centre"]
    assert m.save_folder().endswith("sortiesVBA\\")


def test_get_or_create_sheet():
    wb = Workbook()
    ws = m.get_or_create_sheet(wb, "X")
    assert wb.sheetnames[-1] == "X"
    assert m.get_or_create_sheet(wb, "X") is ws


def test_log_msg_cree_entete_et_ajoute():
    wb = Workbook()
    m.log_msg(wb, "INFO", "a", MAINTENANT)
    m.log_msg(wb, "ERREUR", "b", MAINTENANT)
    ws = wb["Log"]
    assert [c.value for c in ws[1]] == ["Horodatage", "Niveau", "Message"]
    assert ws["A2"].value == MAINTENANT and ws["B2"].value == "INFO" and ws["C2"].value == "a"
    assert ws["B3"].value == "ERREUR"
    assert ws["A2"].number_format == "dd/mm/yyyy hh:mm:ss"


# --- simulation -----------------------------------------------------------
@pytest.mark.parametrize("reg,base", [("Nord", 120), ("Sud", 95), ("Est", 80), ("Ouest", 70), ("Centre", 60), ("Autre", 60)])
def test_simuler_volume_dans_la_fourchette(reg, base):
    rng = random.Random(1)
    for _ in range(50):
        v = m.simuler_volume(reg, rng)
        # base * (0,85 .. 1,15), arrondi à 0,1
        assert base * 0.85 - 0.05 <= v <= base * 1.15 + 0.05
        assert v == round(v, 1)


def test_simuler_temperature_dans_la_fourchette():
    rng = random.Random(2)
    assert all(2 <= m.simuler_temperature(rng) <= 16 for _ in range(50))


# --- structure ------------------------------------------------------------
def test_ensure_structure_cree_14_jours(classeur_vierge):
    wb = load_workbook(classeur_vierge)
    assert m.ensure_structure(wb, JOUR, random.Random(3)) is True
    ws = wb["Feuille2"]
    # 14 jours x 5 régions + en-tête = 71 lignes
    assert ws.max_row == 71
    assert [c.value for c in ws[1]] == m.ENTETES
    assert ws["A2"].value == date(2026, 10, 11)  # Date - 14
    assert ws["A71"].value == date(2026, 10, 24)  # Date - 1
    assert [ws.cell(row=r, column=2).value for r in range(2, 7)] == m.regions()
    assert ws["E2"].value == "Historique"
    assert ws.tables["tblDonnees"].ref == "A1:E71"
    assert "Log" in wb.sheetnames


def test_ensure_structure_ne_touche_pas_une_table_existante(wb):
    avant = m.lire_table(wb)
    assert m.ensure_structure(wb, JOUR, random.Random(3)) is False
    assert m.lire_table(wb).equals(avant)


# --- ajout du jour --------------------------------------------------------
def test_ajouter_donnees_du_jour_nominal(wb):
    nb = m.ajouter_donnees_du_jour(wb, MAINTENANT, random.Random(4))
    assert nb == 5
    ws = wb["Feuille2"]
    assert ws.tables["tblDonnees"].ref == "A1:E21"  # 15 lignes + 5
    assert [ws.cell(row=r, column=2).value for r in range(17, 22)] == m.regions()
    assert ws["E17"].value == "Phase de test"
    assert ws["A17"].value == JOUR
    assert wb["Log"]["C3"].value == "5 ligne(s) ajoutee(s) pour le 25/10/2026"


def test_ajouter_donnees_du_jour_idempotent(wb):
    m.ajouter_donnees_du_jour(wb, MAINTENANT, random.Random(4))
    assert m.ajouter_donnees_du_jour(wb, MAINTENANT, random.Random(5)) == 0
    assert len(m.lire_table(wb)) == 20


def test_ajouter_donnees_du_jour_complete_les_regions_manquantes(classeur_jour_present):
    wb = load_workbook(classeur_jour_present)
    # Nord et Sud existent déjà pour le 25/10 : seules Est, Ouest, Centre sont ajoutées
    assert m.ajouter_donnees_du_jour(wb, MAINTENANT, random.Random(6)) == 3
    df = m.lire_table(wb)
    assert len(df) == 15 + 2 + 3
    assert sorted(df[df["Statut"] == "Phase de test"]["Region"]) == sorted(m.regions())


def test_ligne_existe(wb):
    assert m.ligne_existe(wb, date(2026, 10, 22), "Nord") is True
    assert m.ligne_existe(wb, date(2026, 10, 22), "nord") is False  # `=` VBA sensible à la casse
    assert m.ligne_existe(wb, JOUR, "Nord") is False


def test_ajouter_sans_table_leve_keyerror(classeur_vierge):
    with pytest.raises(KeyError):
        m.ajouter_donnees_du_jour(load_workbook(classeur_vierge), MAINTENANT)


# --- filtre ---------------------------------------------------------------
def test_est_region_valide():
    assert m.est_region_valide("nord") and m.est_region_valide("OUEST")
    assert not m.est_region_valide("Paris") and not m.est_region_valide("")


def test_appliquer_filtre_masque_les_autres_regions(wb):
    assert m.appliquer_filtre(wb, "  sud ", MAINTENANT) == "Sud"
    ws = wb["Feuille2"]
    visibles = [r for r in range(2, 17) if not ws.row_dimensions[r].hidden]
    assert [ws.cell(row=r, column=2).value for r in visibles] == ["Sud"] * 3
    assert ws.tables["tblDonnees"].autoFilter.filterColumn[0].filters.filter == ["Sud"]


@pytest.mark.parametrize("saisie", ["", None, "   "])
def test_appliquer_filtre_vide_enleve_le_filtre(wb, saisie):
    m.appliquer_filtre(wb, "Nord", MAINTENANT)
    assert m.appliquer_filtre(wb, saisie, MAINTENANT) is None
    ws = wb["Feuille2"]
    assert not any(ws.row_dimensions[r].hidden for r in range(2, 17))
    assert not ws.tables["tblDonnees"].autoFilter.filterColumn


def test_appliquer_filtre_region_inconnue(wb):
    assert m.appliquer_filtre(wb, "Paris", MAINTENANT) is None
    assert not any(wb["Feuille2"].row_dimensions[r].hidden for r in range(2, 17))


# --- visuel ---------------------------------------------------------------
def test_calculer_kpi(wb):
    kpi = m.calculer_kpi(m.lire_table(wb))
    # jour k : 400 + 50k -> 400 + 450 + 500 = 1350 ; 3 jours -> 450 ; pic : Nord le dernier jour (120)
    assert kpi == {"volume_total": 1350.0, "moyenne_par_jour": 450.0, "region_pic": "Nord"}


def test_calculer_kpi_table_vide():
    import pandas as pd

    kpi = m.calculer_kpi(pd.DataFrame(columns=["Date", "Region", "Volume_GWh"]))
    assert kpi["volume_total"] == 0.0 and kpi["region_pic"] is None


def test_calculer_kpi_une_seule_ligne():
    import pandas as pd

    df = pd.DataFrame({"Date": [JOUR], "Region": ["Est"], "Volume_GWh": [80.0]})
    assert m.calculer_kpi(df) == {"volume_total": 80.0, "moyenne_par_jour": 80.0, "region_pic": "Est"}


def test_construire_tcd(wb):
    tcd = m.construire_tcd(m.lire_table(wb))
    assert tcd.loc["Total général", "Total général"] == 1350.0
    assert tcd.loc[datetime(2026, 10, 23), "Nord"] == 110.0  # 100 + 10
    assert tcd.loc[datetime(2026, 10, 24), "Total général"] == 500.0


def test_creer_visuel(wb, tmp_path):
    png = m.creer_visuel(wb, tmp_path / "graph", MAINTENANT)
    assert png.exists() and png.stat().st_size > 0
    ws = wb["Feuille1"]
    assert ws["A1"].value == "Suivi distribution gaz - mis a jour le 25/10/2026 11:15"
    assert ws["A1"].font.bold and ws["A1"].font.size == 14
    assert ws["B2"].value == "=SUM(tblDonnees[Volume_GWh])"
    assert ws["H2"].value.startswith("=INDEX(tblDonnees[Region]")
    assert "SUMPRODUCT" in ws["E2"].value
    assert ws["A6"].value == "Étiquettes de lignes"
    assert [ws.cell(row=6, column=c).value for c in range(2, 8)] == [
        "Centre", "Est", "Nord", "Ouest", "Sud", "Total général",
    ]
    assert ws["G9"].value == 500.0  # total du 24/10
    assert ws["A10"].value == "Total général" and ws["G10"].value == 1350.0
    assert len(ws._images) == 1


def test_creer_visuel_regenere_sans_doublon(wb, tmp_path):
    m.creer_visuel(wb, tmp_path, MAINTENANT)
    m.creer_visuel(wb, tmp_path, MAINTENANT)
    assert len(wb["Feuille1"]._images) == 1


# --- sauvegarde -----------------------------------------------------------
def test_sauvegarder_version_cree_dossier_et_nom(wb, tmp_path):
    chemin = m.sauvegarder_version(wb, tmp_path / "a" / "b", MAINTENANT)
    assert chemin.name == "VBA_Natran_Test_20261025_111530_VBATF.xlsx"  # .xlsx : classeur chargé sans projet VBA
    assert chemin.exists()
    load_workbook(chemin)


# --- mail -----------------------------------------------------------------
def test_construire_mail():
    objet, corps = m.construire_mail("C:\\dossier\\fichier.xlsm", MAINTENANT)
    assert objet == "[PHASE DE TEST] Mise a jour suivi distribution - 25/10/2026"
    assert corps == (
        "Bonjour,\r\n\r\nCeci est un envoi automatique en phase de test.\r\n"
        "Le graphique et les donnees ont ete mis a jour.\r\nFichier : fichier.xlsm\r\n\r\n"
        "Cordialement,\r\nRobot VBATF"
    )


@pytest.mark.parametrize("mode", ["envoyer", "afficher"])
def test_envoyer_mail_pj(wb, mode):
    faux = FauxMail()
    assert m.envoyer_mail_pj(wb, "/x/f.xlsx", "a@b.fr", "c@d.fr", mode, faux, MAINTENANT) is True
    appel = faux.appels[0]
    assert (appel["destinataire"], appel["cc"], appel["mode"], appel["pj"]) == ("a@b.fr", "c@d.fr", mode, ["/x/f.xlsx"])
    assert "Fichier : f.xlsx" in appel["corps"]
    log = wb["Log"]
    assert log.cell(row=log.max_row, column=3).value.endswith("a@b.fr")  # la fixture a déjà 1 ligne de log


def test_envoyer_mail_valeurs_par_defaut_du_vba(wb):
    faux = FauxMail()
    m.envoyer_mail_pj(wb, "f.xlsx", envoyer_mail=faux, maintenant=MAINTENANT)
    assert faux.appels[0]["destinataire"] == "jdavoust201@headmind.com"
    assert faux.appels[0]["cc"] == "" and faux.appels[0]["mode"] == "envoyer"


def test_envoyer_mail_mode_aucun(wb):
    faux = FauxMail()
    assert m.envoyer_mail_pj(wb, "f.xlsx", mode_mail="aucun", envoyer_mail=faux) is False
    assert faux.appels == []


def test_envoyer_mail_mode_invalide(wb):
    with pytest.raises(ValueError):
        m.envoyer_mail_pj(wb, "f.xlsx", mode_mail="zzz", envoyer_mail=FauxMail())


def test_envoyer_mail_erreur_journalisee(wb):
    assert m.envoyer_mail_pj(wb, "f.xlsx", envoyer_mail=FauxMail(RuntimeError("boum")), maintenant=MAINTENANT) is False
    ws = wb["Log"]
    assert ws.cell(row=ws.max_row, column=2).value == "ERREUR" and "boum" in ws.cell(row=ws.max_row, column=3).value


def test_connecteur_par_defaut_est_le_connecteur_outlook(wb):
    # la fixture autouse remplace outlook.envoyer_mail : l'appel par défaut doit donc être intercepté (erreur journalisée)
    assert m.envoyer_mail_pj(wb, "f.xlsx", maintenant=MAINTENANT) is False
    assert "Appel réel" in wb["Log"].cell(row=wb["Log"].max_row, column=3).value


def test_connecteur_outlook_mode_invalide(monkeypatch):
    monkeypatch.undo()  # retire le blocage autouse pour tester la validation, avant tout import COM
    with pytest.raises(ValueError):
        outlook.envoyer_mail("a", "", "o", "c", [], "zzz")


# --- pipeline -------------------------------------------------------------
def test_pipeline_complet(classeur_complet, tmp_path):
    faux = FauxMail()
    chemin = m.lancer_pipeline_complet(
        classeur_complet, tmp_path / "out", "Nord", "x@y.fr", "", "afficher",
        envoyer_mail=faux, maintenant=MAINTENANT, rng=random.Random(7),
    )
    assert chemin is not None and chemin.exists()
    assert (tmp_path / "out" / "graphique_volumes.png").stat().st_size > 0
    assert faux.appels[0]["mode"] == "afficher" and faux.appels[0]["pj"] == [str(chemin)]
    res = load_workbook(chemin)
    assert res["Feuille2"].tables["tblDonnees"].ref == "A1:E21"
    niveaux = [r[1].value for r in res["Log"].iter_rows(min_row=2)]
    assert "ERREUR" not in niveaux
    # comme en VBA, la copie est enregistrée avant les logs « Sauvegarde » et « Pipeline termine »
    assert res["Log"].cell(row=res["Log"].max_row, column=3).value == "Visuel regenere"


def test_pipeline_ne_modifie_pas_le_classeur_source(classeur_complet, tmp_path):
    avant = classeur_complet.read_bytes()
    m.lancer_pipeline_complet(classeur_complet, tmp_path, "", mode_mail="aucun", maintenant=MAINTENANT)
    assert classeur_complet.read_bytes() == avant


def test_pipeline_sur_classeur_vierge(classeur_vierge, tmp_path):
    chemin = m.lancer_pipeline_complet(classeur_vierge, tmp_path, "", mode_mail="aucun", maintenant=MAINTENANT)
    # 14 jours d'historique + aujourd'hui = 15 jours x 5 régions
    assert load_workbook(chemin)["Feuille2"].tables["tblDonnees"].ref == "A1:E76"


def test_pipeline_erreur_est_journalisee(classeur_complet, tmp_path, monkeypatch):
    def boum(*a, **k):
        raise RuntimeError("panne")

    monkeypatch.setattr(m, "creer_visuel", boum)
    assert m.lancer_pipeline_complet(
        classeur_complet, tmp_path, "", mode_mail="aucun", maintenant=MAINTENANT, enregistrer_classeur=False
    ) is None


def test_pipeline_enregistrer_classeur(classeur_complet, tmp_path):
    copie = tmp_path / "src.xlsx"
    copie.write_bytes(classeur_complet.read_bytes())
    m.lancer_pipeline_complet(copie, tmp_path / "o", "", mode_mail="aucun", maintenant=MAINTENANT, enregistrer_classeur=True)
    assert load_workbook(copie)["Feuille2"].tables["tblDonnees"].ref == "A1:E21"


def test_pipeline_classeur_introuvable(tmp_path):
    with pytest.raises(FileNotFoundError):
        m.lancer_pipeline_complet(tmp_path / "absent.xlsx", tmp_path, mode_mail="aucun")


def test_main_cli(classeur_complet, tmp_path, capsys):
    m.main([str(classeur_complet), "--dossier-sauvegarde", str(tmp_path), "--mode-mail", "aucun", "--region", "Est"])
    assert ".xlsx" in capsys.readouterr().out
