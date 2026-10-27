"""Crée tests/fixtures/*.xlsx (données fictives, même forme que le classeur réel). Idempotent.

Usage : python tests/generer_donnees_fictives.py [dossier_sortie]
"""

import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

REGIONS = ["Nord", "Sud", "Est", "Ouest", "Centre"]
BASES = {"Nord": 100, "Sud": 90, "Est": 80, "Ouest": 70, "Centre": 60}
JOURS = [date(2026, 10, 22), date(2026, 10, 23), date(2026, 10, 24)]


def _classeur_base() -> Workbook:
    wb = Workbook()
    wb.active.title = "Dashboard"
    wb.create_sheet("Feuille2")
    wb.create_sheet("Feuille1")
    return wb


def creer_classeur_complet(chemin: Path) -> None:
    """3 jours x 5 régions. Volume = base région + 10 par jour écoulé ; température = 5.0 + indice du jour."""
    wb = _classeur_base()
    ws = wb["Feuille2"]
    ws.append(["Date", "Region", "Volume_GWh", "Temperature_C", "Statut"])
    for k, jour in enumerate(JOURS):
        for reg in REGIONS:
            ws.append([jour, reg, float(BASES[reg] + 10 * k), 5.0 + k, "Historique"])
            ws.cell(row=ws.max_row, column=1).number_format = "dd/mm/yyyy"
    table = Table(displayName="tblDonnees", ref=f"A1:E{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    wb["Feuille1"]["A1"] = "Suivi distribution gaz - mis a jour le 24/10/2026 09:00"
    log = wb.create_sheet("Log")
    log.append(["Horodatage", "Niveau", "Message"])
    log.append([datetime(2026, 10, 24, 9, 0, 0), "INFO", "Structure initiale creee"])
    wb.save(chemin)


def creer_classeur_du_jour_present(chemin: Path) -> None:
    """Comme le classeur complet, mais Nord et Sud ont déjà une ligne le 25/10/2026."""
    creer_classeur_complet(chemin)
    from openpyxl import load_workbook

    wb = load_workbook(chemin)
    ws = wb["Feuille2"]
    ws.append([date(2026, 10, 25), "Nord", 111.0, 6.0, "Phase de test"])
    ws.append([date(2026, 10, 25), "Sud", 99.0, 6.0, "Phase de test"])
    ws.tables["tblDonnees"].ref = f"A1:E{ws.max_row}"
    wb.save(chemin)


def creer_classeur_vierge(chemin: Path) -> None:
    """Classeur sans table ni données (déclenche la création de la structure initiale)."""
    wb = Workbook()
    wb.active.title = "Feuille1"
    wb.save(chemin)


def generer(dossier: Path) -> None:
    dossier.mkdir(parents=True, exist_ok=True)
    creer_classeur_complet(dossier / "classeur_complet.xlsx")
    creer_classeur_du_jour_present(dossier / "classeur_jour_present.xlsx")
    creer_classeur_vierge(dossier / "classeur_vierge.xlsx")


if __name__ == "__main__":
    generer(Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "fixtures")
