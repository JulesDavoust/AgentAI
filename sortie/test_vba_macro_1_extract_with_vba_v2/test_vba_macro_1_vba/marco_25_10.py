"""Traduction du module VBA marco_25_10.bas (pipeline : données -> filtre -> visuel -> sauvegarde -> mail)."""

import argparse
import logging
import random
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable, Optional

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.drawing.image import Image as XlImage  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.utils.cell import range_boundaries  # noqa: E402
from openpyxl.worksheet.filters import AutoFilter, FilterColumn, Filters  # noqa: E402
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: E402

from connecteurs import outlook  # noqa: E402

logger = logging.getLogger(__name__)

# --- CONFIGURATION (constantes VBA) ---------------------------------------
SHEET_VISUEL = "Feuille1"
SHEET_DATA = "Feuille2"
SHEET_LOG = "Log"
TABLE_NAME = "tblDonnees"
FILE_TITLE = "VBA_Natran_Test"
FILE_SUFFIX = "VBATF"
MAIL_TO = "jdavoust201@headmind.com"
MAIL_CC = ""
SEND_MAIL = True
SAVE_FOLDER = "C:\\Users\\Jules\\Documents\\Jules\\Projects\\AgentAI\\examples\\sortiesVBA\\"
ENTETES = ["Date", "Region", "Volume_GWh", "Temperature_C", "Statut"]
FORMAT_DATE = "dd/mm/yyyy"
FORMAT_HORODATAGE = "dd/mm/yyyy hh:mm:ss"


def regions() -> list:
    """VBA : Function Regions()."""
    return ["Nord", "Sud", "Est", "Ouest", "Centre"]


def save_folder() -> str:
    """VBA : Function SaveFolder() (valeur par défaut, surchargée par le paramètre `dossier`)."""
    return SAVE_FOLDER


# --- OUTILS ----------------------------------------------------------------
def get_or_create_sheet(wb: Workbook, nom: str):
    """VBA : GetOrCreateSheet. Retourne la feuille, créée en dernière position si absente."""
    if nom in wb.sheetnames:
        return wb[nom]
    return wb.create_sheet(nom)


def _derniere_ligne_col_a(ws) -> int:
    """VBA : Cells(Rows.Count, 1).End(xlUp).Row (1 si la colonne est vide)."""
    for r in range(ws.max_row, 0, -1):
        if ws.cell(row=r, column=1).value not in (None, ""):
            return r
    return 1


def log_msg(wb: Workbook, niveau: str, msg: str, maintenant: Optional[datetime] = None) -> None:
    """VBA : LogMsg. Ajoute une ligne (horodatage, niveau, message) dans la feuille Log."""
    maintenant = maintenant or datetime.now()
    ws = get_or_create_sheet(wb, SHEET_LOG)
    if ws["A1"].value in (None, ""):
        for col, titre in enumerate(["Horodatage", "Niveau", "Message"], start=1):
            ws.cell(row=1, column=col, value=titre)
    n = _derniere_ligne_col_a(ws) + 1
    cell = ws.cell(row=n, column=1, value=maintenant)
    cell.number_format = FORMAT_HORODATAGE
    ws.cell(row=n, column=2, value=niveau)
    ws.cell(row=n, column=3, value=msg)


def _autofit(ws, colonnes: str) -> None:
    """VBA : Columns("A:H").AutoFit (approximation : largeur = longueur max du texte affiché)."""
    debut, fin = colonnes.split(":")
    c1, c2 = range_boundaries(f"{debut}1:{fin}1")[0], range_boundaries(f"{debut}1:{fin}1")[2]
    for c in range(c1, c2 + 1):
        largeur = 0
        for row in ws.iter_rows(min_col=c, max_col=c):
            v = row[0].value
            if v is None or (isinstance(v, str) and v.startswith("=")):
                continue
            if isinstance(v, (datetime, date)):
                texte = "dd/mm/yyyy hh:mm"
            else:
                texte = str(v)
            largeur = max(largeur, len(texte))
        ws.column_dimensions[get_column_letter(c)].width = max(8, largeur + 2)


def _table(wb: Workbook):
    """VBA : ThisWorkbook.Worksheets(SHEET_DATA).ListObjects(TABLE_NAME). KeyError si absente."""
    return wb[SHEET_DATA].tables[TABLE_NAME]


def _maj_ref_table(table: Table, ref: str) -> None:
    table.ref = ref
    if table.autoFilter is not None:
        table.autoFilter.ref = ref


def lire_table(wb: Workbook) -> pd.DataFrame:
    """Lit les lignes de la table tblDonnees dans un DataFrame (colonnes = en-têtes de la table)."""
    ws = wb[SHEET_DATA]
    c1, r1, c2, r2 = range_boundaries(_table(wb).ref)
    lignes = [[ws.cell(row=r, column=c).value for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]
    df = pd.DataFrame(lignes[1:], columns=lignes[0])
    df["Date"] = pd.to_datetime(df["Date"])  # openpyxl relit des datetime, mais les lignes fraîches sont des date
    return df


# --- 1. STRUCTURE / DONNEES ------------------------------------------------
def simuler_volume(reg: str, rng: random.Random) -> float:
    """VBA : SimulerVolume. Base par région x (0,85 + Rnd*0,3), arrondi à 1 décimale."""
    base = {"Nord": 120, "Sud": 95, "Est": 80, "Ouest": 70}.get(reg, 60)
    return round(base * (0.85 + rng.random() * 0.3), 1)


def simuler_temperature(rng: random.Random) -> float:
    """VBA : SimulerTemperature. 2 + Rnd*14 arrondi à 1 décimale."""
    return round(2 + rng.random() * 14, 1)


def ensure_structure(wb: Workbook, aujourdhui: Optional[date] = None, rng: Optional[random.Random] = None) -> bool:
    """VBA : EnsureStructure. Crée les feuilles et, si la table est absente, 14 jours d'historique.

    Retourne True si la structure initiale a été créée.
    """
    aujourdhui = aujourdhui or date.today()
    rng = rng or random.Random()
    wsd = get_or_create_sheet(wb, SHEET_DATA)
    get_or_create_sheet(wb, SHEET_VISUEL)

    if TABLE_NAME in wsd.tables:
        return False

    wsd.delete_rows(1, wsd.max_row)  # Cells.Clear
    for col, titre in enumerate(ENTETES, start=1):
        wsd.cell(row=1, column=col, value=titre)
    ligne = 2
    for i in range(14, 0, -1):
        for r in regions():
            c = wsd.cell(row=ligne, column=1, value=aujourdhui - timedelta(days=i))
            c.number_format = FORMAT_DATE
            wsd.cell(row=ligne, column=2, value=r)
            wsd.cell(row=ligne, column=3, value=simuler_volume(r, rng))
            wsd.cell(row=ligne, column=4, value=simuler_temperature(rng))
            wsd.cell(row=ligne, column=5, value="Historique")
            ligne += 1
    ref = f"A1:E{ligne - 1}"
    table = Table(displayName=TABLE_NAME, ref=ref)
    table.autoFilter = AutoFilter(ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    wsd.add_table(table)
    _autofit(wsd, "A:E")
    log_msg(wb, "INFO", "Structure initiale creee")
    return True


def ligne_existe(wb: Workbook, d: date, reg: str) -> bool:
    """VBA : LigneExiste. Vrai si la table contient déjà (date, région) ; région comparée en respectant la casse (VBA `=`)."""
    ws = wb[SHEET_DATA]
    c1, r1, c2, r2 = range_boundaries(_table(wb).ref)
    for r in range(r1 + 1, r2 + 1):
        v = ws.cell(row=r, column=c1).value
        if isinstance(v, datetime):
            v = v.date()
        if v == d and ws.cell(row=r, column=c1 + 1).value == reg:
            return True
    return False


def ajouter_donnees_du_jour(
    wb: Workbook, maintenant: Optional[datetime] = None, rng: Optional[random.Random] = None
) -> int:
    """VBA : AjouterDonneesDuJour. Ajoute une ligne par région absente pour aujourd'hui ; retourne le nombre ajouté."""
    maintenant = maintenant or datetime.now()
    rng = rng or random.Random()
    jour = maintenant.date()
    ws = wb[SHEET_DATA]
    table = _table(wb)
    nb = 0
    for r in regions():
        if not ligne_existe(wb, jour, r):
            c1, r1, c2, r2 = range_boundaries(table.ref)
            n = r2 + 1
            ws.cell(row=n, column=c1, value=jour).number_format = FORMAT_DATE
            ws.cell(row=n, column=c1 + 1, value=r)
            ws.cell(row=n, column=c1 + 2, value=simuler_volume(r, rng))
            ws.cell(row=n, column=c1 + 3, value=simuler_temperature(rng))
            ws.cell(row=n, column=c1 + 4, value="Phase de test")
            _maj_ref_table(table, f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{n}")
            nb += 1
    log_msg(wb, "INFO", f"{nb} ligne(s) ajoutee(s) pour le {jour.strftime('%d/%m/%Y')}", maintenant)
    return nb


# --- 2. FILTRE -------------------------------------------------------------
def est_region_valide(s: str) -> bool:
    """VBA : EstRegionValide (StrComp en vbTextCompare : insensible à la casse)."""
    return any(r.casefold() == s.casefold() for r in regions())


def appliquer_filtre(wb: Workbook, saisie: Optional[str], maintenant: Optional[datetime] = None) -> Optional[str]:
    """VBA : AppliquerFiltre. `saisie` remplace l'InputBox (vide/None = Annuler = aucun filtre).

    Retourne la région filtrée, ou None si aucun filtre n'a été appliqué (saisie vide ou région inconnue).
    Le filtre est posé sur la colonne Region de la table et les lignes non retenues sont masquées.
    """
    ws = wb[SHEET_DATA]
    table = _table(wb)
    c1, r1, c2, r2 = range_boundaries(table.ref)

    # If lo.AutoFilter.FilterMode Then ShowAllData
    if table.autoFilter is not None:
        table.autoFilter.filterColumn = []
    for r in range(r1 + 1, r2 + 1):
        ws.row_dimensions[r].hidden = False

    saisie = (saisie or "").strip()
    if saisie == "":
        return None
    if not est_region_valide(saisie):
        logger.warning("Region inconnue : %s", saisie)  # MsgBox vbExclamation
        return None
    saisie = next(r for r in regions() if r.casefold() == saisie.casefold())  # nom canonique (Excel filtre sans casse)

    if table.autoFilter is None:
        table.autoFilter = AutoFilter(ref=table.ref)
    table.autoFilter.filterColumn = [FilterColumn(colId=1, filters=Filters(filter=[saisie]))]
    for r in range(r1 + 1, r2 + 1):
        v = ws.cell(row=r, column=c1 + 1).value
        ws.row_dimensions[r].hidden = not (isinstance(v, str) and v.casefold() == saisie.casefold())
    log_msg(wb, "INFO", f"Filtre applique : {saisie}", maintenant)
    return saisie


# --- 3. VISUEL -------------------------------------------------------------
def calculer_kpi(df: pd.DataFrame) -> dict:
    """Valeurs que les formules KPI de CreerVisuel donnent dans Excel (calculées sur toute la table)."""
    if df.empty:
        return {"volume_total": 0.0, "moyenne_par_jour": None, "region_pic": None}
    total = float(df["Volume_GWh"].sum())
    nb_jours = df["Date"].nunique()
    pic = df.loc[df["Volume_GWh"].idxmax(), "Region"]  # première occurrence, comme MATCH(...,0)
    return {"volume_total": total, "moyenne_par_jour": total / nb_jours, "region_pic": pic}


def construire_tcd(df: pd.DataFrame) -> pd.DataFrame:
    """Équivalent du TCD TCD_Volume : lignes = Date, colonnes = Region, valeurs = somme de Volume_GWh (avec totaux)."""
    return df.pivot_table(
        index="Date", columns="Region", values="Volume_GWh", aggfunc="sum", margins=True, margins_name="Total général"
    )


def creer_visuel(wb: Workbook, dossier_graphique: Path, maintenant: Optional[datetime] = None) -> Path:
    """VBA : CreerVisuel. Régénère Feuille1 (titre, KPI, TCD statique, graphique) ; retourne le chemin du PNG."""
    maintenant = maintenant or datetime.now()
    dossier_graphique = Path(dossier_graphique)
    wsv = wb[SHEET_VISUEL]
    df = lire_table(wb)

    # Nettoyage (ChartObjects / PivotTables / Cells.Clear)
    wsv._images = []
    wsv._charts = []
    wsv.delete_rows(1, wsv.max_row)

    wsv["A1"] = "Suivi distribution gaz - mis a jour le " + maintenant.strftime("%d/%m/%Y %H:%M")
    wsv["A1"].font = Font(bold=True, size=14)
    wsv["A2"] = "Volume total (GWh)"
    wsv["B2"] = f"=SUM({TABLE_NAME}[Volume_GWh])"
    wsv["D2"] = "Moyenne / jour (GWh)"
    wsv["E2"] = (
        f"=SUM({TABLE_NAME}[Volume_GWh])/SUMPRODUCT(1/COUNTIF({TABLE_NAME}[Date],{TABLE_NAME}[Date]))"
    )
    wsv["G2"] = "Region du pic"
    for ref in ("A2", "D2", "G2"):
        wsv[ref].font = Font(bold=True)
    for ref in ("B2", "E2"):
        wsv[ref].number_format = "#,##0.0"
    wsv["H2"] = (
        f"=INDEX({TABLE_NAME}[Region],MATCH(MAX({TABLE_NAME}[Volume_GWh]),{TABLE_NAME}[Volume_GWh],0))"
    )

    # TCD (openpyxl ne sait pas créer de tableau croisé dynamique : copie statique des valeurs, à partir de A5)
    tcd = construire_tcd(df)
    wsv["A5"] = "Volume (GWh)"
    wsv["B5"] = "Étiquettes de colonnes"
    wsv["A6"] = "Étiquettes de lignes"
    for j, nom in enumerate(tcd.columns, start=2):
        wsv.cell(row=6, column=j, value=nom)
    for i, (idx, ligne) in enumerate(tcd.iterrows(), start=7):
        c = wsv.cell(row=i, column=1, value=idx.to_pydatetime() if hasattr(idx, "to_pydatetime") else idx)
        if idx != "Total général":
            c.number_format = "dd/mm"
        for j, v in enumerate(ligne, start=2):
            wsv.cell(row=i, column=j, value=None if pd.isna(v) else float(v)).number_format = "#,##0.0"

    # Graphique : lignes par région (sans les totaux), exporté en PNG puis intégré en J5
    dossier_graphique.mkdir(parents=True, exist_ok=True)
    png = dossier_graphique / "graphique_volumes.png"
    fig, ax = plt.subplots(figsize=(5.2, 3.0), dpi=100)
    donnees = tcd.drop(index="Total général", columns="Total général")
    for reg in donnees.columns:
        ax.plot(list(donnees.index), donnees[reg].values, label=reg)
    ax.set_title("Volumes distribues par region (GWh)")
    ax.legend(fontsize=6)
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(png)
    plt.close(fig)
    img = XlImage(str(png))
    wsv.add_image(img, "J5")

    _autofit(wsv, "A:H")
    log_msg(wb, "INFO", "Visuel regenere", maintenant)
    return png


# --- 4. SAUVEGARDE ---------------------------------------------------------
def sauvegarder_version(wb: Workbook, dossier, maintenant: Optional[datetime] = None) -> Path:
    """VBA : SauvegarderVersion (SaveCopyAs). Nom : titre_yyyymmdd_hhnnss_VBATF.xlsm ; crée le dossier si besoin.

    Si le classeur n'a pas été chargé avec son projet VBA, l'extension est .xlsx (un .xlsm sans macros serait illisible).
    """
    maintenant = maintenant or datetime.now()
    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    ext = ".xlsm" if getattr(wb, "vba_archive", None) is not None else ".xlsx"
    nom = f"{FILE_TITLE}_{maintenant.strftime('%Y%m%d_%H%M%S')}_{FILE_SUFFIX}{ext}"
    chemin = dossier / nom
    wb.save(chemin)
    log_msg(wb, "INFO", f"Sauvegarde : {chemin}", maintenant)
    return chemin


# --- 5. MAIL ---------------------------------------------------------------
def construire_mail(chemin_pj, maintenant: Optional[datetime] = None) -> tuple:
    """Objet et corps du mail de EnvoyerMail (retours à la ligne vbCrLf = \\r\\n)."""
    maintenant = maintenant or datetime.now()
    nom_fichier = str(chemin_pj).replace("\\", "/").rsplit("/", 1)[-1]  # Mid$(..., InStrRev(..., "\") + 1)
    objet = "[PHASE DE TEST] Mise a jour suivi distribution - " + maintenant.strftime("%d/%m/%Y")
    corps = (
        "Bonjour,\r\n\r\n"
        "Ceci est un envoi automatique en phase de test.\r\n"
        "Le graphique et les donnees ont ete mis a jour.\r\n"
        f"Fichier : {nom_fichier}\r\n\r\n"
        "Cordialement,\r\nRobot VBATF"
    )
    return objet, corps


def envoyer_mail_pj(
    wb: Workbook,
    chemin_pj,
    destinataire: str = MAIL_TO,
    cc: str = MAIL_CC,
    mode_mail: str = "envoyer",
    envoyer_mail: Optional[Callable] = None,
    maintenant: Optional[datetime] = None,
) -> bool:
    """VBA : EnvoyerMail. mode_mail : envoyer (SEND_MAIL=True) | afficher (SEND_MAIL=False) | aucun.

    Comme en VBA, une erreur d'envoi est journalisée (feuille Log) et n'interrompt pas le pipeline.
    Retourne True si le mail a été envoyé/affiché.
    """
    if mode_mail == "aucun":
        return False
    if mode_mail not in ("envoyer", "afficher"):
        raise ValueError(f"Mode de mail inconnu : {mode_mail!r}")
    if envoyer_mail is None:
        envoyer_mail = outlook.envoyer_mail  # résolu à l'appel (permet de le remplacer dans les tests)
    objet, corps = construire_mail(chemin_pj, maintenant)
    try:
        envoyer_mail(destinataire, cc, objet, corps, [str(chemin_pj)], mode_mail)
    except Exception as exc:  # VBA : On Error GoTo MailErr
        log_msg(wb, "ERREUR", f"Mail impossible : {exc}", maintenant)
        logger.error("Mail impossible : %s", exc)
        return False
    log_msg(wb, "INFO", ("Mail envoye a " if mode_mail == "envoyer" else "Brouillon cree pour ") + destinataire, maintenant)
    return True


# --- POINT D'ENTREE --------------------------------------------------------
def lancer_pipeline_complet(
    chemin_classeur,
    dossier_sauvegarde=SAVE_FOLDER,
    region_filtre: Optional[str] = None,
    destinataire: str = MAIL_TO,
    cc: str = MAIL_CC,
    mode_mail: str = "envoyer",
    dossier_graphique=None,
    envoyer_mail: Optional[Callable] = None,
    maintenant: Optional[datetime] = None,
    rng: Optional[random.Random] = None,
    enregistrer_classeur: bool = False,
) -> Optional[Path]:
    """VBA : LancerPipelineComplet.

    Ouvre le classeur, enchaîne structure -> données du jour -> filtre -> visuel -> sauvegarde -> mail.
    `region_filtre` remplace l'InputBox. Retourne le chemin de la copie sauvegardée, ou None en cas d'erreur
    (l'erreur est journalisée dans la feuille Log, comme le gestionnaire ErrHandler du VBA).
    `enregistrer_classeur` : le VBA ne sauvegarde que la copie (SaveCopyAs) ; mettre True pour aussi écrire le classeur source.
    """
    chemin_classeur = Path(chemin_classeur)
    maintenant = maintenant or datetime.now()
    wb = load_workbook(chemin_classeur, keep_vba=chemin_classeur.suffix.lower() == ".xlsm")
    if dossier_graphique is None:
        dossier_graphique = Path(dossier_sauvegarde)
    try:
        ensure_structure(wb, maintenant.date(), rng)
        ajouter_donnees_du_jour(wb, maintenant, rng)
        appliquer_filtre(wb, region_filtre, maintenant)
        creer_visuel(wb, dossier_graphique, maintenant)
        chemin = sauvegarder_version(wb, dossier_sauvegarde, maintenant)
        envoyer_mail_pj(wb, chemin, destinataire, cc, mode_mail, envoyer_mail, maintenant)
        log_msg(wb, "INFO", f"Pipeline termine : {chemin}", maintenant)
        logger.info("Pipeline termine.\n%s", chemin)  # MsgBox vbInformation
    except Exception as exc:  # VBA : On Error GoTo ErrHandler
        log_msg(wb, "ERREUR", f"{type(exc).__name__} - {exc}", maintenant)
        logger.error("Erreur : %s", exc)
        chemin = None
    if enregistrer_classeur:
        wb.save(chemin_classeur)
    return chemin


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description="Pipeline de suivi de distribution gaz (traduction de marco_25_10.bas)")
    parser.add_argument("classeur", help="Classeur .xlsm/.xlsx à traiter")
    parser.add_argument("--dossier-sauvegarde", default=SAVE_FOLDER)
    parser.add_argument("--dossier-graphique", default=None, help="Dossier du PNG (défaut : dossier de sauvegarde)")
    parser.add_argument("--region", default="", help="Région à filtrer (vide = aucun filtre)")
    parser.add_argument("--destinataire", default=MAIL_TO)
    parser.add_argument("--cc", default=MAIL_CC)
    parser.add_argument("--mode-mail", choices=["envoyer", "afficher", "aucun"], default="envoyer")
    parser.add_argument("--enregistrer-classeur", action="store_true", help="Réécrit aussi le classeur source")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    chemin = lancer_pipeline_complet(
        args.classeur,
        args.dossier_sauvegarde,
        args.region,
        args.destinataire,
        args.cc,
        args.mode_mail,
        args.dossier_graphique,
        enregistrer_classeur=args.enregistrer_classeur,
    )
    print(chemin if chemin else "Échec : voir la feuille Log")


if __name__ == "__main__":
    main()
