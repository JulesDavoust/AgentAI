import random
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

import marco_25_10 as m
from bibliotheque.excel_com import EnregistreurCOM

MAINTENANT = datetime(2026, 10, 25, 9, 30, 15)


def _journal():
    lignes = []
    return m.Journal(lambda niveau, msg: lignes.append((niveau, msg))), lignes


def _lignes_table(wb, nom="tblDonnees"):
    ws = wb["Feuille2"]
    return list(ws.iter_rows(min_row=2, values_only=True)), ws.tables[nom]


# ---- fonctions pures


def test_regions():
    assert m.regions() == ("Nord", "Sud", "Est", "Ouest", "Centre")


@pytest.mark.parametrize("reg,base", [("Nord", 120), ("Sud", 95), ("Est", 80), ("Ouest", 70),
                                      ("Centre", 60), ("Inconnue", 60)])
def test_simuler_volume_bornes(reg, base):
    # base * (0.85 .. 1.15), arrondi à 1 décimale
    for graine in range(20):
        v = m.simuler_volume(reg, random.Random(graine))
        assert base * 0.85 - 0.05 <= v <= base * 1.15 + 0.05
        assert v == round(v, 1)


def test_simuler_volume_valeur_exacte():
    class Fixe:
        def random(self):
            return 0.5
    # 120 * (0.85 + 0.5*0.3) = 120 * 1.0 = 120.0
    assert m.simuler_volume("Nord", Fixe()) == 120.0
    # 2 + 0.5*14 = 9.0
    assert m.simuler_temperature(Fixe()) == 9.0


def test_est_region_valide_insensible_casse():
    assert m.est_region_valide("nord") and m.est_region_valide("OUEST")
    assert not m.est_region_valide("Paris") and not m.est_region_valide("")


def test_preparer_saisie_filtre():
    assert m.preparer_saisie_filtre("  sud ") == "sud"
    assert m.preparer_saisie_filtre("") == ""
    assert m.preparer_saisie_filtre("   ") == ""
    assert m.preparer_saisie_filtre("Paris") == ""  # région inconnue : pas de filtre


def test_nom_fichier_version_et_chemin(tmp_path, monkeypatch):
    assert m.nom_fichier_version("T", "S", ".xlsx", MAINTENANT) == "T_20261025_093015_S.xlsx"
    monkeypatch.chdir(tmp_path)
    chemin = m.chemin_sauvegarde("rel/sous", "a.xlsx")
    assert Path(chemin).is_absolute() and Path(chemin).parent.is_dir()


# ---- structure et données


def test_ensure_structure_cree_14_jours(classeur_vide):
    wb = openpyxl.load_workbook(classeur_vide)
    journal, lignes = _journal()
    assert m.ensure_structure(wb, "Feuille2", "Feuille1", "tblDonnees", journal,
                              date(2026, 10, 25), random.Random(1)) is True
    assert "Feuille1" in wb.sheetnames
    donnees, tbl = _lignes_table(wb)
    assert len(donnees) == 70  # 14 jours x 5 régions
    assert tbl.ref == "A1:E71"
    assert m._en_jour(donnees[0][0]) == date(2026, 10, 11)  # Date - 14
    assert m._en_jour(donnees[-1][0]) == date(2026, 10, 24)  # Date - 1
    assert {d[4] for d in donnees} == {"Historique"}
    assert [c.value for c in wb["Feuille2"][1]] == list(m.EN_TETES)
    assert lignes == [("INFO", "Structure initiale creee")]


def test_ensure_structure_ne_touche_pas_a_une_table_existante(classeur_complet):
    wb = openpyxl.load_workbook(classeur_complet)
    journal, lignes = _journal()
    assert m.ensure_structure(wb, "Feuille2", "Feuille1", "tblDonnees", journal,
                              date(2026, 10, 25), random.Random(1)) is False
    assert len(_lignes_table(wb)[0]) == 10 and lignes == []


def test_ajouter_donnees_du_jour(classeur_complet):
    wb = openpyxl.load_workbook(classeur_complet)
    journal, lignes = _journal()
    nb = m.ajouter_donnees_du_jour(wb, "Feuille2", "tblDonnees", journal, date(2026, 10, 25),
                                   random.Random(2))
    donnees, tbl = _lignes_table(wb)
    assert nb == 5 and len(donnees) == 15 and tbl.ref == "A1:E16"
    assert tbl.autoFilter.ref == "A1:E16"
    assert [d[1] for d in donnees[-5:]] == list(m.regions())
    assert {d[4] for d in donnees[-5:]} == {"Phase de test"}
    assert lignes == [("INFO", "5 ligne(s) ajoutee(s) pour le 25/10/2026")]


def test_ajouter_donnees_du_jour_idempotent_et_partiel(classeur_complet):
    wb = openpyxl.load_workbook(classeur_complet)
    journal, _ = _journal()
    # 24/10 existe déjà pour les 5 régions : rien à ajouter
    assert m.ajouter_donnees_du_jour(wb, "Feuille2", "tblDonnees", journal, date(2026, 10, 24),
                                     random.Random(2)) == 0
    assert m.ajouter_donnees_du_jour(wb, "Feuille2", "tblDonnees", journal, date(2026, 10, 25),
                                     random.Random(2)) == 5
    assert m.ajouter_donnees_du_jour(wb, "Feuille2", "tblDonnees", journal, date(2026, 10, 25),
                                     random.Random(2)) == 0


def test_ligne_existe_region_sensible_a_la_casse(classeur_complet):
    wb = openpyxl.load_workbook(classeur_complet)
    ws = wb["Feuille2"]
    tbl = ws.tables["tblDonnees"]
    assert m.ligne_existe(ws, tbl, date(2026, 10, 24), "Nord")
    assert not m.ligne_existe(ws, tbl, date(2026, 10, 24), "nord")  # VBA : "=" binaire


def test_ecrire_log_feuille():
    ws = openpyxl.Workbook().active
    m.ecrire_log_feuille(ws, "INFO", "a", MAINTENANT)
    m.ecrire_log_feuille(ws, "ERREUR", "b", MAINTENANT)
    assert [c.value for c in ws[1]] == ["Horodatage", "Niveau", "Message"]
    assert [c.value for c in ws[3]] == [MAINTENANT, "ERREUR", "b"]


# ---- mail


def test_envoyer_mail_envoyer(tmp_path):
    pj = tmp_path / "f.xlsx"
    pj.write_text("x")
    recu = []

    def faux(to, cc, sujet, corps, pjs, mode):
        assert all(Path(p).is_absolute() and Path(p).exists() for p in pjs)
        recu.append((to, cc, sujet, corps, pjs, mode))

    journal, lignes = _journal()
    assert m.envoyer_mail(str(pj), "a@b.c", "", "envoyer", journal, date(2026, 10, 25), faux)
    to, cc, sujet, corps, pjs, mode = recu[0]
    assert (to, cc, mode) == ("a@b.c", "", "envoyer")
    assert sujet == "[PHASE DE TEST] Mise a jour suivi distribution - 25/10/2026"
    assert "Fichier : f.xlsx\r\n" in corps and corps.endswith("Robot VBATF")
    assert pjs == [str(pj.resolve())]
    assert lignes == [("INFO", "Mail envoye a a@b.c")]


def test_envoyer_mail_afficher_aucun_et_erreur(tmp_path):
    pj = tmp_path / "f.xlsx"
    pj.write_text("x")
    appels = []
    journal, lignes = _journal()
    assert m.envoyer_mail(str(pj), "a@b.c", "", "afficher", journal, date(2026, 10, 25),
                          lambda *a: appels.append(a[-1]))
    assert appels == ["afficher"] and lignes == [("INFO", "Brouillon cree pour a@b.c")]
    assert not m.envoyer_mail(str(pj), "a@b.c", "", "aucun", journal, date(2026, 10, 25),
                              lambda *a: appels.append("non"))
    assert appels == ["afficher"]

    def panne(*a):
        raise RuntimeError("Outlook absent")

    assert not m.envoyer_mail(str(pj), "a@b.c", "", "envoyer", journal, date(2026, 10, 25), panne)
    assert lignes[-1] == ("ERREUR", "Mail impossible : Outlook absent")


def test_envoyer_mail_utilise_le_connecteur_par_defaut_bloque(tmp_path):
    journal, lignes = _journal()
    # conftest remplace le connecteur par une fonction qui lève : l'erreur est journalisée
    assert not m.envoyer_mail(str(tmp_path / "f.xlsx"), "a@b.c", "", "envoyer", journal,
                              date(2026, 10, 25))
    assert lignes[-1][0] == "ERREUR"


# ---- pipeline


class FauxExcel:
    """Fonction d'entrée simulée : exécute les étapes sur un EnregistreurCOM."""

    def __init__(self):
        self.chemin = None
        self.wb = EnregistreurCOM("wb", reponses={
            "wb.Worksheets('Log').Rows.Count": 1048576,
            "wb.Worksheets('Log').Cells(1048576, 1).End(-4162).Row": 3,
        })
        self.fichier_existait = None

    def __call__(self, chemin, fonction):
        self.chemin = chemin
        self.fichier_existait = Path(chemin).exists()  # enregistré avant l'ouverture dans Excel
        return fonction(self.wb)


def _lancer(classeur, tmp_path, **kw):
    faux = FauxExcel()
    mails = []

    def faux_mail(to, cc, sujet, corps, pjs, mode):
        assert all(Path(p).is_absolute() for p in pjs)
        mails.append((to, cc, sujet, pjs, mode))

    chemin = m.lancer_pipeline_complet(
        classeur, kw.pop("dossier_sortie", tmp_path / "out"), envoyer_mail_fn=faux_mail,
        executer_excel=faux, maintenant=lambda: MAINTENANT, alea=random.Random(3), **kw)
    return chemin, faux, mails


def test_pipeline_nominal(classeur_complet, tmp_path):
    chemin, faux, mails = _lancer(classeur_complet, tmp_path, saisie="sud")
    assert chemin == str((tmp_path / "out").resolve() / "VBA_Natran_Test_20261025_093015_VBATF.xlsx")
    assert faux.fichier_existait and Path(faux.chemin) == classeur_complet
    assert mails and mails[0][0] == "jdavoust201@headmind.com" and mails[0][4] == "envoyer"
    # données ajoutées dans le classeur enregistré avant Excel
    wb = openpyxl.load_workbook(classeur_complet)
    assert len(_lignes_table(wb)[0]) == 15
    # étapes Excel dans l'ordre
    instr = faux.wb.instructions()
    assert any("AutoFilter(Field=2, Criteria1='sud')" in i for i in instr)
    ordre = [next(k for k, i in enumerate(instr) if cle in i) for cle in
             ("Criteria1", "TCD_Volume", "SaveCopyAs")]
    assert ordre == sorted(ordre)
    assert faux.wb.appels("SaveCopyAs") == [((chemin,), {})]
    messages = [v for t, c, v in faux.wb.journal if t == "affectation" and c.endswith("Cells(4, 3).Value")]
    assert messages[-1] == f"Pipeline termine : {chemin}"


def test_pipeline_sans_filtre_et_mode_aucun(classeur_complet, tmp_path):
    chemin, faux, mails = _lancer(classeur_complet, tmp_path, mode_mail="aucun", saisie="Paris")
    assert chemin and mails == []
    assert faux.wb.appels("AutoFilter") == []


def test_pipeline_cree_la_structure(classeur_vide, tmp_path):
    chemin, _, _ = _lancer(classeur_vide, tmp_path)
    wb = openpyxl.load_workbook(classeur_vide)
    assert len(_lignes_table(wb)[0]) == 70 + 5
    assert [c.value for c in wb["Log"][1]] == ["Horodatage", "Niveau", "Message"]
    assert wb["Log"]["C2"].value == "Structure initiale creee"


def test_pipeline_erreur_excel_journalisee(classeur_complet, tmp_path):
    class Panne:
        def __init__(self):
            self.log = FauxExcel().wb

        def Worksheets(self, nom):
            if nom == "Feuille2":
                raise RuntimeError("feuille introuvable")
            return self.log.Worksheets(nom)

    def excel_en_panne(chemin, fonction):
        return fonction(Panne())

    chemin = m.lancer_pipeline_complet(classeur_complet, tmp_path / "out",
                                       executer_excel=excel_en_panne, maintenant=lambda: MAINTENANT)
    assert chemin is None  # l'erreur est journalisée, rien n'est sauvegardé ni envoyé


def test_cli_help_et_interface(capsys):
    with pytest.raises(SystemExit):
        m.main(["--help"])
    sortie = capsys.readouterr().out
    import json
    contrat = json.loads((Path(__file__).parent.parent / "interface.json").read_text(encoding="utf-8"))
    for arg in contrat["arguments"]:
        if arg["nom"].startswith("--"):
            assert arg["nom"] in sortie
