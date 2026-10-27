import random
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

import marco_25_10 as m
from bibliotheque.excel_com import excel_disponible
from connecteurs import excel as cx

pytestmark = [pytest.mark.excel_reel,
              pytest.mark.skipif(not excel_disponible(), reason="Excel indisponible")]

FIXTURE = Path(__file__).parent / "fixtures" / "classeur_complet.xlsx"


def test_pipeline_sur_excel_reel(tmp_path):
    classeur = tmp_path / "c.xlsx"
    shutil.copy(FIXTURE, classeur)
    maintenant = datetime(2026, 10, 25, 9, 30, 15)
    chemin = m.lancer_pipeline_complet(
        classeur, tmp_path / "out", mode_mail="aucun", saisie="Sud",
        maintenant=lambda: maintenant, alea=random.Random(1))
    assert chemin and Path(chemin).exists()
    for fichier in (classeur, Path(chemin)):
        wb = openpyxl.load_workbook(fichier)
        ws = wb["Feuille1"]
        assert [p.name for p in ws._pivots] == ["TCD_Volume"]
        assert len(ws._charts) == 1
        assert ws["B2"].value == "=SUM(tblDonnees[Volume_GWh])"
        assert ws["H2"].value.startswith("=INDEX(tblDonnees[Region],MATCH(MAX(")
        assert ws["A2"].font.b and ws["B2"].number_format != "General"
    log = [r[2] for r in openpyxl.load_workbook(classeur)["Log"].iter_rows(min_row=2, values_only=True)]
    assert any(x == f"Pipeline termine : {chemin}" for x in log)


def test_format_local_accepte_par_excel(tmp_path):
    from bibliotheque.excel_com import session_excel
    classeur = tmp_path / "c.xlsx"
    shutil.copy(FIXTURE, classeur)
    with session_excel(classeur) as wb:
        cellule = wb.Worksheets("Log").Cells(10, 1)
        for code in ("dd/mm/yyyy hh:mm:ss", "dd/mm", "#,##0.0"):
            cellule.NumberFormat = cx.format_excel(wb, code)  # ne doit pas lever ErreurExcel


def test_executer_dans_excel_reel(tmp_path):
    classeur = tmp_path / "c.xlsx"
    shutil.copy(FIXTURE, classeur)
    maintenant = datetime(2026, 10, 25, 9, 30)
    cx.executer_dans_excel(classeur, lambda wb: cx.ecrire_log_excel(wb, "Log", "INFO", "x", maintenant))
    ws = openpyxl.load_workbook(classeur)["Log"]
    assert ws["C3"].value == "x"
