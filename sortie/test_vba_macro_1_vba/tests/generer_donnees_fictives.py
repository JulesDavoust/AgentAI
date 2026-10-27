"""Crée les classeurs fictifs de tests/fixtures/ (idempotent, valeurs inventées)."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

FIXTURES = Path(__file__).parent / "fixtures"
REGIONS = ("Nord", "Sud", "Est", "Ouest", "Centre")
BASES = (120.0, 95.0, 80.0, 70.0, 60.0)


def creer_classeur_complet(chemin: Path) -> None:
    """Dashboard, Feuille2 (tblDonnees, 2 jours x 5 régions), Feuille1 (visuel), Log."""
    wb = openpyxl.Workbook()
    wb.active.title = "Dashboard"
    ws = wb.create_sheet("Feuille2")
    ws.append(["Date", "Region", "Volume_GWh", "Temperature_C", "Statut"])
    for jour in (date(2026, 10, 23), date(2026, 10, 24)):
        for reg, base in zip(REGIONS, BASES):
            ws.append([jour, reg, base, 8.5, "Historique"])
            ws.cell(row=ws.max_row, column=1).number_format = "dd/mm/yyyy"
    tbl = Table(displayName="tblDonnees", ref="A1:E11")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)
    wb.create_sheet("Feuille1")["A1"] = "Suivi distribution gaz - mis a jour le 24/10/2026 08:00"
    log = wb.create_sheet("Log")
    log.append(["Horodatage", "Niveau", "Message"])
    log.append([datetime(2026, 10, 24, 8, 0, 0), "INFO", "Structure initiale creee"])
    wb.save(chemin)


def creer_classeur_vide(chemin: Path) -> None:
    """Classeur sans aucune des feuilles de travail : la structure est créée par le pipeline."""
    wb = openpyxl.Workbook()
    wb.active.title = "Dashboard"
    wb.save(chemin)


def generer(dossier: Path = FIXTURES) -> None:
    dossier.mkdir(parents=True, exist_ok=True)
    creer_classeur_complet(dossier / "classeur_complet.xlsx")
    creer_classeur_vide(dossier / "classeur_vide.xlsx")


if __name__ == "__main__":
    generer()
