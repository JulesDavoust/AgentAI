"""Traduction du module VBA `marco_25_10` (suivi distribution gaz).

Pipeline : Maj données -> Filtre -> Visuel -> Sauvegarde -> Mail.
Openpyxl fait la structure et les données ; Excel (connecteurs/excel.py) applique le filtre,
crée le TCD et le graphique ; Outlook (connecteurs/outlook.py) envoie le mail.
"""

from __future__ import annotations

import argparse
import logging
import random
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable

import openpyxl
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableStyleInfo

from connecteurs import excel as connecteur_excel
from connecteurs import outlook as connecteur_outlook

logger = logging.getLogger(__name__)

# --- CONFIGURATION (constantes VBA, devenues des options de ligne de commande) ---------------
SHEET_VISUEL = "Feuille1"
SHEET_DATA = "Feuille2"
SHEET_LOG = "Log"
TABLE_NAME = "tblDonnees"
FILE_TITLE = "VBA_Natran_Test"
FILE_SUFFIX = "VBATF"
MAIL_TO = "jdavoust201@headmind.com"
MAIL_CC = ""
DOSSIER_SORTIE = "sortie"
EN_TETES = ("Date", "Region", "Volume_GWh", "Temperature_C", "Statut")


def regions() -> tuple[str, ...]:
    """VBA Regions()."""
    return ("Nord", "Sud", "Est", "Ouest", "Centre")


# --------------------------------------------------------------------------- outils


def get_or_create_sheet(wb, nom: str):
    """VBA GetOrCreateSheet : feuille existante, sinon créée après la dernière."""
    if nom in wb.sheetnames:
        return wb[nom]
    return wb.create_sheet(nom)


def ecrire_log_feuille(ws, niveau: str, msg: str, horodatage: datetime) -> None:
    """VBA LogMsg (côté openpyxl)."""
    if ws["A1"].value in (None, ""):
        for col, titre in enumerate(("Horodatage", "Niveau", "Message"), start=1):
            ws.cell(row=1, column=col, value=titre)
    n = 1  # End(xlUp) sur une colonne vide renvoie la ligne 1
    for ligne in range(ws.max_row, 0, -1):
        if ws.cell(row=ligne, column=1).value not in (None, ""):
            n = ligne
            break
    n += 1
    ws.cell(row=n, column=1, value=horodatage).number_format = "dd/mm/yyyy hh:mm:ss"
    ws.cell(row=n, column=2, value=niveau)
    ws.cell(row=n, column=3, value=msg)


class Journal:
    """VBA LogMsg : écrit dans la feuille de log (via `ecrire`) et dans le logging Python."""

    def __init__(self, ecrire: Callable[[str, str], None]):
        self._ecrire = ecrire

    def ecrire(self, niveau: str, msg: str) -> None:
        logger.log(logging.ERROR if niveau == "ERREUR" else logging.INFO, "%s %s", niveau, msg)
        self._ecrire(niveau, msg)


def _en_jour(valeur) -> date | None:
    """Équivalent de CLng(valeur) pour comparer des dates (jour seulement)."""
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    return None


# --------------------------------------------------------------------------- 1. structure / données


def simuler_volume(reg: str, alea: random.Random) -> float:
    """VBA SimulerVolume."""
    base = {"Nord": 120, "Sud": 95, "Est": 80, "Ouest": 70}.get(reg, 60)
    return round(base * (0.85 + alea.random() * 0.3), 1)


def simuler_temperature(alea: random.Random) -> float:
    """VBA SimulerTemperature."""
    return round(2 + alea.random() * 14, 1)


def _largeurs_auto(ws, nb_colonnes: int) -> None:
    """Approximation openpyxl de Columns("A:E").AutoFit (largeur selon le texte le plus long)."""
    for col in range(1, nb_colonnes + 1):
        longueur = max((len(str(c.value)) for c in ws[openpyxl.utils.get_column_letter(col)]
                        if c.value is not None), default=8)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = longueur + 2


def ensure_structure(wb, sheet_data: str, sheet_visuel: str, table_name: str,
                     journal: Journal, aujourdhui: date, alea: random.Random) -> bool:
    """VBA EnsureStructure : crée les feuilles et le tableau initial (14 jours). Retourne True si créé."""
    ws_d = get_or_create_sheet(wb, sheet_data)
    get_or_create_sheet(wb, sheet_visuel)

    if table_name in ws_d.tables:
        return False

    # Cells.Clear (y compris d'éventuels tableaux orphelins)
    for nom in list(ws_d.tables):
        del ws_d.tables[nom]
    for row in ws_d.iter_rows():
        for cell in row:
            cell.value = None
            cell.style = "Normal"
    for col, titre in enumerate(EN_TETES, start=1):
        ws_d.cell(row=1, column=col, value=titre)
    ligne = 2
    for i in range(14, 0, -1):
        for r in regions():
            ws_d.cell(row=ligne, column=1, value=aujourdhui - timedelta(days=i))
            ws_d.cell(row=ligne, column=2, value=r)
            ws_d.cell(row=ligne, column=3, value=simuler_volume(r, alea))
            ws_d.cell(row=ligne, column=4, value=simuler_temperature(alea))
            ws_d.cell(row=ligne, column=5, value="Historique")
            ligne += 1
    ref = f"A1:E{ligne - 1}"
    tbl = Table(displayName=table_name, ref=ref, autoFilter=AutoFilter(ref=ref))
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws_d.add_table(tbl)
    for row in range(2, ligne):
        ws_d.cell(row=row, column=1).number_format = "dd/mm/yyyy"
    _largeurs_auto(ws_d, 5)
    journal.ecrire("INFO", "Structure initiale creee")
    return True


def ligne_existe(ws, tbl: Table, jour: date, reg: str) -> bool:
    """VBA LigneExiste : une ligne (date, région) existe-t-elle déjà dans le tableau ?"""
    _, min_row, _, max_row = openpyxl.utils.range_boundaries(tbl.ref)
    for row in range(min_row + 1, max_row + 1):
        if _en_jour(ws.cell(row=row, column=1).value) == jour and ws.cell(row=row, column=2).value == reg:
            return True
    return False


def ajouter_donnees_du_jour(wb, sheet_data: str, table_name: str, journal: Journal,
                            aujourdhui: date, alea: random.Random) -> int:
    """VBA AjouterDonneesDuJour : une ligne par région absente pour aujourd'hui. Retourne le nombre ajouté."""
    ws = wb[sheet_data]
    tbl = ws.tables[table_name]
    nb = 0
    for r in regions():
        if not ligne_existe(ws, tbl, aujourdhui, r):
            min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(tbl.ref)
            nouvelle = max_row + 1
            ws.cell(row=nouvelle, column=1, value=aujourdhui).number_format = "dd/mm/yyyy"
            ws.cell(row=nouvelle, column=2, value=r)
            ws.cell(row=nouvelle, column=3, value=simuler_volume(r, alea))
            ws.cell(row=nouvelle, column=4, value=simuler_temperature(alea))
            ws.cell(row=nouvelle, column=5, value="Phase de test")
            tbl.ref = f"A{min_row}:E{nouvelle}"
            if tbl.autoFilter is not None:
                tbl.autoFilter.ref = tbl.ref
            nb += 1
    journal.ecrire("INFO", f"{nb} ligne(s) ajoutee(s) pour le {aujourdhui:%d/%m/%Y}")
    return nb


# --------------------------------------------------------------------------- 2. filtre


def est_region_valide(s: str) -> bool:
    """VBA EstRegionValide : comparaison insensible à la casse (vbTextCompare)."""
    return any(r.lower() == s.lower() for r in regions())


def preparer_saisie_filtre(saisie: str) -> str:
    """VBA AppliquerFiltre (partie InputBox) : Trim$ puis validation.

    Retourne la région à filtrer, ou "" (aucun filtre : saisie vide ou région inconnue).
    Le MsgBox « Region inconnue » devient un avertissement journalisé.
    """
    saisie = saisie.strip(" ")
    if saisie == "":
        return ""
    if not est_region_valide(saisie):
        logger.warning("Region inconnue : %s", saisie)
        return ""
    return saisie


# --------------------------------------------------------------------------- 4. sauvegarde


def nom_fichier_version(file_title: str, file_suffix: str, extension: str, maintenant: datetime) -> str:
    """VBA SauvegarderVersion : titre_yyyymmdd_hhnnss_suffixe + extension."""
    return f"{file_title}_{maintenant:%Y%m%d_%H%M%S}_{file_suffix}{extension}"


def chemin_sauvegarde(dossier_sortie, nom: str) -> str:
    """VBA SauvegarderVersion : crée le dossier (MkDir) et renvoie le chemin absolu de la copie."""
    dossier = Path(dossier_sortie).resolve()
    dossier.mkdir(parents=True, exist_ok=True)
    return str(dossier / nom)


# --------------------------------------------------------------------------- 5. mail


def envoyer_mail(chemin_pj: str, mail_to: str, mail_cc: str, mode_mail: str, journal: Journal,
                 aujourdhui: date, envoyer_mail_fn: Callable | None = None) -> bool:
    """VBA EnvoyerMail. `mode_mail` : envoyer, afficher ou aucun. Retourne True si le mail est parti/affiché."""
    if mode_mail == "aucun":
        logger.info("Mail non envoyé (mode aucun)")
        return False
    envoyer_mail_fn = envoyer_mail_fn or connecteur_outlook.envoyer_mail
    sujet = f"[PHASE DE TEST] Mise a jour suivi distribution - {aujourdhui:%d/%m/%Y}"
    corps = ("Bonjour,\r\n\r\n"
             "Ceci est un envoi automatique en phase de test.\r\n"
             "Le graphique et les donnees ont ete mis a jour.\r\n"
             f"Fichier : {Path(chemin_pj).name}\r\n\r\n"
             "Cordialement,\r\nRobot VBATF")
    try:
        envoyer_mail_fn(mail_to, mail_cc, sujet, corps, [str(Path(chemin_pj).resolve())], mode_mail)
    except Exception as erreur:  # VBA : On Error GoTo MailErr
        journal.ecrire("ERREUR", f"Mail impossible : {erreur}")
        return False
    journal.ecrire("INFO", ("Mail envoye a " if mode_mail == "envoyer" else "Brouillon cree pour ") + mail_to)
    return True


# --------------------------------------------------------------------------- pipeline


def lancer_pipeline_complet(classeur, dossier_sortie=DOSSIER_SORTIE, sheet_visuel=SHEET_VISUEL,
                            sheet_data=SHEET_DATA, sheet_log=SHEET_LOG, table_name=TABLE_NAME,
                            file_title=FILE_TITLE, file_suffix=FILE_SUFFIX, mail_to=MAIL_TO,
                            mail_cc=MAIL_CC, mode_mail="envoyer", saisie="",
                            envoyer_mail_fn: Callable | None = None,
                            executer_excel: Callable | None = None,
                            maintenant: Callable[[], datetime] = datetime.now,
                            alea: random.Random | None = None) -> str | None:
    """VBA LancerPipelineComplet. Retourne le chemin de la copie sauvegardée, ou None en cas d'erreur.

    Le classeur est modifié en place (comme ThisWorkbook). Application.ScreenUpdating et
    DisplayAlerts sont supprimés ; MsgBox devient un message de log.
    """
    alea = alea or random.Random()
    chemin_classeur = Path(classeur)
    wb = openpyxl.load_workbook(chemin_classeur, keep_vba=chemin_classeur.suffix.lower() == ".xlsm")

    def log_openpyxl(niveau, msg):
        ecrire_log_feuille(get_or_create_sheet(wb, sheet_log), niveau, msg, maintenant())

    journal = Journal(log_openpyxl)
    aujourdhui = maintenant().date()

    try:
        ensure_structure(wb, sheet_data, sheet_visuel, table_name, journal, aujourdhui, alea)
        ajouter_donnees_du_jour(wb, sheet_data, table_name, journal, aujourdhui, alea)
    except Exception as erreur:  # VBA : ErrHandler
        journal.ecrire("ERREUR", f"{type(erreur).__name__} - {erreur}")
        wb.save(chemin_classeur)
        return None
    wb.save(chemin_classeur)

    region = preparer_saisie_filtre(saisie)
    nom = nom_fichier_version(file_title, file_suffix, chemin_classeur.suffix, maintenant())
    chemin_fichier = chemin_sauvegarde(dossier_sortie, nom)
    resultat: list[str | None] = [None]

    def etapes_excel(xwb):
        def log_excel(niveau, msg):
            connecteur_excel.ecrire_log_excel(xwb, sheet_log, niveau, msg, maintenant())

        journal_xl = Journal(log_excel)
        try:
            connecteur_excel.appliquer_filtre_excel(xwb, sheet_data, table_name, region)
            if region:
                journal_xl.ecrire("INFO", f"Filtre applique : {region}")
            connecteur_excel.creer_visuel_excel(xwb, sheet_visuel, sheet_data, table_name, maintenant())
            journal_xl.ecrire("INFO", "Visuel regenere")
            connecteur_excel.sauvegarder_copie_excel(xwb, chemin_fichier)
            journal_xl.ecrire("INFO", f"Sauvegarde : {chemin_fichier}")
            envoyer_mail(chemin_fichier, mail_to, mail_cc, mode_mail, journal_xl,
                         maintenant().date(), envoyer_mail_fn)
            journal_xl.ecrire("INFO", f"Pipeline termine : {chemin_fichier}")
            logger.info("Pipeline termine.\n%s", chemin_fichier)
            resultat[0] = chemin_fichier
        except Exception as erreur:  # VBA : ErrHandler (ErreurExcel incluse, message complet)
            journal_xl.ecrire("ERREUR", str(erreur))

    (executer_excel or connecteur_excel.executer_dans_excel)(chemin_classeur, etapes_excel)
    return resultat[0]


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description="VBA LancerPipelineComplet", allow_abbrev=False)
    parser.add_argument("classeur", help="Classeur à traiter (remplace ThisWorkbook)")
    parser.add_argument("--dossier-sortie", default=DOSSIER_SORTIE)
    parser.add_argument("--sheet-visuel", default=SHEET_VISUEL)
    parser.add_argument("--sheet-data", default=SHEET_DATA)
    parser.add_argument("--sheet-log", default=SHEET_LOG)
    parser.add_argument("--table-name", default=TABLE_NAME)
    parser.add_argument("--file-title", default=FILE_TITLE)
    parser.add_argument("--file-suffix", default=FILE_SUFFIX)
    parser.add_argument("--mail-to", default=MAIL_TO)
    parser.add_argument("--mail-cc", default=MAIL_CC)
    parser.add_argument("--mode-mail", choices=["envoyer", "afficher", "aucun"], default="envoyer")
    parser.add_argument("--saisie", default="", help="Région à filtrer (vide = aucun filtre)")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    chemin = lancer_pipeline_complet(
        args.classeur, args.dossier_sortie, args.sheet_visuel, args.sheet_data, args.sheet_log,
        args.table_name, args.file_title, args.file_suffix, args.mail_to, args.mail_cc,
        args.mode_mail, args.saisie)
    if chemin is None:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
