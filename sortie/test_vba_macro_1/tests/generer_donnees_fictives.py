"""Crée tests/fixtures/*.xlsx (données inventées, même forme que test_vba_macro_1.xlsm). Idempotent."""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

DOSSIER = Path(__file__).parent / "fixtures"
REGIONS = ["Nord", "Sud", "Est", "Ouest", "Centre"]
# Volumes inventés par jour (ordre : Nord, Sud, Est, Ouest, Centre)
JOURS = {
    date(2026, 10, 20): [120.0, 95.0, 80.0, 70.0, 60.0],
    date(2026, 10, 21): [110.5, 100.0, 82.5, 65.0, 58.0],
    date(2026, 10, 22): [130.0, 90.0, 78.0, 72.5, 61.5],
}


def _entetes_log(wb):
    ws = wb.create_sheet("Log")
    ws.append(["Horodatage", "Niveau", "Message"])


def creer_classeur_complet(chemin: Path) -> None:
    """Classeur avec Dashboard, Feuille2 (tableau tblDonnees, 15 lignes), Feuille1 et Log."""
    wb = Workbook()
    wb.active.title = "Dashboard"
    ws = wb.create_sheet("Feuille2")
    ws.append(["Date", "Region", "Volume_GWh", "Temperature_C", "Statut"])
    temp = 5.0
    for d, volumes in JOURS.items():
        for reg, vol in zip(REGIONS, volumes):
            ws.append([d, reg, vol, temp, "Historique"])
            ws.cell(row=ws.max_row, column=1).number_format = "dd/mm/yyyy"
            temp += 0.5
    tbl = Table(displayName="tblDonnees", ref=f"A1:E{ws.max_row}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)
    wb.create_sheet("Feuille1")
    _entetes_log(wb)
    wb.save(chemin)


def creer_classeur_vide(chemin: Path) -> None:
    """Classeur sans tableau : le pipeline doit recréer la structure initiale."""
    wb = Workbook()
    wb.active.title = "Dashboard"
    wb.save(chemin)


def generer(dossier: Path = DOSSIER) -> Path:
    dossier.mkdir(parents=True, exist_ok=True)
    creer_classeur_complet(dossier / "classeur_complet.xlsx")
    creer_classeur_vide(dossier / "classeur_vide.xlsx")
    return dossier


if __name__ == "__main__":
    print(generer(Path(sys.argv[1]) if len(sys.argv) > 1 else DOSSIER))
