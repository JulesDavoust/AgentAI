import random
from datetime import date, datetime
from email import message_from_bytes, policy

import pytest
from openpyxl import load_workbook

import marco_25_10 as m

AUJOURDHUI = date(2026, 10, 25)
MAINTENANT = datetime(2026, 10, 25, 11, 15, 30)


class RngMilieu:
    """Faux générateur : random() vaut toujours 0.5 (facteur 1.0 pour les volumes)."""

    def random(self):
        return 0.5


@pytest.fixture
def wb(classeur_complet):
    return load_workbook(classeur_complet)


# --- simulation ------------------------------------------------------------
@pytest.mark.parametrize("reg,attendu", [("Nord", 120.0), ("Sud", 95.0), ("Est", 80.0),
                                         ("Ouest", 70.0), ("Centre", 60.0), ("Inconnue", 60.0)])
def test_simuler_volume_milieu(reg, attendu):
    # random=0.5 -> base * (0.85 + 0.15) = base
    assert m.simuler_volume(reg, RngMilieu()) == attendu


def test_simuler_volume_bornes():
    rng = random.Random(1)
    for _ in range(200):
        v = m.simuler_volume("Nord", rng)
        assert 102.0 <= v <= 138.0  # 120 * [0.85 ; 1.15]


def test_simuler_temperature():
    assert m.simuler_temperature(RngMilieu()) == 9.0  # 2 + 0.5*14
    rng = random.Random(2)
    assert all(2.0 <= m.simuler_temperature(rng) <= 16.0 for _ in range(200))


# --- structure -----------------------------------------------------------------
def test_ensure_structure_cree_historique(classeur_vide):
    wb = load_workbook(classeur_vide)
    assert m.ensure_structure(wb, RngMilieu(), AUJOURDHUI, MAINTENANT) is True
    assert {"Feuille1", "Feuille2", "Log"} <= set(wb.sheetnames)
    ws = wb["Feuille2"]
    assert [c.value for c in ws[1]] == m.ENTETES
    assert ws.tables["tblDonnees"].ref == "A1:E71"  # 14 jours * 5 régions + en-tête
    assert m._en_date(ws["A2"].value) == date(2026, 10, 11)  # aujourd'hui - 14
    assert m._en_date(ws["A71"].value) == date(2026, 10, 24)  # aujourd'hui - 1
    assert ws["E2"].value == "Historique"
    assert ws["C2"].value == 120.0 and ws["B2"].value == "Nord"


def test_ensure_structure_existante_ne_change_rien(wb):
    avant = m.lire_table(wb)
    assert m.ensure_structure(wb, RngMilieu(), AUJOURDHUI, MAINTENANT) is False
    assert m.lire_table(wb).equals(avant)


def test_get_or_create_sheet(wb):
    assert m.get_or_create_sheet(wb, "Feuille2") is wb["Feuille2"]
    n = len(wb.sheetnames)
    ws = m.get_or_create_sheet(wb, "Nouvelle")
    assert wb.sheetnames[-1] == "Nouvelle" and len(wb.sheetnames) == n + 1 and ws.title == "Nouvelle"


# --- ajout du jour ---------------------------------------------------------------
def test_ajouter_donnees_du_jour(wb):
    assert m.ajouter_donnees_du_jour(wb, RngMilieu(), AUJOURDHUI, MAINTENANT) == 5
    ws = wb["Feuille2"]
    assert ws.tables["tblDonnees"].ref == "A1:E21"  # 15 + 5 lignes
    # lignes 17..21 = nouvelles lignes, dans l'ordre Nord, Sud, ... : Sud en ligne 18
    assert ws["B18"].value == "Sud" and ws["C18"].value == 95.0
    assert ws["E21"].value == "Phase de test"
    assert wb["Log"]["C2"].value == "5 ligne(s) ajoutee(s) pour le 25/10/2026"


def test_ajouter_donnees_du_jour_idempotent(wb):
    m.ajouter_donnees_du_jour(wb, RngMilieu(), AUJOURDHUI, MAINTENANT)
    assert m.ajouter_donnees_du_jour(wb, RngMilieu(), AUJOURDHUI, MAINTENANT) == 0
    assert len(m.lire_table(wb)) == 20


def test_ajouter_partiel(wb):
    # Les 5 régions existent déjà le 21/10 : rien à ajouter
    assert m.ajouter_donnees_du_jour(wb, RngMilieu(), date(2026, 10, 21), MAINTENANT) == 0


def test_ligne_existe(wb):
    assert m.ligne_existe(wb, date(2026, 10, 20), "Nord")
    assert not m.ligne_existe(wb, date(2026, 10, 20), "nord")  # '=' VBA sensible à la casse
    assert not m.ligne_existe(wb, date(2026, 10, 25), "Nord")


# --- filtre ------------------------------------------------------------------------
def test_est_region_valide():
    assert m.est_region_valide("nord") and m.est_region_valide("OUEST")
    assert not m.est_region_valide("Lune") and not m.est_region_valide("")


def test_appliquer_filtre_valide(wb):
    assert m.appliquer_filtre(wb, "  sud ", MAINTENANT) == "sud"
    ws = wb["Feuille2"]
    visibles = [r for r in range(2, 17) if not ws.row_dimensions[r].hidden]
    assert [ws.cell(row=r, column=2).value for r in visibles] == ["Sud"] * 3
    assert ws.tables["tblDonnees"].autoFilter.filterColumn[0].colId == 1
    assert wb["Log"]["C2"].value == "Filtre applique : sud"


@pytest.mark.parametrize("saisie", [None, "", "   "])
def test_appliquer_filtre_vide(wb, saisie):
    m.appliquer_filtre(wb, "Nord", MAINTENANT)
    assert m.appliquer_filtre(wb, saisie, MAINTENANT) is None
    ws = wb["Feuille2"]
    assert not any(ws.row_dimensions[r].hidden for r in range(2, 17))  # ShowAllData
    assert ws.tables["tblDonnees"].autoFilter.filterColumn == []


def test_appliquer_filtre_region_inconnue(wb):
    assert m.appliquer_filtre(wb, "Lune", MAINTENANT) is None
    assert not any(wb["Feuille2"].row_dimensions[r].hidden for r in range(2, 17))


# --- KPI / visuel ----------------------------------------------------------------------
def test_calculer_kpi(wb):
    k = m.calculer_kpi(m.lire_table(wb))
    # 425 + 416 + 432 = 1273 ; 3 jours distincts ; max 130.0 = Nord
    assert k["volume_total"] == pytest.approx(1273.0)
    assert k["moyenne_par_jour"] == pytest.approx(1273.0 / 3)
    assert k["region_pic"] == "Nord"


def test_calculer_kpi_vide_et_une_ligne(wb):
    df = m.lire_table(wb)
    assert m.calculer_kpi(df.iloc[0:0])["region_pic"] is None
    k = m.calculer_kpi(df.iloc[[2]])  # Est, 80.0
    assert k == {"volume_total": 80.0, "moyenne_par_jour": 80.0, "region_pic": "Est"}


def test_calculer_kpi_egalite_premiere_occurrence(wb):
    df = m.lire_table(wb)
    # Sud (ligne 1, le 20/10) = Nord (ligne 10, le 22/10) = 130 ; MATCH renvoie la 1re -> Sud
    df.loc[1, "Volume_GWh"] = 130.0
    assert m.calculer_kpi(df)["region_pic"] == "Sud"


def test_tableau_croise(wb):
    t = m.tableau_croise(m.lire_table(wb))
    assert t.shape == (3, 5)
    assert t.loc[date(2026, 10, 21), "Sud"] == 100.0


def test_creer_visuel(wb, tmp_path):
    png = m.creer_visuel(wb, tmp_path / "out", MAINTENANT)
    assert png.exists() and png.stat().st_size > 0
    ws = wb["Feuille1"]
    assert ws["A1"].value == "Suivi distribution gaz - mis a jour le 25/10/2026 11:15"
    assert ws["B2"].value == "=SUM(tblDonnees[Volume_GWh])"
    assert ws["H2"].value.startswith("=INDEX(tblDonnees[Region],MATCH(MAX(")
    assert ws["A1"].font.bold and ws["A1"].font.size == 14
    assert [ws.cell(row=5, column=c).value for c in range(1, 8)] == [
        "Volume (GWh)", "Centre", "Est", "Nord", "Ouest", "Sud", "Total general"]
    assert ws["G6"].value == 425.0  # total du 20/10 (colonne G = Total general)


def test_creer_visuel_ecrase_ancien_contenu(wb, tmp_path):
    wb["Feuille1"]["Z50"] = "vieux"
    m.creer_visuel(wb, tmp_path, MAINTENANT)
    assert wb["Feuille1"]["Z50"].value is None


# --- sauvegarde / mail -------------------------------------------------------------------
def test_sauvegarder_version(wb, tmp_path):
    chemin = m.sauvegarder_version(wb, tmp_path / "sous" / "dossier", MAINTENANT)
    assert chemin.name == "VBA_Natran_Test_20261025_111530_VBATF.xlsx"
    assert chemin.exists()
    assert "Feuille2" in load_workbook(chemin).sheetnames


def test_composer_mail(tmp_path):
    msg = m.composer_mail(tmp_path / "x" / "fichier.xlsx", AUJOURDHUI)
    assert msg["To"] == m.MAIL_TO and msg["Cc"] is None
    assert msg["Subject"] == "[PHASE DE TEST] Mise a jour suivi distribution - 25/10/2026"
    assert "Fichier : fichier.xlsx" in msg.get_content()


def test_envoyer_mail_ecrit_brouillon(tmp_path):
    pj = tmp_path / "pj.xlsx"
    pj.write_bytes(b"abc")
    eml = m.envoyer_mail(pj, tmp_path / "mails", AUJOURDHUI)
    msg = message_from_bytes(eml.read_bytes(), policy=policy.default)
    assert msg["To"] == m.MAIL_TO
    assert [p.get_filename() for p in msg.iter_attachments()] == ["pj.xlsx"]


# --- pipeline -----------------------------------------------------------------------------
def test_pipeline_complet(classeur_complet, tmp_path):
    sortie = tmp_path / "sorties"
    chemin = m.lancer_pipeline_complet(classeur_complet, sortie, "Nord", RngMilieu(), MAINTENANT)
    assert chemin.exists() and (sortie / "volumes_par_region.png").stat().st_size > 0
    assert list(sortie.glob("*.eml"))
    wb = load_workbook(chemin)
    assert len(m.lire_table(wb)) == 20
    messages = [r[2].value for r in wb["Log"].iter_rows(min_row=2)]
    assert messages[-1].startswith("Pipeline termine")
    assert "Filtre applique : Nord" in messages


def test_pipeline_classeur_absent(tmp_path):
    chemin = m.lancer_pipeline_complet(tmp_path / "nexiste_pas.xlsx", tmp_path / "s", None,
                                       random.Random(3), MAINTENANT)
    assert len(m.lire_table(load_workbook(chemin))) == 14 * 5 + 5


def test_pipeline_ne_modifie_pas_la_source(classeur_complet, tmp_path):
    avant = classeur_complet.read_bytes()
    m.lancer_pipeline_complet(classeur_complet, tmp_path, None, RngMilieu(), MAINTENANT)
    assert classeur_complet.read_bytes() == avant


def test_pipeline_erreur_relancee(tmp_path):
    corrompu = tmp_path / "mauvais.xlsx"
    corrompu.write_text("pas un classeur")
    with pytest.raises(Exception):
        m.lancer_pipeline_complet(corrompu, tmp_path / "s")


# --- journal ---------------------------------------------------------------------------------
def test_log_msg_ajoute_en_dessous(wb):
    m.log_msg(wb, "INFO", "a", MAINTENANT)
    m.log_msg(wb, "ERREUR", "b", MAINTENANT)
    ws = wb["Log"]
    assert [c.value for c in ws[1]] == ["Horodatage", "Niveau", "Message"]
    assert (ws["B2"].value, ws["C3"].value, ws["A2"].value) == ("INFO", "b", MAINTENANT)


def test_log_msg_feuille_absente():
    from openpyxl import Workbook
    w = Workbook()
    m.log_msg(w, "INFO", "x", MAINTENANT)
    assert w["Log"]["C2"].value == "x"
