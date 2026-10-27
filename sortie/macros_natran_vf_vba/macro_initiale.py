"""Traduction du module VBA ``MacroInitiale`` : pipeline gaz (structure, données du jour, filtre,
visuel avec tableau croisé et graphique, sauvegarde datée, mail).

Les parties qu'openpyxl ne sait pas faire (tableau croisé, graphique, filtre appliqué, AutoFit)
sont dans ``connecteurs/excel.py`` ; l'envoi du mail dans ``connecteurs/outlook.py``.
"""

from __future__ import annotations

import argparse
import logging
import random
import sys
from datetime import datetime, time, timedelta
from pathlib import Path

from openpyxl.styles import Font

import commun
from commun import ErreurVBA, Resultat, log_msg
from connecteurs import excel as connecteur_excel
from bibliotheque.excel_com import preparer_feuille

SHEET_VISUEL = "Feuille1"
SHEET_DATA = "Feuille2"
TABLE_NAME = "tblDonnees"
FILE_TITLE = "VBA_Natran_Test"
FILE_SUFFIX = "VBATF"
MAIL_TO = "jdavoust201@headmind.com"
MAIL_CC = ""
REGIONS = ("Nord", "Sud", "Est", "Ouest", "Centre")


def regions() -> tuple:
    """VBA ``Regions``"""
    return REGIONS


def simuler_volume(reg: str, rng) -> float:
    """VBA ``SimulerVolume``"""
    base = {"Nord": 120, "Sud": 95, "Est": 80, "Ouest": 70}.get(reg, 60)
    return round(base * (0.85 + rng.random() * 0.3), 1)


def simuler_temperature(rng) -> float:
    """VBA ``SimulerTemperature``"""
    return round(2 + rng.random() * 14, 1)


def _minuit(d) -> datetime:
    return datetime.combine(d, time())


# ----------------------------------------------------------------------------- 1. structure / données

def ensure_structure(wb, maintenant: datetime, rng) -> bool:
    """VBA ``EnsureStructure`` : crée Feuille2 + tableau tblDonnees (14 jours d'historique) s'il
    n'existe pas. Retourne True si la structure a été créée (AutoFit des colonnes à faire)."""
    ws_d = commun.get_or_create_sheet(wb, SHEET_DATA)
    commun.get_or_create_sheet(wb, SHEET_VISUEL)
    if TABLE_NAME in ws_d.tables:
        return False

    preparer_feuille(ws_d)  # wsD.Cells.Clear
    for i, titre in enumerate(("Date", "Region", "Volume_GWh", "Temperature_C", "Statut"), 1):
        ws_d.cell(row=1, column=i, value=titre)
    ligne = 2
    for i in range(14, 0, -1):
        for r in REGIONS:
            ws_d.cell(row=ligne, column=1, value=_minuit(maintenant.date() - timedelta(days=i)))
            ws_d.cell(row=ligne, column=2, value=r)
            ws_d.cell(row=ligne, column=3, value=simuler_volume(r, rng))
            ws_d.cell(row=ligne, column=4, value=simuler_temperature(rng))
            ws_d.cell(row=ligne, column=5, value="Historique")
            ligne += 1
    commun.creer_table(ws_d, TABLE_NAME, f"A1:E{ligne - 1}")
    commun.formater_colonne(ws_d, "A", "dd/mm/yyyy")
    log_msg(wb, "INFO", "Structure initiale creee", maintenant)
    return True


def ajouter_donnees_du_jour(wb, maintenant: datetime, rng) -> int:
    """VBA ``AjouterDonneesDuJour`` : une ligne par région absente pour la date du jour."""
    ws = commun.feuille(wb, SHEET_DATA)
    lo = commun.table_de(ws, TABLE_NAME)
    nb = 0
    for r in REGIONS:
        if not ligne_existe(ws, lo, maintenant.date(), r):
            ligne = commun.ajouter_ligne_table(ws, lo, [
                _minuit(maintenant.date()), r, simuler_volume(r, rng), simuler_temperature(rng), "Phase de test"])
            ws.cell(row=ligne, column=1).number_format = "dd/mm/yyyy"
            nb += 1
    log_msg(wb, "INFO", f"{nb} ligne(s) ajoutee(s) pour le {commun.fmt_date(maintenant)}", maintenant)
    return nb


def ligne_existe(ws, lo, d, reg: str) -> bool:
    """VBA ``LigneExiste``"""
    for valeurs in commun.lignes_table(ws, lo):
        if commun.en_jour(valeurs[0]) == d and valeurs[1] == reg:
            return True
    return False


# ----------------------------------------------------------------------------- 2. filtre

def est_region_valide(s: str) -> bool:
    """VBA ``EstRegionValide`` (``StrComp`` en comparaison de texte : insensible à la casse)."""
    return any(r.casefold() == s.casefold() for r in REGIONS)


def appliquer_filtre(wb, saisie: str, maintenant: datetime | None = None) -> str | None:
    """VBA ``AppliquerFiltre`` (partie classeur). ``saisie`` remplace l'InputBox (vide = aucun
    filtre). Retourne la région à filtrer (telle que saisie), ou None. Le filtre lui-même est
    appliqué par Excel (``connecteurs.excel.appliquer_filtre_excel``)."""
    commun.table_de(commun.feuille(wb, SHEET_DATA), TABLE_NAME)
    saisie = (saisie or "").strip()
    if saisie == "":
        return None
    if not est_region_valide(saisie):
        commun.journal.warning("Region inconnue : %s", saisie)  # MsgBox vbExclamation
        return None
    log_msg(wb, "INFO", f"Filtre applique : {saisie}", maintenant)
    return saisie


# ----------------------------------------------------------------------------- 3. visuel

def creer_visuel_cellules(wb, maintenant: datetime) -> None:
    """VBA ``CreerVisuel`` (partie openpyxl) : nettoyage de Feuille1, titre, KPI en formules.
    Le tableau croisé, le graphique et l'AutoFit sont faits par Excel."""
    ws_v = commun.feuille(wb, SHEET_VISUEL)
    commun.table_de(commun.feuille(wb, SHEET_DATA), TABLE_NAME)
    preparer_feuille(ws_v)  # ChartObjects.Delete, PivotTables.TableRange2.Clear, Cells.Clear

    ws_v["A1"] = "Suivi distribution gaz - mis a jour le " + commun.fmt_date_heure(maintenant)
    ws_v["A1"].font = Font(bold=True, size=14)
    ws_v["A2"] = "Volume total (GWh)"
    ws_v["B2"] = f"=SUM({TABLE_NAME}[Volume_GWh])"
    ws_v["D2"] = "Moyenne / jour (GWh)"
    ws_v["E2"] = (f"=SUM({TABLE_NAME}[Volume_GWh])/SUMPRODUCT(1/COUNTIF({TABLE_NAME}[Date],{TABLE_NAME}[Date]))")
    ws_v["G2"] = "Region du pic"
    for ref in ("A2", "D2", "G2"):
        ws_v[ref].font = Font(bold=True)
    for ref in ("B2", "E2"):
        ws_v[ref].number_format = "#,##0.0"
    ws_v["H2"] = (f"=INDEX({TABLE_NAME}[Region],MATCH(MAX({TABLE_NAME}[Volume_GWh]),{TABLE_NAME}[Volume_GWh],0))")
    commun.activer_feuille(wb, ws_v)
    log_msg(wb, "INFO", "Visuel regenere", maintenant)


# ----------------------------------------------------------------------------- 4. sauvegarde / 5. mail

def sauvegarder_version(wb, dossier, extension: str, maintenant: datetime) -> Path:
    """VBA ``SauvegarderVersion`` : crée le dossier, enregistre la copie ``titre_date_VBATF`` et
    retourne son chemin absolu (``SaveCopyAs``)."""
    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    chemin = (dossier / commun.nom_version(FILE_TITLE, maintenant, FILE_SUFFIX, extension)).resolve()
    wb.save(chemin)
    return chemin


def envoyer_mail_pipeline(wb, chemin_pj: Path, mail_to: str, mail_cc: str, mode_mail: str,
                          maintenant: datetime, envoyer_mail=None) -> bool:
    """VBA ``EnvoyerMail`` (MacroInitiale)."""
    sujet = "[PHASE DE TEST] Mise a jour suivi distribution - " + commun.fmt_date(maintenant)
    corps = ("Bonjour,\r\n\r\n"
             "Ceci est un envoi automatique en phase de test.\r\n"
             "Le graphique et les donnees ont ete mis a jour.\r\n"
             f"Fichier : {Path(chemin_pj).name}\r\n\r\n"
             "Cordialement,\r\nRobot VBATF")
    return commun.expedier_mail(wb, envoyer_mail, mode_mail, mail_to, mail_cc, sujet, corps, [chemin_pj], maintenant)


# ----------------------------------------------------------------------------- point d'entrée

def lancer_pipeline_complet(classeur, dossier_sortie="sortie", mail_to=MAIL_TO, mail_cc=MAIL_CC,
                            mode_mail="envoyer", saisie="", *, envoyer_mail=None, excel=None, rng=None,
                            maintenant=datetime.now) -> Resultat:
    """VBA ``LancerPipelineComplet`` : structure, données du jour, filtre, visuel, sauvegarde, mail.

    ``classeur`` = ``ThisWorkbook`` (jamais modifié : le résultat est la copie enregistrée,
    comme ``SaveCopyAs``). ``saisie`` = InputBox du filtre. ``envoyer_mail`` et ``excel``
    sont injectables (tests)."""
    rng = rng or random.Random()
    excel = excel or connecteur_excel.finaliser_initiale
    wb = commun.charger_classeur(classeur)
    try:
        creee = ensure_structure(wb, maintenant(), rng)
        ajouter_donnees_du_jour(wb, maintenant(), rng)
        region = appliquer_filtre(wb, saisie, maintenant())
        creer_visuel_cellules(wb, maintenant())
        chemin = sauvegarder_version(wb, dossier_sortie, Path(classeur).suffix, maintenant())
        excel(chemin, creee, region)
        log_msg(wb, "INFO", f"Sauvegarde : {chemin}", maintenant())
        envoyer_mail_pipeline(wb, chemin, mail_to, mail_cc, mode_mail, maintenant(), envoyer_mail)
        log_msg(wb, "INFO", f"Pipeline termine : {chemin}", maintenant())
        message = f"Pipeline termine.\n{chemin}"
        commun.journal.info(message)  # MsgBox vbInformation
        return Resultat(True, [chemin], message)
    except Exception as erreur:  # ErrHandler : consigne dans le Log et affiche l'erreur
        message = f"Erreur : {getattr(erreur, 'description', erreur)}"
        log_msg(wb, "ERREUR", f"{commun.numero_erreur(erreur)} - {getattr(erreur, 'description', erreur)}",
                maintenant())
        commun.journal.error(message)
        return Resultat(False, [], message)


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description="VBA MacroInitiale.LancerPipelineComplet")
    p.add_argument("classeur", help="classeur de départ (remplace ThisWorkbook)")
    p.add_argument("--dossier-sortie", default="sortie", help="remplace SaveFolder()")
    p.add_argument("--mail-to", default=MAIL_TO)
    p.add_argument("--mail-cc", default=MAIL_CC)
    p.add_argument("--mode-mail", choices=["envoyer", "afficher", "aucun"], default="envoyer",
                   help="remplace SEND_MAIL (True = envoyer, False = afficher)")
    p.add_argument("--saisie", default="", help="région à filtrer (InputBox ; vide = aucun filtre)")
    a = p.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    res = lancer_pipeline_complet(a.classeur, a.dossier_sortie, a.mail_to, a.mail_cc, a.mode_mail, a.saisie)
    return 0 if res.ok else 1


if __name__ == "__main__":
    sys.exit(main())
