"""Tests de macro_initiale.py (valeurs attendues calculées à la main, commentées)."""

import logging
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

import commun
import macro_initiale as mi
from commun import ErreurVBA

NOW = datetime(2026, 10, 27, 14, 30, 5)


def _logs(wb):
    return [wb["Log"].cell(row=r, column=3).value for r in range(2, wb["Log"].max_row + 1)]


# ------------------------------------------------------------------ données

def test_ensure_structure_cree_14_jours_de_5_regions(copie, rng):
    wb = openpyxl.load_workbook(copie("classeur_complet.xlsx"))
    assert mi.ensure_structure(wb, NOW, rng) is True
    ws = wb["Feuille2"]
    # 14 jours x 5 régions = 70 lignes de données + en-tête => A1:E71
    assert ws.tables["tblDonnees"].ref == "A1:E71"
    assert [c.value for c in ws[1]] == ["Date", "Region", "Volume_GWh", "Temperature_C", "Statut"]
    assert ws["A2"].value == datetime(2026, 10, 13)      # Date - 14
    assert ws["A71"].value == datetime(2026, 10, 26)     # Date - 1
    assert ws["A2"].number_format == "dd/mm/yyyy"
    # rng = 0.5 : volume = base * (0.85 + 0.15) = base ; température = 2 + 7 = 9
    assert (ws["B2"].value, ws["C2"].value, ws["D2"].value, ws["E2"].value) == ("Nord", 120.0, 9.0, "Historique")
    assert ws["C6"].value == 60.0                         # Centre : base 60 (Case Else)
    assert "Feuille1" in wb.sheetnames and "Structure initiale creee" in _logs(wb)


def test_ensure_structure_ne_touche_pas_un_tableau_existant(copie, rng):
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    assert mi.ensure_structure(wb, NOW, rng) is False
    assert wb["Feuille2"].tables["tblDonnees"].ref == "A1:E13"


def test_ajouter_donnees_du_jour_complete_les_regions_manquantes(copie, rng):
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    # le 27/10 contient déjà Nord et Sud : 3 régions à ajouter (Est, Ouest, Centre)
    assert mi.ajouter_donnees_du_jour(wb, NOW, rng) == 3
    ws = wb["Feuille2"]
    assert ws.tables["tblDonnees"].ref == "A1:E16"
    assert ws.tables["tblDonnees"].autoFilter.ref == "A1:E16"
    assert [ws.cell(row=r, column=2).value for r in (14, 15, 16)] == ["Est", "Ouest", "Centre"]
    assert ws["E14"].value == "Phase de test" and ws["A14"].value == datetime(2026, 10, 27)
    assert "3 ligne(s) ajoutee(s) pour le 27/10/2026" in _logs(wb)
    assert mi.ajouter_donnees_du_jour(wb, NOW, rng) == 0          # idempotent : plus rien à ajouter


def test_ajouter_donnees_sans_tableau_erreur_9(copie, rng):
    wb = openpyxl.load_workbook(copie("classeur_complet.xlsx"))
    wb.create_sheet("Feuille2")
    with pytest.raises(ErreurVBA) as e:
        mi.ajouter_donnees_du_jour(wb, NOW, rng)
    assert e.value.numero == 9


def test_ligne_existe(copie):
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    ws = wb["Feuille2"]
    lo = ws.tables["tblDonnees"]
    assert mi.ligne_existe(ws, lo, datetime(2026, 10, 27).date(), "Sud")
    assert not mi.ligne_existe(ws, lo, datetime(2026, 10, 27).date(), "Est")
    assert not mi.ligne_existe(ws, lo, datetime(2026, 10, 24).date(), "Nord")
    assert not mi.ligne_existe(ws, lo, datetime(2026, 10, 27).date(), "nord")   # comparaison binaire


def test_simuler_valeurs_bornes():
    class Zero:
        def random(self):
            return 0.0

    class Presque1:
        def random(self):
            return 0.999999

    assert mi.simuler_volume("Nord", Zero()) == 102.0       # 120 * 0.85
    assert mi.simuler_volume("Autre", Zero()) == 51.0       # Case Else : base 60 * 0.85
    assert mi.simuler_volume("Nord", Presque1()) == 138.0   # 120 * 1.15
    assert mi.simuler_temperature(Zero()) == 2.0
    assert mi.simuler_temperature(Presque1()) == 16.0


# ------------------------------------------------------------------ filtre

@pytest.mark.parametrize("saisie, attendu", [("", None), ("   ", None), ("Atlantide", None), ("sud", "sud"),
                                              (" Nord ", "Nord")])
def test_appliquer_filtre(copie, saisie, attendu):
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    assert mi.appliquer_filtre(wb, saisie, NOW) == attendu
    if attendu:
        assert f"Filtre applique : {attendu}" in _logs(wb)


def test_est_region_valide():
    assert mi.est_region_valide("OUEST") and not mi.est_region_valide("Paris") and not mi.est_region_valide("")


# ------------------------------------------------------------------ visuel

def test_creer_visuel_cellules(copie):
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    wb["Feuille1"]["Z9"] = "ancien contenu"
    mi.creer_visuel_cellules(wb, NOW)
    ws = wb["Feuille1"]
    assert ws["A1"].value == "Suivi distribution gaz - mis a jour le 27/10/2026 14:30"
    assert ws["A1"].font.b and ws["A1"].font.sz == 14
    assert ws["Z9"].value is None                                   # Cells.Clear
    assert ws["B2"].value == "=SUM(tblDonnees[Volume_GWh])"
    assert ws["E2"].value.startswith("=SUM(tblDonnees[Volume_GWh])/SUMPRODUCT(1/COUNTIF(")
    assert ws["H2"].value == "=INDEX(tblDonnees[Region],MATCH(MAX(tblDonnees[Volume_GWh]),tblDonnees[Volume_GWh],0))"
    assert ws["B2"].number_format == "#,##0.0" and ws["D2"].font.b
    assert wb.active.title == "Feuille1" and _logs(wb)[-1] == "Visuel regenere"


# ------------------------------------------------------------------ pipeline

class Espions:
    """Fausses fonctions d'entrée : enregistrent leurs appels et vérifient fichiers / chemins."""

    def __init__(self):
        self.appels = []

    def excel(self, chemin, ajuster_donnees, region):
        p = Path(chemin)
        assert p.is_absolute() and p.exists()                       # fichier enregistré AVANT la session Excel
        wb = openpyxl.load_workbook(p)                               # types attendus par Excel
        ws = wb["Feuille2"]
        n = ws.max_row
        assert ws.tables["tblDonnees"].ref == f"A1:E{n}"           # la plage s'arrête à la dernière ligne
        assert all(isinstance(ws.cell(row=r, column=1).value, datetime) for r in range(2, n + 1))
        # openpyxl relit 120.0 comme l'entier 120 : des nombres, jamais du texte
        assert all(isinstance(ws.cell(row=r, column=3).value, (int, float)) for r in range(2, n + 1))
        self.appels.append(("excel", p.name, ajuster_donnees, region))

    def mail(self, destinataire, cc, sujet, corps, pieces_jointes, mode):
        for pj in pieces_jointes:
            assert Path(pj).is_absolute() and Path(pj).exists()      # mail APRÈS l'enregistrement
        self.appels.append(("mail", destinataire, cc, sujet, corps, list(pieces_jointes), mode))


def test_pipeline_complet(tmp_path, copie, rng, maintenant, monkeypatch):
    monkeypatch.chdir(tmp_path)
    classeur = copie("classeur_complet.xlsx")
    avant = classeur.read_bytes()
    e = Espions()
    res = mi.lancer_pipeline_complet(classeur, "sortie", "a@b.fr", "c@d.fr", "afficher", "Nord",
                                     envoyer_mail=e.mail, excel=e.excel, rng=rng, maintenant=maintenant)
    assert res.ok
    nom = "VBA_Natran_Test_20261027_143005_VBATF.xlsx"
    assert res.chemins == [(tmp_path / "sortie" / nom).resolve()] and res.chemins[0].exists()
    # structure créée (True) => AutoFit des données ; région filtrée "Nord" ; Excel avant le mail
    assert [a[0] for a in e.appels] == ["excel", "mail"]
    assert e.appels[0] == ("excel", nom, True, "Nord")
    _, dest, cc, sujet, corps, pj, mode = e.appels[1]
    assert (dest, cc, mode) == ("a@b.fr", "c@d.fr", "afficher")
    assert sujet == "[PHASE DE TEST] Mise a jour suivi distribution - 27/10/2026"
    assert corps == ("Bonjour,\r\n\r\nCeci est un envoi automatique en phase de test.\r\n"
                     "Le graphique et les donnees ont ete mis a jour.\r\n"
                     f"Fichier : {nom}\r\n\r\nCordialement,\r\nRobot VBATF")
    assert pj == [str(res.chemins[0])]
    assert classeur.read_bytes() == avant                            # le classeur d'origine n'est pas modifié
    copie_wb = openpyxl.load_workbook(res.chemins[0])
    assert copie_wb["Feuille2"].tables["tblDonnees"].ref == "A1:E76"  # 70 historiques + 5 du jour
    assert copie_wb["Feuille1"]["B2"].value == "=SUM(tblDonnees[Volume_GWh])"


def test_pipeline_sans_filtre_ni_mail(tmp_path, copie, rng, maintenant):
    e = Espions()
    res = mi.lancer_pipeline_complet(copie("classeur_donnees.xlsx"), tmp_path / "s", mode_mail="aucun",
                                     saisie="", envoyer_mail=e.mail, excel=e.excel, rng=rng, maintenant=maintenant)
    assert res.ok
    assert e.appels == [("excel", "VBA_Natran_Test_20261027_143005_VBATF.xlsx", False, None)]  # pas de mail


def test_pipeline_mail_en_echec_ne_bloque_pas(tmp_path, copie, rng, maintenant, caplog):
    caplog.set_level(logging.INFO, logger="vba2py")

    def mail_ko(*a):
        raise OSError("Outlook indisponible")

    res = mi.lancer_pipeline_complet(copie("classeur_donnees.xlsx"), tmp_path / "s", envoyer_mail=mail_ko,
                                     excel=Espions().excel, rng=rng, maintenant=maintenant)
    assert res.ok                                                    # MailErr : consigné, la macro continue
    assert "Mail impossible : Outlook indisponible" in caplog.text


def test_pipeline_erreur_excel_est_signalee(tmp_path, copie, rng, maintenant, caplog):
    def excel_ko(*a):
        raise RuntimeError("Excel refuse")

    e = Espions()
    res = mi.lancer_pipeline_complet(copie("classeur_donnees.xlsx"), tmp_path / "s", envoyer_mail=e.mail,
                                     excel=excel_ko, rng=rng, maintenant=maintenant)
    assert not res.ok and "Excel refuse" in res.message
    assert e.appels == []                                            # ErrHandler : pas de mail
    assert "Excel refuse" in caplog.text


def test_pipeline_extension_conservee(tmp_path, copie, rng, maintenant):
    res = mi.lancer_pipeline_complet(copie("classeur_donnees.xlsx"), tmp_path / "s", mode_mail="aucun",
                                     excel=Espions().excel, rng=rng, maintenant=maintenant)
    assert res.chemins[0].suffix == ".xlsx"                          # ThisWorkbook.SaveCopyAs garde le format


def test_ligne_de_commande_declare_les_options(capsys):
    with pytest.raises(SystemExit):
        mi.main(["--help"])
    sortie = capsys.readouterr().out
    for option in ("--dossier-sortie", "--mail-to", "--mail-cc", "--mode-mail", "--saisie"):
        assert option in sortie
