"""Tests de utilisation_csv.py (fixture : tests/fixtures/csv/DAT_20261027.csv, 4 lignes valides, 4 invalides)."""

import logging
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

import commun
import utilisation_csv as uc
from commun import ErreurVBA

NOW = datetime(2026, 10, 27, 14, 30, 5)


def _logs(wb):
    return [wb["Log"].cell(row=r, column=c).value for r in range(2, wb["Log"].max_row + 1) for c in (2, 3)]


@pytest.fixture
def csv_fixture(donnees_fictives):
    return donnees_fictives / "csv" / "DAT_20261027.csv"


def test_nom_csv():
    assert uc.nom_csv(datetime(2026, 10, 5).date(), "csv") == Path("csv") / "DAT_20261005.csv"


def test_lire_lignes_csv(csv_fixture, tmp_path):
    lignes = uc.lire_lignes_csv(csv_fixture)
    assert len(lignes) == 10 and lignes[3] == "2026-10-27;Orléans;47.9;1.9;60.5;10.4"   # cp1252 relu correctement
    assert lignes[8] == ""
    mixte = tmp_path / "m.csv"
    mixte.write_bytes(b"a\r\nb\rc\nd")                      # CRLF, CR, LF, sans saut final
    assert uc.lire_lignes_csv(mixte) == ["a", "b", "c", "d"]


@pytest.mark.parametrize("texte, attendu", [("12.5", True), ("12,5", False), ("", False), ("abc", False),
                                            ("-0.37", True), ("  7 ", True)])
def test_est_nombre(texte, attendu):
    assert uc.est_nombre(texte) is attendu


@pytest.mark.parametrize("ligne, attendu", [
    ("2026-10-27;Lille;50.63;3.06;98.4;9.1", True),
    ("2026-10-27;Lille;50.63;3.06;98.4", False),             # 5 champs
    ("2026-10-27;;50.63;3.06;98.4;9.1", False),              # ville vide
    ("27/10/2026;Caen;49.18;-0.37;55;7", False),             # date mal formée
    ("2026-10-27;Nice;43,7;7.27;70;12", False),              # virgule décimale
    ("2026-10-27;Madrid;40.4;-3.7;200;15", False),           # latitude < 41
    ("2026-10-27;Brest;48.4;-5.6;80;12", False),             # longitude < -5.5
    ("2026-10-27;Metz;49.1;10.5;80;12", False),              # longitude > 10
    ("2026-10-27;Metz;52;10;80;12", True),                   # bornes incluses
    ("2026-10-27;X;abc;;xx;", False),
])
def test_ligne_valide(ligne, attendu):
    assert uc.ligne_valide(ligne.split(";")) is attendu


def test_importer_csv(copie, csv_fixture):
    wb = openpyxl.load_workbook(copie("classeur_complet.xlsx"))
    assert uc.importer_csv(wb, csv_fixture, NOW) == 4
    ws = wb["Donnees_V3"]
    assert ws.tables["tblCSV_V3"].ref == "A1:F5"
    assert [c.value for c in ws[1]] == ["Date", "Ville", "Latitude", "Longitude", "Volume_GWh", "Temperature_C"]
    assert [ws.cell(row=r, column=2).value for r in range(2, 6)] == ["Lille", "Paris", "Orléans", "Lyon"]
    assert ws["A2"].value == datetime(2026, 10, 27) and ws["A2"].number_format == "dd/mm/yyyy"
    assert (ws["C2"].value, ws["D2"].value, ws["E2"].value, ws["F2"].value) == (50.63, 3.06, 98.4, 9.1)
    assert ws["E5"].value == 112.0 and isinstance(ws["E5"].value, float)
    journal = _logs(wb)
    assert journal.count("AVERTISSEMENT") == 4              # VilleInconnue, Madrid, Nice, Caen (la ligne vide est ignorée)
    assert any(m.startswith("Ligne ignoree : 2026-10-27;Madrid") for m in journal)
    assert any(m.startswith("CSV importe : 4 valide(s), 4 ignoree(s)") for m in journal)


def test_importer_csv_remplace_les_donnees_existantes(copie, csv_fixture):
    wb = openpyxl.load_workbook(copie("classeur_complet.xlsx"))
    uc.importer_csv(wb, csv_fixture, NOW)
    uc.importer_csv(wb, csv_fixture, NOW)                   # relance : ListObjects(1).Delete + Cells.Clear
    assert list(wb["Donnees_V3"].tables) == ["tblCSV_V3"] and wb["Donnees_V3"].tables["tblCSV_V3"].ref == "A1:F5"


def test_importer_csv_sans_ligne_valide(copie, tmp_path):
    csv = tmp_path / "vide.csv"
    csv.write_text("Date;Ville;Latitude;Longitude;Volume_GWh;Temperature_C\r\nn'importe quoi\r\n", encoding="cp1252")
    wb = openpyxl.load_workbook(copie("classeur_complet.xlsx"))
    with pytest.raises(ErreurVBA) as e:
        uc.importer_csv(wb, csv, NOW)
    assert e.value.numero == commun.VB_OBJECT_ERROR + 11 and e.value.description == f"Aucune ligne valide dans {csv}"


@pytest.fixture
def wb_importe(copie, csv_fixture):
    wb = openpyxl.load_workbook(copie("classeur_complet.xlsx"))
    uc.importer_csv(wb, csv_fixture, NOW)
    return wb


@pytest.mark.parametrize("saisie, attendu", [("", ""), ("paris", "Paris"), ("  LYON ", "Lyon"), ("Rome", "")])
def test_demander_ville(wb_importe, saisie, attendu):
    assert uc.demander_ville(wb_importe, saisie, NOW) == attendu


def test_generer_carte_html(wb_importe, tmp_path):
    chemin = uc.generer_carte_html(wb_importe, "Paris", tmp_path, NOW)
    assert chemin.name == "VBA_Natran_Test_Carte_20261027_143005_VBATF_V3.html"
    html = chemin.read_bytes().decode("cp1252")
    assert "<b>Distribution gaz - donnees du 27/10/2026</b>" in html
    assert "var DATA=[['Lille',50.63,3.06,98.4,9.1],['Paris',48.86,2.35,151.2,11.3],['Orléans',47.9,1.9,60.5,10.4]," in html
    assert "['Lyon',45.76,4.84,112,8.8]];" in html
    assert "selV.value='Paris';" in html and "- VBATF_V3</small>" in html


def test_carte_ville_inconnue_et_apostrophe(wb_importe, tmp_path):
    wb_importe["Donnees_V3"]["B3"] = "Villers-l'Evêque"
    html = uc.generer_carte_html(wb_importe, "", tmp_path, NOW).read_text(encoding="cp1252")
    assert "['Villers-l\\'Evêque',48.86" in html and "selV.value='*';" in html


def test_creer_synthese(wb_importe):
    uc.creer_synthese(wb_importe, "C:/c.html", "C:/d.csv", NOW)
    ws = wb_importe["Carte_V3"]
    assert ws["D2"].value == "Moyenne / ville (GWh)" and ws["E2"].value == "=AVERAGE(tblCSV_V3[Volume_GWh])"
    assert ws["A5"].value == "=INDEX(tblCSV_V3[Ville],1)" and ws["C8"].value == "=INDEX(tblCSV_V3[Temperature_C],4)"
    assert ws["A9"].value is None                           # 4 villes : lignes 5 à 8
    assert ws["A10"].value == "Ouvrir la carte interactive" and ws["A10"].hyperlink.target == "C:/c.html"
    assert ws["A11"].value == "Source CSV : C:/d.csv" and ws["A12"].value == "Carte : C:/c.html"
    assert [str(r.sqref) for r in ws.conditional_formatting] == ["B5:B8"]


class Espions:
    def __init__(self):
        self.appels = []

    def excel(self, chemin):
        assert Path(chemin).is_absolute() and Path(chemin).exists()
        self.appels.append(("excel", Path(chemin).name))

    def mail(self, dest, cc, sujet, corps, pj, mode):
        assert all(Path(p).is_absolute() and Path(p).exists() for p in pj)
        self.appels.append(("mail", dest, sujet, corps, [Path(p).name for p in pj], mode))

    def ouvrir(self, chemin):
        self.appels.append(("ouvrir", Path(chemin).name))


def test_pipeline_csv(tmp_path, copie, csv_fixture, maintenant, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "csv").mkdir()
    shutil.copy(csv_fixture, tmp_path / "csv" / "DAT_20261027.csv")
    e = Espions()
    res = uc.lancer_pipeline_csv_v3(copie("classeur_complet.xlsx"), "sortie", "csv", mode_mail="afficher",
                                    saisie="Lyon", envoyer_mail=e.mail, excel=e.excel, ouvrir_fichier=e.ouvrir,
                                    maintenant=maintenant)
    assert res.ok and "4 ville(s) importee(s)" in res.message
    xlsm, carte = res.chemins
    assert xlsm.name == "VBA_Natran_Test_20261027_143005_VBATF_V3.xlsx"
    assert [a[0] for a in e.appels] == ["excel", "mail", "ouvrir"]
    _, dest, sujet, corps, pj, mode = e.appels[1]
    assert dest == "tom.ferrandis@hotmail.com" and mode == "afficher"       # destinataire d'origine conservé
    assert pj == [xlsm.name, carte.name, "DAT_20261027.csv"]
    assert "Les donnees ont ete importees depuis : DAT_20261027.csv\r\n" in corps
    assert sujet == "[PHASE DE TEST V3] Mise a jour suivi distribution - 27/10/2026"
    wb = openpyxl.load_workbook(xlsm)
    assert wb["Donnees_V3"].tables["tblCSV_V3"].ref == "A1:F5"              # plage = dernière ligne remplie
    assert all(isinstance(wb["Donnees_V3"].cell(row=r, column=1).value, datetime) for r in range(2, 6))
    assert "selV.value='Lyon';" in carte.read_text(encoding="cp1252")


def test_pipeline_csv_du_jour_introuvable(tmp_path, copie, maintenant, caplog):
    caplog.set_level(logging.INFO, logger="vba2py")
    e = Espions()
    res = uc.lancer_pipeline_csv_v3(copie("classeur_complet.xlsx"), tmp_path / "s", tmp_path / "vide",
                                    envoyer_mail=e.mail, excel=e.excel, maintenant=maintenant)
    assert not res.ok and "CSV du jour introuvable" in res.message
    assert "Genere d'abord le CSV avec le module de creation." in caplog.text
    assert e.appels == [] and not (tmp_path / "s").exists()


def test_ligne_de_commande_declare_les_options(capsys):
    with pytest.raises(SystemExit):
        uc.main(["--help"])
    sortie = capsys.readouterr().out
    for option in ("--dossier-sortie", "--csv-folder", "--mail-to", "--mail-cc", "--mode-mail", "--open-map",
                   "--saisie"):
        assert option in sortie
