"""Tests des fonctions qui reproduisent le comportement de VBA (Val, IsNumeric, Str$, CDate...)."""

from datetime import date, datetime

import pytest

import commun
from commun import ErreurVBA


@pytest.mark.parametrize("texte, attendu", [
    ("12,5", 12.0),          # Val lit le point, pas la virgule : s'arrête à la virgule
    ("12.5", 12.5),
    ("1 234", 1234.0),       # les espaces sont ignorés
    ("abc", 0.0),
    ("", 0.0),
    ("-3.5kg", -3.5),
    ("1e3", 1000.0),
])
def test_val(texte, attendu):
    assert commun.vba_val(texte) == attendu


@pytest.mark.parametrize("texte, attendu", [
    ("12,5", True), ("12.5", False), ("2026", True), ("", False), ("abc", False), ("-4", True), ("&H1F", True),
    ("1,2,3", False),
])
def test_is_numeric_reglages_francais(texte, attendu):
    assert commun.is_numeric(texte) is attendu


def test_str_nb():
    assert commun.str_nb(50.63) == "50.63"
    assert commun.str_nb(95) == "95"          # Str$(95#) = " 95" puis Trim$
    assert commun.str_nb(-1.55) == "-1.55"
    assert commun.str_nb(0.00001) == "1E-05"


def test_date_serial_deborde_comme_vba():
    assert commun.date_serial(2026, 13, 1) == datetime(2027, 1, 1)   # mois 13 = janvier suivant
    assert commun.date_serial(2026, 2, 30) == datetime(2026, 3, 2)
    assert commun.date_serial(26, 10, 5) == datetime(2026, 10, 5)    # année sur 2 chiffres


def test_c_date_jour_mois_annee():
    assert commun.c_date("05/10/2026") == date(2026, 10, 5)          # le 5 octobre, pas le 10 mai
    assert commun.c_date("2026-10-05") == date(2026, 10, 5)
    assert commun.is_date("31/12/2026") and not commun.is_date("31/02/2026") and not commun.is_date("hier")
    with pytest.raises(ErreurVBA) as e:
        commun.c_date("hier")
    assert e.value.numero == 13


def test_cstr_francais():
    assert commun.cstr(None) == ""
    assert commun.cstr(12.5) == "12,5"
    assert commun.cstr(12.0) == "12"
    assert commun.cstr(datetime(2026, 10, 27)) == "27/10/2026"


def test_derniere_ligne_et_journal(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    assert commun.derniere_ligne(ws) == 1                              # colonne vide : End(xlUp) donne 1
    commun.log_msg(wb, "INFO", "un", datetime(2026, 10, 27, 8, 0, 0))
    commun.log_msg(wb, "ERREUR", "deux", datetime(2026, 10, 27, 8, 0, 1))
    log = wb["Log"]
    assert [c.value for c in log[1]] == ["Horodatage", "Niveau", "Message"]
    assert [log.cell(row=3, column=i).value for i in (2, 3)] == ["ERREUR", "deux"]
    assert log["A2"].number_format == "dd/mm/yyyy hh:mm:ss"


def test_table_helpers_et_erreur_9():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, t in enumerate(("a", "b"), 1):
        ws.cell(row=1, column=i, value=t)
    ws.append([1, 2])
    t = commun.creer_table(ws, "tbl", "A1:B2")
    assert commun.ajouter_ligne_table(ws, t, [3, 4]) == 3
    assert t.ref == "A1:B3" and t.autoFilter.ref == "A1:B3"
    assert commun.lignes_table(ws, t) == [[1, 2], [3, 4]]
    with pytest.raises(ErreurVBA) as e:
        commun.table_de(ws, "absent")
    assert e.value.numero == 9
    with pytest.raises(ErreurVBA):
        commun.feuille(wb, "absente")
