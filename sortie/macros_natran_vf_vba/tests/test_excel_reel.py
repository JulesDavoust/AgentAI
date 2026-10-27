"""Tests sur Excel réel : chaque fonction d'entrée de connecteurs/excel.py est exécutée par Excel, puis le
fichier est relu avec openpyxl. Ignorés si Excel n'est pas disponible."""

import shutil
from datetime import datetime

import openpyxl
import pytest

import macro_carte as mc
import macro_initiale as mi
import utilisation_csv as uc
from bibliotheque.excel_com import excel_disponible
from connecteurs import excel as ce

pytestmark = [pytest.mark.excel_reel,
              pytest.mark.skipif(not excel_disponible(), reason="Excel indisponible")]

NOW = datetime(2026, 10, 27, 14, 30, 5)


def _preparer_donnees(copie, tmp_path, ecrire):
    """Reproduit le pipeline : étapes openpyxl puis enregistrement, avant la session Excel."""
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    ecrire(wb)
    chemin = tmp_path / "pipeline.xlsx"
    wb.save(chemin)
    # types vérifiés juste avant Excel : vraies dates, plage du tableau exacte
    ws = openpyxl.load_workbook(chemin)["Feuille2"]
    assert ws.tables["tblDonnees"].ref == f"A1:E{ws.max_row}"
    assert all(isinstance(ws.cell(row=r, column=1).value, datetime) for r in range(2, ws.max_row + 1))
    return chemin


def test_finaliser_initiale_cree_tableau_croise_et_graphique(copie, tmp_path):
    chemin = _preparer_donnees(copie, tmp_path, lambda wb: mi.creer_visuel_cellules(wb, NOW))
    ce.finaliser_initiale(chemin, False, "Nord")
    wb = openpyxl.load_workbook(chemin)
    ws = wb["Feuille1"]
    assert [p.name for p in ws._pivots] == ["TCD_Volume"]
    assert ws._pivots[0].location.ref.startswith("A5")
    assert len(ws._charts) == 1 and type(ws._charts[0]).__name__ == "LineChart"
    assert ws._charts[0].anchor._from.col == 9 and ws._charts[0].anchor._from.row == 4     # J5
    filtre = wb["Feuille2"].tables["tblDonnees"].autoFilter.filterColumn
    assert filtre and filtre[0].colId == 1 and filtre[0].filters.filter == ["Nord"]


def test_finaliser_carte_filtre_et_largeur(copie, tmp_path):
    chemin = _preparer_donnees(copie, tmp_path, lambda wb: mc.creer_synthese(wb, "c.html", NOW))
    ce.finaliser_carte(chemin, True, "Sud")
    wb = openpyxl.load_workbook(chemin)
    # ColumnWidth = 24 : Excel l'enregistre dans le fichier avec son remplissage (24 + 5/7)
    assert wb["Feuille1"].column_dimensions["A"].width == pytest.approx(24.71, abs=0.05)
    assert wb["Feuille2"].tables["tblDonnees"].autoFilter.filterColumn[0].filters.filter == ["Sud"]
    cachees = [r for r in range(2, 14) if wb["Feuille2"].row_dimensions[r].hidden]
    assert len(cachees) == 12 - 3                       # 3 lignes « Sud » restent visibles (25, 26 et 27/10)


def test_finaliser_csv_v3_largeurs(copie, tmp_path, donnees_fictives):
    wb = openpyxl.load_workbook(copie("classeur_complet.xlsx"))
    csv = donnees_fictives / "csv" / "DAT_20261027.csv"
    uc.importer_csv(wb, csv, NOW)
    uc.creer_synthese(wb, "c.html", csv, NOW)
    chemin = tmp_path / "v3.xlsx"
    wb.save(chemin)
    ce.finaliser_csv_v3(chemin)
    wb = openpyxl.load_workbook(chemin)
    assert wb["Carte_V3"].column_dimensions["A"].width == pytest.approx(24.71, abs=0.05)
    ville = wb["Donnees_V3"].column_dimensions["B"]                    # « Orléans » : 7 caractères
    assert ville.bestFit and 7 < ville.width < 9                       # largeur calculée par Excel (AutoFit)


def test_ajuster_feuille_mvs(copie, tmp_path):
    chemin = tmp_path / "mvs.xlsx"
    shutil.copy(copie("classeur_mvs.xlsx"), chemin)
    ce.ajuster_feuille(chemin, "MVS_Base", "A:F")
    wb = openpyxl.load_workbook(chemin)
    assert wb["MVS_Base"].column_dimensions["B"].width > 15            # libellés de fonds longs
    assert wb["MVS_Base"].sheet_state == "hidden"
