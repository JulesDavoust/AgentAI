"""Crée les classeurs et CSV fictifs de tests/fixtures/ (idempotent). À lancer sans argument.

Valeurs inventées, de même forme que le vrai classeur :
- classeur_complet.xlsx : Feuil1 (feuille des boutons) + MVS_Base masquée (en-têtes ISIN... et 6 fonds) ;
- classeur_donnees.xlsx : classeur_complet + Feuille1 vide + Feuille2 avec le tableau tblDonnees
  (25 et 26/10/2026 : 5 régions ; 27/10/2026 : Nord et Sud seulement) ;
- classeur_mvs.xlsx : classeur_complet avec l'écran MVS déjà dessiné ;
- csv/DAT_20261027.csv : 4 lignes valides et 4 lignes invalides (Windows-1252, fins de ligne Windows).
"""

from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

RACINE = Path(__file__).resolve().parent.parent
FIXTURES = Path(__file__).resolve().parent / "fixtures"
if str(RACINE) not in sys.path:
    sys.path.insert(0, str(RACINE))

REGIONS = ("Nord", "Sud", "Est", "Ouest", "Centre")
BASES = {"Nord": 120, "Sud": 95, "Est": 80, "Ouest": 70, "Centre": 60}

BASE_MVS = (
    ("FR0010000001", "FONDS ACTIONS EUROPE", "ACTIF", "1250.5", "01/10/2026", ""),
    ("FR0010000002", "FONDS OBLIG CORPORATE", "ACTIF", "890.2", "28/09/2026", "RAS"),
    ("FR0010000003", "FONDS MONETAIRE EURO", "ACTIF", "2310.0", "15/10/2026", ""),
    ("FR0010000004", "FONDS DIVERSIFIE PRUDENT", "ACTIF", "455.8", "03/10/2026", "A SURVEILLER"),
    ("FR0010000005", "FONDS ACTIONS USA", "FERME", "120.4", "30/06/2026", "LIQUIDATION"),
    ("FR0010000006", "FONDS INFRA ENERGIE", "ACTIF", "780.0", "20/10/2026", ""),
)

CSV_DU_JOUR = [
    "Date;Ville;Latitude;Longitude;Volume_GWh;Temperature_C",
    "2026-10-27;Lille;50.63;3.06;98.4;9.1",
    "2026-10-27;Paris;48.86;2.35;151.2;11.3",
    "2026-10-27;Orléans;47.9;1.9;60.5;10.4",
    "2026-10-27;Lyon;45.76;4.84;112;8.8",
    "2026-10-27;VilleInconnue;abc;;xx;",          # invalide : 6 champs mais non numériques
    "2026-10-27;Madrid;40.4;-3.7;200;15",          # invalide : hors de France (latitude < 41)
    "2026-10-27;Nice;43,7;7.27;70;12.0",           # invalide : virgule décimale
    "",                                            # ligne vide : ignorée sans avertissement
    "27/10/2026;Caen;49.18;-0.37;55;7",            # invalide : date au mauvais format
]


def _classeur_complet() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Feuil1"
    ws["A1"] = "Boutons : voir EXTRACTION.md"
    base = wb.create_sheet("MVS_Base")
    for i, t in enumerate(("ISIN", "Libelle", "Statut", "Encours_MEUR", "Date_MAJ", "Commentaire"), 1):
        base.cell(row=1, column=i, value=t).number_format = "@"
    for r, ligne in enumerate(BASE_MVS, 2):
        for c, v in enumerate(ligne, 1):
            cell = base.cell(row=r, column=c, value=v if v != "" else None)
            cell.number_format = "@"
    base.sheet_state = "hidden"
    return wb


def generer() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    (FIXTURES / "csv").mkdir(exist_ok=True)

    wb = _classeur_complet()
    wb.save(FIXTURES / "classeur_complet.xlsx")

    wb = _classeur_complet()
    wb.create_sheet("Feuille1")
    ws = wb.create_sheet("Feuille2")
    for i, t in enumerate(("Date", "Region", "Volume_GWh", "Temperature_C", "Statut"), 1):
        ws.cell(row=1, column=i, value=t)
    ligne = 2
    for jour_idx, jour in enumerate((25, 26, 27)):
        for reg_idx, reg in enumerate(REGIONS):
            if jour == 27 and reg not in ("Nord", "Sud"):
                continue
            ws.cell(row=ligne, column=1, value=datetime(2026, 10, jour)).number_format = "dd/mm/yyyy"
            ws.cell(row=ligne, column=2, value=reg)
            ws.cell(row=ligne, column=3, value=BASES[reg] + jour_idx * 2.5)
            ws.cell(row=ligne, column=4, value=6.0 + jour_idx * 1.5 + reg_idx * 0.5)
            ws.cell(row=ligne, column=5, value="Historique")
            ligne += 1
    table = Table(displayName="tblDonnees", name="tblDonnees", ref=f"A1:E{ligne - 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    wb.save(FIXTURES / "classeur_donnees.xlsx")

    from simu_mainframe import creer_ecran_mvs  # écran MVS dessiné par le code traduit
    wb = _classeur_complet()
    creer_ecran_mvs(wb, date(2026, 10, 27))
    wb.save(FIXTURES / "classeur_mvs.xlsx")

    with open(FIXTURES / "csv" / "DAT_20261027.csv", "w", encoding="cp1252", newline="\r\n") as f:
        for ligne in CSV_DU_JOUR:
            f.write(ligne + "\n")


if __name__ == "__main__":
    generer()
    print(f"Données fictives créées dans {FIXTURES}")
