"""Tests de test_interface_mvs.py (API terminal et démonstrations sur l'écran MVS simulé)."""

import logging
from datetime import date, datetime

import openpyxl
import pytest

import commun
import test_interface_mvs as ti
from commun import ErreurVBA

JOUR = lambda: date(2026, 10, 27)  # noqa: E731
HEURE = lambda: datetime(2026, 10, 27, 14, 30, 5)  # noqa: E731


@pytest.fixture
def terminal(copie):
    t = ti.TerminalMVS(openpyxl.load_workbook(copie("classeur_mvs.xlsx")), JOUR)
    assert t.init_mvs()
    return t


def test_init_mvs(copie):
    sans_ecran = ti.TerminalMVS(openpyxl.load_workbook(copie("classeur_complet.xlsx")), JOUR)
    assert sans_ecran.init_mvs() is False and sans_ecran.modele_session() == 0
    avec = ti.TerminalMVS(openpyxl.load_workbook(copie("classeur_mvs.xlsx")), JOUR)
    assert avec.init_mvs() is True and avec.modele_session() == 3 and (avec.cur_ligne, avec.cur_col) == (1, 1)


def test_ecrire_puis_lire(terminal):
    terminal.ecrire(6, 27, "ABC")
    assert terminal.lire(6, 27, 5) == "ABC  "                      # complété par des espaces jusqu'à la taille
    assert (terminal.cur_ligne, terminal.cur_col) == (6, 30)       # curseur après le texte
    terminal.ecrire(6, 28, " ")                                    # un espace efface la cellule
    assert terminal.lire(6, 27, 3) == "A C"


def test_ecrire_au_dela_de_la_colonne_80(terminal):
    terminal.ecrire(3, 79, "XYZ")
    assert terminal.lire(3, 79, 2) == "XY" and terminal.cur_col == 80


@pytest.mark.parametrize("l, c", [(0, 1), (33, 1), (1, 0), (1, 81)])
def test_coordonnees_hors_ecran(terminal, l, c):
    with pytest.raises(ErreurVBA) as e:
        terminal.lire(l, c, 1)
    assert e.value.numero == commun.VB_OBJECT_ERROR + 20
    assert e.value.description == f"Coordonnees hors ecran : ({l},{c})"


def test_touches(terminal):
    terminal.se_positionner(6, 27)
    terminal.touche("fr0010000001")                               # un texte est saisi à la position du curseur
    terminal.touche("@E")                                         # ENTREE
    assert terminal.lire(8, 27, 20).strip() == "FONDS ACTIONS EUROPE"   # ISIN mis en majuscules puis trouvé
    terminal.touche("@C")                                         # CLEAR
    assert terminal.lire(6, 27, 12).strip() == "" and terminal.lire(30, 13, 12) == "ECRAN EFFACE"
    terminal.touche("@3")
    assert terminal.lire(30, 13, 13) == "PF3 - SESSION"
    terminal.touche("@T")
    terminal.touche("")                                           # sans effet
    terminal.touche("@Z")                                         # touche non gérée : journal debug seulement


def test_touche_f_efface_la_fin_de_ligne(terminal):
    terminal.ecrire(16, 27, "ABCDEF")
    terminal.se_positionner(16, 29)
    terminal.touche("@f")                                         # UCase$ : insensible à la casse
    assert terminal.lire(16, 27, 6) == "AB    "


def test_touche_macro_introuvable_ou_en_erreur(terminal, caplog):
    terminal.appeler_mvs("MVS_Inconnue")
    assert "Macro 'MVS_Inconnue' introuvable ou en erreur" in caplog.text
    del terminal.wb["MVS_Base"]                                   # MVS_Effacer échoue : erreur signalée, pas levée
    terminal.touche("@C")
    assert caplog.text.count("introuvable ou en erreur") == 2


def test_recherche_insensible_a_la_casse(terminal):
    assert terminal.recherche("code isin") and terminal.recherche("PF3=Quitter")
    assert not terminal.recherche("INCONNU")


def test_consulter_isin(terminal):
    terminal.consulter_isin("FR0010000004")
    assert terminal.lire(ti.L_LIB, ti.C_CHAMP, ti.LEN_LIB).strip() == "FONDS DIVERSIFIE PRUDENT"
    assert terminal.lire(ti.L_COM, ti.C_CHAMP, ti.LEN_COM).strip() == "A SURVEILLER"
    terminal.consulter_isin("FR0000000000")
    assert terminal.recherche("INCONNU")


@pytest.mark.parametrize("saisie, attendu", [("5", 5), (" 32 ", 32), ("1,0", 1), ("", None), (None, None)])
def test_demander_entier_valide_ou_annule(saisie, attendu):
    assert ti.demander_entier(saisie, 1, 32) == attendu


@pytest.mark.parametrize("saisie", ["0", "33", "2,5", "abc"])
def test_demander_entier_invalide(saisie):
    with pytest.raises(ValueError, match="entier entre 1 et 32"):
        ti.demander_entier(saisie, 1, 32)


def test_interagir_lecture(copie):
    classeur = copie("classeur_mvs.xlsx")
    # le titre commence en colonne 2 : « NATRAN - ... »
    res = ti.interagir_mainframe(classeur, "1", "2", "5", None, date_du_jour=JOUR)
    assert res.ok and res.message == "Ligne 1, colonne 2, longueur 5\n\nContenu : [NATRA]"
    ws = openpyxl.load_workbook(classeur)["MVS"]
    assert ws.sheet_view.selection[0].sqref == "B1:F1"             # zone lue sélectionnée


def test_interagir_ecriture_valide_avec_entree(copie):
    classeur = copie("classeur_mvs.xlsx")
    res = ti.interagir_mainframe(classeur, "6", "27", "10", "fr0010000003", date_du_jour=JOUR)
    assert res.ok and res.message == "Reponse du mainframe :\nFSIC - ENREGISTREMENT TROUVE"
    wb = openpyxl.load_workbook(classeur)                          # état enregistré en place
    assert wb["MVS_Base"]["H1"].value == "FR0010000003"


def test_interagir_texte_tronque_en_fin_de_ligne(copie, caplog):
    res = ti.interagir_mainframe(copie("classeur_mvs.xlsx"), "3", "79", None, "ABCDE", date_du_jour=JOUR)
    assert res.ok and "Texte tronque a 2 caracteres" in caplog.text


@pytest.mark.parametrize("args, ok, message", [
    (("", "1"), True, ""),                                          # Annuler : la macro s'arrête
    (("1", ""), True, ""),
    (("40", "1"), False, "Valeur invalide : saisir un entier entre 1 et 32."),
    (("1", "81"), False, "Valeur invalide : saisir un entier entre 1 et 80."),
])
def test_interagir_saisies(copie, args, ok, message):
    res = ti.interagir_mainframe(copie("classeur_mvs.xlsx"), *args, date_du_jour=JOUR)
    assert (res.ok, res.message) == (ok, message)


def test_interagir_mainframe_non_ouvert(copie):
    res = ti.interagir_mainframe(copie("classeur_complet.xlsx"), "1", "1", date_du_jour=JOUR)
    assert not res.ok and res.message.startswith("Le mainframe n'est pas encore ouvert.")


def test_demo_ecrire_et_lire(copie):
    classeur = copie("classeur_mvs.xlsx")
    res = ti.demo_ecrire_et_lire(classeur, date_du_jour=JOUR, maintenant=HEURE)
    assert res.ok
    assert res.message.startswith("Lecture mainframe :\nFONDS OBLIG CORPORATE\nStatut : ACTIF\nEncours (MEUR) : 890.2")
    assert res.message.endswith("Reponse du mainframe : ENREGISTREMENT MIS A JOUR")
    base = openpyxl.load_workbook(classeur)["MVS_Base"]
    assert base["F3"].value == "MAJ PAR VBA 14:30:05" and base["E3"].value == "27/10/2026"


def test_demo_lire_saisie_manuelle(copie):
    res = ti.demo_lire_saisie_manuelle(copie("classeur_mvs.xlsx"), date_du_jour=JOUR)
    assert res.message == "ISIN saisi : \nCommentaire : \nMessage : PRET - SAISIR UN CODE ISIN PUIS ENTREE"


@pytest.mark.parametrize("isin, attendu", [
    ("fr0010000003", "FR0010000003 existe dans FSIC :\nFONDS MONETAIRE EURO"),
    ("FR0000000000", "FR0000000000 n'existe pas dans FSIC."),
])
def test_demo_verifier_existence(copie, isin, attendu):
    assert ti.demo_verifier_existence_fsic(copie("classeur_mvs.xlsx"), isin, date_du_jour=JOUR).message == attendu


def test_demo_verifier_existence_saisie_vide_ou_sans_ecran(copie):
    assert ti.demo_verifier_existence_fsic(copie("classeur_mvs.xlsx"), "  ").message == ""
    assert not ti.demo_verifier_existence_fsic(copie("classeur_complet.xlsx"), "FR0010000003").ok


def test_demo_extraction_vers_feuille(copie):
    classeur = copie("classeur_mvs.xlsx")
    appels = []
    res = ti.demo_extraction_vers_feuille(classeur, date_du_jour=JOUR, excel=lambda *a: appels.append(a))
    assert res.message == "7 ISIN interroges sur le mainframe."
    assert appels == [(classeur, "Extraction_MVS", "A:F", True)]    # AutoFit + Activate par Excel
    wb = openpyxl.load_workbook(classeur)
    ws = wb["Extraction_MVS"]
    assert [c.value for c in ws[1]] == ["ISIN", "Libelle", "Statut", "Encours_MEUR", "Date_MAJ", "Resultat"]
    assert ws["A1"].font.b
    assert [ws.cell(row=2, column=c).value for c in range(1, 7)] == [
        "FR0010000001", "FONDS ACTIONS EUROPE", "ACTIF", 1250.5, "01/10/2026", "OK"]     # Val("1250.5") : nombre
    assert ws["D4"].value == 2310.0                                  # Val("2310.0") : nombre, pas texte
    assert [ws.cell(row=8, column=c).value for c in range(1, 7)] == ["FR0000000000", None, None, None, None, "INCONNU"]
    assert wb.active.title == "Extraction_MVS"


def test_ligne_de_commande_declare_les_options(capsys):
    with pytest.raises(SystemExit):
        ti.main(["interagir-mainframe", "--help"])
    sortie = capsys.readouterr().out
    for option in ("--ligne", "--colonne", "--longueur", "--texte"):
        assert option in sortie
    with pytest.raises(SystemExit):
        ti.main(["demo-verifier-existence-fsic", "--help"])
    assert "--isin" in capsys.readouterr().out
