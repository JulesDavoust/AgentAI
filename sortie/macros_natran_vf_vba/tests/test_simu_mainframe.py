"""Tests de simu_mainframe.py."""

from datetime import date, datetime

import openpyxl
import pytest

import simu_mainframe as sm

AUJOURDHUI = date(2026, 10, 27)
MAINTENANT = lambda: datetime(2026, 10, 27, 14, 30, 5)  # noqa: E731


@pytest.fixture
def wb(copie):
    return openpyxl.load_workbook(copie("classeur_mvs.xlsx"))


def saisir(wb, isin=None, commentaire=None):
    ws = wb["MVS"]
    if isin is not None:
        sm.ecrire_champ(ws, sm.L_ISIN, sm.C_CHAMP, sm.LEN_ISIN, isin)
    if commentaire is not None:
        sm.ecrire_champ(ws, sm.L_COM, sm.C_CHAMP, sm.LEN_COM, commentaire)


def champ(wb, l, c, n):
    return sm.lire_champ(wb["MVS"], l, c, n).strip()


def test_creer_ecran_dessine_l_ecran(copie):
    wb = openpyxl.load_workbook(copie("classeur_complet.xlsx"))
    assert sm.creer_ecran_mvs(wb, AUJOURDHUI) is False                # la base existait
    ws = wb["MVS"]
    assert sm.buffer_ligne(ws, 1).startswith(" NATRAN - SIMULATION MAINFRAME MVS (3278 MODELE 3 - 32x80)")
    assert sm.buffer_ligne(ws, 1)[61:76] == "DATE 27/10/2026"        # colonne 62
    assert sm.buffer_ligne(ws, 2) == "-" * 80
    assert sm.buffer_ligne(ws, 6)[2:23] == "CODE ISIN . . . . . :"
    assert champ(wb, sm.L_MSG, sm.C_MSG, sm.LEN_MSG) == "PRET - SAISIR UN CODE ISIN PUIS ENTREE"
    assert ws["AA6"].fill.start_color.rgb.endswith(sm.C_SAISIE)        # zone de saisie de l'ISIN
    assert ws["A1"].fill.start_color.rgb.endswith(sm.C_FOND) and ws["A1"].font.name == "Courier New"
    assert ws["B1"].number_format == "@" and ws.row_dimensions[3].height == 15
    assert ws.sheet_view.showGridLines is False and ws.sheet_view.zoomScale == 100
    assert ws.sheet_view.selection[0].activeCell == "AA6" and wb.active.title == "MVS"
    assert sm.buffer_ligne(ws, 35) == "1234567890" * 8                # règle de coordonnées


def test_creer_ecran_recree_la_base_absente():
    wb = openpyxl.Workbook()
    assert sm.creer_ecran_mvs(wb, AUJOURDHUI) is True
    base = wb["MVS_Base"]
    assert base.sheet_state == "hidden"
    assert [c.value for c in base[1][:6]] == ["ISIN", "Libelle", "Statut", "Encours_MEUR", "Date_MAJ", "Commentaire"]
    assert base["A7"].value == "FR0010000006" and base["D2"].value == "1250.5" and base["F2"].value is None
    assert base["D2"].number_format == "@" and isinstance(base["D2"].value, str)      # tout est du texte


def test_ouvrir_mainframe_cree_puis_affiche(copie):
    classeur = copie("classeur_complet.xlsx")
    appels = []
    assert sm.ouvrir_mainframe(classeur, excel=lambda *a: appels.append(a), maintenant=MAINTENANT) is True
    assert "MVS" in openpyxl.load_workbook(classeur).sheetnames
    assert appels == []                                                # base déjà présente : pas d'AutoFit
    wb = openpyxl.load_workbook(classeur)
    wb["MVS"].sheet_state = "hidden"
    wb["MVS"].sheet_view.selection[0].activeCell = "A1"
    wb.save(classeur)
    assert sm.ouvrir_mainframe(classeur, excel=lambda *a: appels.append(a)) is False      # existe : on l'affiche
    wb = openpyxl.load_workbook(classeur)
    assert wb["MVS"].sheet_state == "visible" and wb.active.title == "MVS"
    assert wb["MVS"].sheet_view.selection[0].activeCell == "AA6"


def test_ouvrir_mainframe_base_recreee_declenche_autofit(tmp_path):
    classeur = tmp_path / "vide.xlsx"
    openpyxl.Workbook().save(classeur)
    appels = []
    sm.ouvrir_mainframe(classeur, excel=lambda *a: appels.append(a), maintenant=MAINTENANT)
    assert appels == [(classeur, "MVS_Base", "A:F")]


def test_entree_consulte_un_isin(wb):
    saisir(wb, isin="fr0010000002")
    assert sm.mvs_entree(wb, AUJOURDHUI) == "FSIC - ENREGISTREMENT TROUVE"
    assert champ(wb, sm.L_ISIN, sm.C_CHAMP, sm.LEN_ISIN) == "FR0010000002"      # mis en majuscules
    assert champ(wb, sm.L_LIB, sm.C_CHAMP, sm.LEN_LIB) == "FONDS OBLIG CORPORATE"
    assert champ(wb, sm.L_STAT, sm.C_CHAMP, sm.LEN_STAT) == "ACTIF"
    assert champ(wb, sm.L_ENC, sm.C_CHAMP, sm.LEN_ENC) == "890.2"
    assert champ(wb, sm.L_DATE, sm.C_CHAMP, sm.LEN_DATE) == "28/09/2026"
    assert champ(wb, sm.L_COM, sm.C_CHAMP, sm.LEN_COM) == "RAS"
    assert champ(wb, sm.L_MSG, sm.C_MSG, sm.LEN_MSG) == "FSIC - ENREGISTREMENT TROUVE"
    assert wb["MVS_Base"]["H1"].value == "FR0010000002"


def test_entree_met_a_jour_le_commentaire(wb):
    saisir(wb, isin="FR0010000002")
    sm.mvs_entree(wb, AUJOURDHUI)
    assert sm.mvs_entree(wb, AUJOURDHUI) == "AUCUNE MODIFICATION"              # même ISIN, commentaire inchangé
    saisir(wb, commentaire="NOUVEAU COMMENTAIRE")
    assert sm.mvs_entree(wb, AUJOURDHUI) == "ENREGISTREMENT MIS A JOUR"
    base = wb["MVS_Base"]
    assert base["F3"].value == "NOUVEAU COMMENTAIRE" and base["E3"].value == "27/10/2026"
    assert champ(wb, sm.L_DATE, sm.C_CHAMP, sm.LEN_DATE) == "27/10/2026"


def test_entree_isin_inconnu_et_vide(wb):
    saisir(wb, isin="FR9999999999")
    assert sm.mvs_entree(wb, AUJOURDHUI) == "ISIN INCONNU DANS FSIC : FR9999999999"
    assert champ(wb, sm.L_LIB, sm.C_CHAMP, sm.LEN_LIB) == ""
    saisir(wb, isin="")
    assert sm.mvs_entree(wb, AUJOURDHUI) == "SAISIR UN CODE ISIN PUIS ENTREE"


def test_entree_sans_ecran_signale_l_erreur(caplog, copie):
    wb = openpyxl.load_workbook(copie("classeur_complet.xlsx"))
    assert sm.mvs_entree(wb, AUJOURDHUI) is None
    assert "Erreur MVS_Entree : L'indice n'appartient pas à la sélection" in caplog.text


def test_effacer_et_pf3(wb):
    saisir(wb, isin="FR0010000001")
    sm.mvs_entree(wb, AUJOURDHUI)
    sm.mvs_effacer(wb)
    assert champ(wb, sm.L_ISIN, sm.C_CHAMP, sm.LEN_ISIN) == "" and champ(wb, sm.L_LIB, sm.C_CHAMP, sm.LEN_LIB) == ""
    assert champ(wb, sm.L_MSG, sm.C_MSG, sm.LEN_MSG) == "ECRAN EFFACE" and wb["MVS_Base"]["H1"].value in (None, "")
    sm.mvs_pf3(wb)
    assert champ(wb, sm.L_MSG, sm.C_MSG, sm.LEN_MSG) == "PF3 - SESSION TERMINEE (SIMULATION)"


def test_ecrire_champ_tronque_et_ignore_les_espaces(wb):
    ws = wb["MVS"]
    sm.ecrire_champ(ws, sm.L_ISIN, sm.C_CHAMP, sm.LEN_ISIN, "AB CD" + "X" * 20)
    assert sm.lire_champ(ws, sm.L_ISIN, sm.C_CHAMP, sm.LEN_ISIN) == "AB CDXXXXXXX"   # tronqué à 12
    assert ws["AC6"].value is None                                                  # l'espace laisse la cellule vide


def test_buffer_ligne_texte_multi_caracteres_deborde():
    ws = openpyxl.Workbook().active
    ws["A1"] = "AB"                  # un texte de 2 caractères dans une cellule occupe aussi la colonne suivante
    ws["C1"] = "Z"
    assert sm.buffer_ligne(ws, 1)[:4] == "ABZ "       # le B de « AB » occupe la colonne 2, Z écrase la colonne 3
    assert len(sm.buffer_ligne(ws, 1)) == 80
    ws["CB1"] = "FIN"                # colonne 80 : un seul caractère tient
    assert sm.buffer_ligne(ws, 1)[79] == "F"


def test_poser_limite_a_80_colonnes():
    ws = openpyxl.Workbook().active
    sm.poser(ws, 1, 79, "ABCD")
    assert (ws.cell(row=1, column=79).value, ws.cell(row=1, column=80).value, ws.cell(row=1, column=81).value) == (
        "A", "B", None)


def test_trouver_isin(wb):
    assert sm.trouver_isin(wb["MVS_Base"], "FR0010000004") == 5
    assert sm.trouver_isin(wb["MVS_Base"], "XX") == 0


def test_ligne_de_commande_sous_commandes(capsys):
    with pytest.raises(SystemExit):
        sm.main(["mvs-entree", "--help"])
    assert "classeur" in capsys.readouterr().out


def test_ligne_de_commande_enregistre_en_place(copie):
    classeur = copie("classeur_mvs.xlsx")
    wb = openpyxl.load_workbook(classeur)
    saisir(wb, isin="FR0010000001")
    wb.save(classeur)
    assert sm.main(["mvs-entree", str(classeur)]) == 0
    wb = openpyxl.load_workbook(classeur)
    assert champ(wb, sm.L_LIB, sm.C_CHAMP, sm.LEN_LIB) == "FONDS ACTIONS EUROPE"
    assert sm.main(["mvs-pf3", str(classeur)]) == 0
    assert champ(openpyxl.load_workbook(classeur), sm.L_MSG, sm.C_MSG, sm.LEN_MSG).startswith("PF3")
