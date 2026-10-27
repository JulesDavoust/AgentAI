"""Tests de macro_carte.py."""

import logging
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

import commun
import macro_carte as mc
from commun import ErreurVBA

NOW = datetime(2026, 10, 27, 14, 30, 5)


def _logs(wb):
    return [wb["Log"].cell(row=r, column=3).value for r in range(2, wb["Log"].max_row + 1)]


def test_coord_region():
    assert mc.coord_region("Centre") == "47.08,2.40"
    assert mc.coord_region("Atlantide") == ""


def test_ajouter_donnees_statut_v2(copie, rng):
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    assert mc.ajouter_donnees(wb, NOW, rng) == 3
    assert wb["Feuille2"]["E14"].value == "Phase de test V2"


@pytest.mark.parametrize("saisie, attendu", [("", None), ("Mars", None), ("nord", "nord"), ("Sud", "Sud")])
def test_appliquer_filtre_v2(copie, saisie, attendu):
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    # la région est reprise telle que saisie (Criteria1 d'Excel), sans le "="
    assert mc.appliquer_filtre_v2(wb, saisie, NOW) == attendu


def test_generer_carte_html(tmp_path, copie):
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    chemin = mc.generer_carte_html(wb, "Nord", tmp_path / "out", NOW)
    assert chemin == (tmp_path / "out" / "VBA_Natran_Test_Carte_20261027_143005_VBATF_V2.html").resolve()
    brut = chemin.read_bytes()
    assert brut.endswith(b"</html>\n\r\n")             # lignes en LF (vbLf) + CRLF final du Print
    assert b"\r\n" not in brut[:-2]
    html = brut.decode("cp1252")
    # 1re ligne de données : 25/10, Nord, 120 GWh (120.0 -> "120"), 6 degrés
    assert "var DATA=[['2026-10-25','Nord',120,6],['2026-10-25','Sud',95,6.5]," in html
    assert html.count("['2026-10-27'") == 2             # 12 lignes au total, 2 pour le 27/10
    assert "var COORD={'Nord':[50.63,3.06],'Sud':[43.30,5.37],'Est':[48.57,7.75],'Ouest':[47.22,-1.55],'Centre':[47.08,2.40]};" in html
    assert "selR.value='Nord';" in html
    assert "Genere le 27/10/2026 14:30 - VBATF_V2" in html
    assert "[3.60,50.35],[3.25,50.75],[2.60,50.85]];" in html
    assert "Carte generee : " + str(chemin) in _logs(wb)


def test_generer_carte_sans_region_et_table_vide(tmp_path, copie):
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    html = mc.generer_carte_html(wb, None, tmp_path, NOW).read_text(encoding="cp1252")
    assert "selR.value='*';" in html
    wb["Feuille2"].tables["tblDonnees"].ref = "A1:E1"       # seulement l'en-tête : DataBodyRange Nothing
    with pytest.raises(ErreurVBA) as e:
        mc.generer_carte_html(wb, None, tmp_path, NOW)
    assert e.value.numero == commun.VB_OBJECT_ERROR + 1 and e.value.description == "Table de donnees vide"


def test_creer_synthese(copie):
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    carte = Path("C:/x/carte.html")
    mc.creer_synthese(wb, carte, NOW)
    ws = wb["Feuille1"]
    assert ws["A1"].value == "Suivi distribution gaz (V2 carte) - mis a jour le 27/10/2026 14:30"
    assert ws["B3"].value == "=MAX(tblDonnees[Date])" and ws["B3"].number_format == "dd/mm/yyyy"
    assert [ws.cell(row=4, column=c).value for c in (1, 2, 3)] == ["Region", "Volume du jour (GWh)",
                                                                  "Temp. moyenne (C)"]
    assert ws["A4"].fill.start_color.rgb.endswith("DDEBF7")
    assert [ws.cell(row=r, column=1).value for r in range(5, 10)] == ["Nord", "Sud", "Est", "Ouest", "Centre"]
    assert ws["B5"].value == "=SUMIFS(tblDonnees[Volume_GWh],tblDonnees[Region],A5,tblDonnees[Date],$B$3)"
    assert ws["C9"].value.startswith("=IFERROR(AVERAGEIFS(tblDonnees[Temperature_C],tblDonnees[Region],A9,")
    # ligne = 10 après la boucle : lien en A11, texte « Fichier » en A12
    assert ws["A11"].value == "Ouvrir la carte interactive" and ws["A11"].hyperlink.target == str(carte)
    assert ws["A12"].value == f"Fichier : {carte}" and ws["A12"].font.i
    assert [str(r.sqref) for r in ws.conditional_formatting] == ["B5:B9"]
    assert wb.active.title == "Feuille1" and _logs(wb)[-1] == "Synthese regeneree"


def test_creer_synthese_efface_l_ancien_contenu(copie):
    wb = openpyxl.load_workbook(copie("classeur_donnees.xlsx"))
    wb["Feuille1"]["A20"] = "reste"
    wb["Feuille1"]["A20"].hyperlink = "http://exemple.invalide"
    mc.creer_synthese(wb, "c.html", NOW)
    assert wb["Feuille1"]["A20"].value is None and wb["Feuille1"]["A20"].hyperlink is None


class Espions:
    def __init__(self, ouvrir_ko=False):
        self.appels, self.ouvrir_ko = [], ouvrir_ko

    def excel(self, chemin, ajuster, region):
        assert Path(chemin).is_absolute() and Path(chemin).exists()
        self.appels.append(("excel", Path(chemin).name, ajuster, region))

    def mail(self, dest, cc, sujet, corps, pj, mode):
        assert all(Path(p).is_absolute() and Path(p).exists() for p in pj)
        self.appels.append(("mail", dest, sujet, corps, [Path(p).name for p in pj], mode))

    def ouvrir(self, chemin):
        assert Path(chemin).is_absolute() and Path(chemin).exists()
        if self.ouvrir_ko:
            raise OSError("pas de navigateur")
        self.appels.append(("ouvrir", Path(chemin).name))


def test_pipeline_carte(tmp_path, copie, rng, maintenant, monkeypatch):
    monkeypatch.chdir(tmp_path)
    e = Espions()
    res = mc.lancer_pipeline_carte_v2(copie("classeur_donnees.xlsx"), "sortie", "x@y.fr", "", "envoyer", True, "Est",
                                      envoyer_mail=e.mail, excel=e.excel, ouvrir_fichier=e.ouvrir, rng=rng,
                                      maintenant=maintenant)
    assert res.ok
    xlsm, carte = res.chemins
    assert xlsm.name == "VBA_Natran_Test_20261027_143005_VBATF_V2.xlsx"
    assert carte.name == "VBA_Natran_Test_Carte_20261027_143005_VBATF_V2.html"
    # la structure existait déjà (False) ; carte et classeur envoyés, puis carte ouverte, dans cet ordre
    assert e.appels[0] == ("excel", xlsm.name, False, "Est")
    assert [a[0] for a in e.appels] == ["excel", "mail", "ouvrir"]
    _, dest, sujet, corps, pj, mode = e.appels[1]
    assert (dest, mode, pj) == ("x@y.fr", "envoyer", [xlsm.name, carte.name])
    assert sujet == "[PHASE DE TEST V2] Mise a jour suivi distribution - 27/10/2026"
    assert f"Classeur : {xlsm.name}\r\nCarte : {carte.name} (a ouvrir dans un navigateur, fonctionne hors ligne)" in corps
    assert "selR.value='Est';" in carte.read_text(encoding="cp1252")
    copie_wb = openpyxl.load_workbook(xlsm)
    assert copie_wb["Feuille1"]["A5"].value == "Nord" and copie_wb["Feuille2"].tables["tblDonnees"].ref == "A1:E16"


def test_pipeline_carte_sans_ouverture(tmp_path, copie, rng, maintenant):
    e = Espions()
    res = mc.lancer_pipeline_carte_v2(copie("classeur_donnees.xlsx"), tmp_path, mode_mail="aucun", open_map=False,
                                      envoyer_mail=e.mail, excel=e.excel, ouvrir_fichier=e.ouvrir, rng=rng,
                                      maintenant=maintenant)
    assert res.ok and [a[0] for a in e.appels] == ["excel"]


def test_pipeline_carte_ouverture_impossible_est_consignee(tmp_path, copie, rng, maintenant, caplog):
    caplog.set_level(logging.INFO, logger="vba2py")
    e = Espions(ouvrir_ko=True)
    res = mc.lancer_pipeline_carte_v2(copie("classeur_donnees.xlsx"), tmp_path, mode_mail="aucun",
                                      excel=e.excel, ouvrir_fichier=e.ouvrir, rng=rng, maintenant=maintenant)
    assert res.ok                                                    # On Error Resume Next + LogMsg
    assert "Ouverture impossible : pas de navigateur" in caplog.text


def test_ligne_de_commande_declare_les_options(capsys):
    with pytest.raises(SystemExit):
        mc.main(["--help"])
    sortie = capsys.readouterr().out
    for option in ("--dossier-sortie", "--mail-to", "--mail-cc", "--mode-mail", "--open-map", "--saisie"):
        assert option in sortie
