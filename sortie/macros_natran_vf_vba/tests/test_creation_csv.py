"""Tests de creation_csv.py."""

from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

import creation_csv as cc

NOW = datetime(2026, 10, 27, 14, 30, 5)


def _log(chemin):
    ws = openpyxl.load_workbook(chemin)["Log"]
    return [(ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value) for r in range(2, ws.max_row + 1)]


def test_v3_nom_csv():
    assert cc.v3_nom_csv(date(2026, 10, 5), "dossier") == Path("dossier") / "DAT_20261005.csv"


def test_ecrire_csv_contenu_exact(tmp_path, rng):
    wb = openpyxl.Workbook()
    chemin = cc.ecrire_csv(wb, date(2026, 10, 27), tmp_path / "a" / "b", rng, lambda: NOW)   # dossiers parents créés
    assert chemin == tmp_path / "a" / "b" / "DAT_20261027.csv"
    brut = chemin.read_bytes()
    assert brut.count(b"\r\n") == 11 and b"\n" not in brut.replace(b"\r\n", b"")     # CRLF Windows partout
    lignes = brut.decode("cp1252").split("\r\n")[:-1]
    assert lignes[0] == "Date;Ville;Latitude;Longitude;Volume_GWh;Temperature_C"
    # rng = 0.5 : volume = base * 1.0 ; température = 2 + 0.5 * 14 = 9
    assert lignes[1] == "2026-10-27;Lille;50.63;3.06;95;9"
    assert lignes[2] == "2026-10-27;Paris;48.86;2.35;150;9"
    assert lignes[4] == "2026-10-27;Nantes;47.22;-1.55;70;9"                          # longitude négative
    assert lignes[10] == "2026-10-27;Dijon;47.32;5.04;45;9"
    assert wb["Log"]["C2"].value == f"CSV genere : {chemin}"


def test_ecrire_csv_bornes_aleatoires(tmp_path):
    class Haut:
        def random(self):
            return 0.9999

    wb = openpyxl.Workbook()
    chemin = cc.ecrire_csv(wb, date(2026, 10, 27), tmp_path, Haut(), lambda: NOW)
    assert chemin.read_text(encoding="cp1252").splitlines()[2].endswith(";172.5;16")  # Paris 150*1.15 ; 2+14


def test_ecrire_csv_existant_ecrase(tmp_path, rng):
    wb = openpyxl.Workbook()
    cc.ecrire_csv(wb, date(2026, 10, 27), tmp_path, rng, lambda: NOW)
    cc.ecrire_csv(wb, date(2026, 10, 27), tmp_path, rng, lambda: NOW)
    messages = [wb["Log"].cell(row=r, column=3).value for r in range(2, wb["Log"].max_row + 1)]
    assert any(m.startswith("CSV existant ecrase : ") for m in messages)


def test_generer_csv_du_jour(tmp_path, copie, rng, maintenant):
    classeur = copie("classeur_complet.xlsx")
    res = cc.generer_csv_du_jour_v3(classeur, tmp_path / "csv", rng=rng, maintenant=maintenant)
    assert res.ok and res.chemins == [tmp_path / "csv" / "DAT_20261027.csv"] and res.chemins[0].exists()
    assert res.message == f"CSV genere :\n{res.chemins[0]}"
    assert _log(classeur) == [("INFO", f"CSV genere : {res.chemins[0]}")]            # classeur enregistré en place


@pytest.mark.parametrize("saisie, fichier", [("05/10/2026", "DAT_20261005.csv"), (None, "DAT_20261027.csv"),
                                              (" 1/2/26 ", "DAT_20260201.csv")])
def test_generer_csv_date(tmp_path, copie, rng, maintenant, saisie, fichier):
    res = cc.generer_csv_date_v3(copie("classeur_complet.xlsx"), tmp_path, saisie, rng=rng, maintenant=maintenant)
    assert res.ok and res.chemins[0].name == fichier       # 05/10/2026 = 5 octobre ; None = date du jour (défaut InputBox)


def test_generer_csv_date_vide_ou_invalide(tmp_path, copie, rng, maintenant, caplog):
    vide = cc.generer_csv_date_v3(copie("classeur_complet.xlsx"), tmp_path, "  ", rng=rng, maintenant=maintenant)
    assert vide.ok and vide.chemins == []                  # saisie vide : la macro s'arrête sans rien faire
    invalide = cc.generer_csv_date_v3(copie("classeur_complet.xlsx"), tmp_path, "31/02/2026", rng=rng,
                                      maintenant=maintenant)
    assert not invalide.ok and invalide.message == "Date invalide : 31/02/2026"
    assert list(tmp_path.glob("DAT_*.csv")) == []


def test_erreur_consignee_dans_le_log(tmp_path, copie, rng, maintenant):
    occupe = tmp_path / "fichier.txt"
    occupe.write_text("x")                                  # un fichier à la place du dossier : MkDir échoue
    classeur = copie("classeur_complet.xlsx")
    res = cc.generer_csv_du_jour_v3(classeur, occupe, rng=rng, maintenant=maintenant)
    assert not res.ok and res.message.startswith("Erreur : ")
    niveau, message = _log(classeur)[0]
    assert niveau == "ERREUR" and message.startswith("Generation CSV : 0 - ")


def test_ligne_de_commande(capsys):
    with pytest.raises(SystemExit):
        cc.main(["generer-csv-date-v3", "--help"])
    sortie = capsys.readouterr().out
    assert "--v3-csv-folder" in sortie and "--saisie" in sortie
