"""Traduction du module VBA ``SimuMainframe`` : simulation d'un écran mainframe MVS (3278 modèle 3,
32x80) dessiné dans la feuille « MVS », avec une base de fonds dans la feuille masquée « MVS_Base ».

Le classeur (``ThisWorkbook``) est l'état de la simulation : chaque commande le charge, agit, puis
l'enregistre **en place** (équivalent du classeur resté ouvert dans Excel).
"""

from __future__ import annotations

import argparse
import logging
import sys
from datetime import date, datetime

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import commun
from commun import cstr
from connecteurs import excel as connecteur_excel
from bibliotheque.excel_com import preparer_feuille

SHEET_MVS = "MVS"
SHEET_BASE = "MVS_Base"
NB_LIGNES = 32
NB_COLS = 80

# Coordonnées des champs (identiques dans TestInterfaceMVS)
C_CHAMP = 27
L_ISIN, LEN_ISIN = 6, 12
L_LIB, LEN_LIB = 8, 40
L_STAT, LEN_STAT = 10, 10
L_ENC, LEN_ENC = 12, 12
L_DATE, LEN_DATE = 14, 10
L_COM, LEN_COM = 16, 40
L_MSG, C_MSG, LEN_MSG = 30, 13, 60

# Palette NATRAN (RGB -> hexadécimal)
C_FOND = "081E3C"      # RGB(8, 30, 60)
C_TEXTE = "EBF2FA"     # RGB(235, 242, 250)
C_TITRE = "FFFFFF"     # RGB(255, 255, 255)
C_LABEL = "00BEAF"     # RGB(0, 190, 175)
C_SAISIE = "00549F"    # RGB(0, 84, 159)
C_SEP = "00788C"       # RGB(0, 120, 140)
C_MESSAGE = "96D250"   # RGB(150, 210, 80)
C_REGLE = "788CAA"     # RGB(120, 140, 170)

LIGNES_BASE = (
    "FR0010000001;FONDS ACTIONS EUROPE;ACTIF;1250.5;01/10/2026;",
    "FR0010000002;FONDS OBLIG CORPORATE;ACTIF;890.2;28/09/2026;RAS",
    "FR0010000003;FONDS MONETAIRE EURO;ACTIF;2310.0;15/10/2026;",
    "FR0010000004;FONDS DIVERSIFIE PRUDENT;ACTIF;455.8;03/10/2026;A SURVEILLER",
    "FR0010000005;FONDS ACTIONS USA;FERME;120.4;30/06/2026;LIQUIDATION",
    "FR0010000006;FONDS INFRA ENERGIE;ACTIF;780.0;20/10/2026;",
)


def feuille_existe(wb, nom: str) -> bool:
    """VBA ``FeuilleExiste``"""
    return nom in wb.sheetnames


# ----------------------------------------------------------------------------- point d'entrée : ouvrir

def ouvrir_mainframe(classeur, *, excel=None, maintenant=datetime.now) -> bool:
    """VBA ``OuvrirMainframe`` : crée l'écran s'il n'existe pas, sinon l'affiche. Le classeur est
    enregistré en place. Retourne True si l'écran a été créé."""
    wb = commun.charger_classeur(classeur)
    cree = not feuille_existe(wb, SHEET_MVS)
    base_recreee = False
    if cree:
        base_recreee = creer_ecran_mvs(wb, maintenant().date())
    else:
        ws = wb[SHEET_MVS]
        ws.sheet_state = "visible"
        commun.activer_feuille(wb, ws)
        commun.selectionner(ws, f"{get_column_letter(C_CHAMP)}{L_ISIN}")
    wb.save(classeur)
    if base_recreee:  # AutoFit de MVS_Base : Excel
        (excel or connecteur_excel.ajuster_feuille)(classeur, SHEET_BASE, "A:F")
    return cree


def creer_ecran_mvs(wb, aujourdhui: date) -> bool:
    """VBA ``CreerEcranMVS`` : (ré)crée la base si besoin, dessine l'écran, vue sans quadrillage, zoom
    100 %, curseur sur le champ ISIN. Retourne True si MVS_Base vient d'être (re)créée."""
    base_creee = False
    if not feuille_existe(wb, SHEET_BASE):
        reinitialiser_base_mvs(wb)
        base_creee = True
    ws = commun.get_or_create_sheet(wb, SHEET_MVS)
    dessiner_ecran(ws, aujourdhui)
    commun.activer_feuille(wb, ws)
    ws.sheet_view.showGridLines = False  # ActiveWindow.DisplayGridlines = False
    ws.sheet_view.zoomScale = 100        # ActiveWindow.Zoom = 100
    commun.selectionner(ws, f"{get_column_letter(C_CHAMP)}{L_ISIN}")
    return base_creee


def reinitialiser_base_mvs(wb) -> None:
    """VBA ``ReinitialiserBaseMVS`` : feuille MVS_Base (texte) avec 6 fonds, masquée. L'AutoFit
    des colonnes A:F est fait par Excel (``connecteurs.excel.ajuster_feuille``)."""
    ws = commun.get_or_create_sheet(wb, SHEET_BASE)
    preparer_feuille(ws)  # Cells.Clear
    for lettre in "ABCDEFGH":
        ws.column_dimensions[lettre].number_format = "@"  # Cells.NumberFormat = "@"
    for i, titre in enumerate(("ISIN", "Libelle", "Statut", "Encours_MEUR", "Date_MAJ", "Commentaire"), 1):
        ws.cell(row=1, column=i, value=titre).number_format = "@"
    for i, ligne in enumerate(LIGNES_BASE):
        for j, champ in enumerate(ligne.split(";")):
            c = ws.cell(row=i + 2, column=j + 1, value=champ if champ != "" else None)
            c.number_format = "@"
    ws["H1"].number_format = "@"  # Range("H1").Value = ""
    ws.sheet_state = "hidden"


# ----------------------------------------------------------------------------- dessin de l'écran

def _fill(couleur: str) -> PatternFill:
    return PatternFill("solid", start_color=couleur, end_color=couleur)


def _police(couleur: str) -> Font:
    return Font(name="Courier New", size=10, color=couleur)


def dessiner_ecran(ws, aujourdhui: date) -> None:
    """VBA ``DessinerEcran`` : fond bleu nuit, textes et libellés de l'écran FSIC, règle de coordonnées."""
    preparer_feuille(ws)  # ws.Cells.Clear
    fond = _fill(C_FOND)
    for c in range(1, NB_COLS + 1):  # ws.Cells.Interior.Color + ColumnWidth = 1.5
        cd = ws.column_dimensions[get_column_letter(c)]
        cd.width = 1.5
        cd.fill = fond
    reste = ws.column_dimensions[get_column_letter(NB_COLS + 1)]
    reste.min, reste.max, reste.fill = NB_COLS + 1, 16384, fond
    for l in range(1, 37):
        ws.row_dimensions[l].height = 15
        for c in range(1, NB_COLS + 1):
            cell = ws.cell(row=l, column=c)
            cell.fill = fond
            cell.number_format = "@"
            cell.font = _police(C_TEXTE)
            cell.alignment = Alignment(horizontal="left")

    poser(ws, 1, 2, "NATRAN - SIMULATION MAINFRAME MVS (3278 MODELE 3 - 32x80)", C_TITRE)
    poser(ws, 1, 62, "DATE " + commun.fmt_date(aujourdhui), C_TITRE)
    poser(ws, 2, 1, "-" * NB_COLS, C_SEP)
    poser(ws, 4, 3, "ECRAN FSIC - CONSULTATION / MISE A JOUR FONDS", C_TITRE)
    poser(ws, L_ISIN, 3, "CODE ISIN . . . . . :", C_LABEL)
    poser(ws, L_LIB, 3, "LIBELLE . . . . . . :", C_LABEL)
    poser(ws, L_STAT, 3, "STATUT . . . . . . :", C_LABEL)
    poser(ws, L_ENC, 3, "ENCOURS (MEUR) . . :", C_LABEL)
    poser(ws, L_DATE, 3, "DERNIERE MAJ . . . :", C_LABEL)
    poser(ws, L_COM, 3, "COMMENTAIRE . . . :", C_LABEL)
    poser(ws, 31, 1, "-" * NB_COLS, C_SEP)
    poser(ws, L_MSG, 3, "MESSAGE :", C_LABEL)
    poser(ws, NB_LIGNES, 3, "ENTREE=Valider  PF3=Quitter  CLEAR=Effacer", C_TITRE)

    style_saisie(ws, L_ISIN, LEN_ISIN)
    style_saisie(ws, L_COM, LEN_COM)
    for c in range(C_MSG, C_MSG + LEN_MSG):
        ws.cell(row=L_MSG, column=c).font = _police(C_MESSAGE)

    for c in range(1, NB_COLS + 1):
        if c % 10 == 0:
            ws.cell(row=34, column=c, value=str(c // 10))
        ws.cell(row=35, column=c, value=str(c % 10))
    for l in (34, 35):
        for c in range(1, NB_COLS + 1):
            ws.cell(row=l, column=c).font = _police(C_REGLE)
    poser(ws, 36, 1, "(regle de coordonnees, hors ecran)", C_REGLE)

    message(ws, "PRET - SAISIR UN CODE ISIN PUIS ENTREE")


def style_saisie(ws, l: int, n: int) -> None:
    """VBA ``StyleSaisie`` : zone de saisie en bleu, texte blanc."""
    for c in range(C_CHAMP, C_CHAMP + n):
        cell = ws.cell(row=l, column=c)
        cell.fill = _fill(C_SAISIE)
        cell.font = _police(C_TITRE)


def poser(ws, l: int, c: int, t: str, couleur: str | None = None) -> None:
    """VBA ``Poser`` : un caractère par cellule (espaces ignorés), colonnes limitées à 80."""
    for k, ch in enumerate(t):
        if c + k <= NB_COLS and ch != " ":
            ws.cell(row=l, column=c + k, value=ch)
    if couleur is not None:
        for col in range(c, min(c + len(t) - 1, NB_COLS) + 1):
            ws.cell(row=l, column=col).font = _police(couleur)


# ----------------------------------------------------------------------------- outils écran

def message(ws, txt: str) -> None:
    """VBA ``Message``"""
    ecrire_champ(ws, L_MSG, C_MSG, LEN_MSG, txt)


def effacer_sorties(ws) -> None:
    """VBA ``EffacerSorties``"""
    for l, n in ((L_LIB, LEN_LIB), (L_STAT, LEN_STAT), (L_ENC, LEN_ENC), (L_DATE, LEN_DATE), (L_COM, LEN_COM)):
        ecrire_champ(ws, l, C_CHAMP, n, "")


def ecrire_champ(ws, l: int, c: int, n: int, t: str) -> None:
    """VBA ``EcrireChamp`` : efface la zone puis écrit ``t`` tronqué à ``n`` caractères."""
    for col in range(c, c + n):
        ws.cell(row=l, column=col).value = None
    for k, ch in enumerate(t[:n]):
        if ch != " ":
            ws.cell(row=l, column=c + k, value=ch)


def buffer_ligne(ws, l: int) -> str:
    """VBA ``BufferLigne`` : la ligne de l'écran sous forme d'une chaîne de 80 caractères."""
    buf = [" "] * NB_COLS
    for c in range(1, NB_COLS + 1):
        for k, ch in enumerate(cstr(ws.cell(row=l, column=c).value)):
            if c + k <= NB_COLS:
                buf[c + k - 1] = ch
    return "".join(buf)


def lire_champ(ws, l: int, c: int, n: int) -> str:
    """VBA ``LireChamp``"""
    return buffer_ligne(ws, l)[c - 1:c - 1 + n]


def normaliser_ligne(ws, l: int) -> None:
    """VBA ``NormaliserLigne`` : une lettre par cellule."""
    buf = buffer_ligne(ws, l)
    for c in range(1, NB_COLS + 1):
        ws.cell(row=l, column=c).value = None
    for c in range(1, NB_COLS + 1):
        if buf[c - 1] != " ":
            ws.cell(row=l, column=c, value=buf[c - 1])


def trouver_isin(ws_b, isin: str) -> int:
    """VBA ``TrouverISIN`` : numéro de ligne dans MVS_Base, 0 si absent."""
    for r in range(2, commun.derniere_ligne(ws_b, 1) + 1):
        if cstr(ws_b.cell(row=r, column=1).value).strip().upper() == isin:
            return r
    return 0


# ----------------------------------------------------------------------------- touches du terminal

def mvs_entree(wb, aujourdhui: date) -> str | None:
    """VBA ``MVS_Entree`` (touche ENTREE) : consultation de l'ISIN saisi ou mise à jour du commentaire.
    Retourne le message affiché sur l'écran (None en cas d'erreur, signalée comme le MsgBox VBA)."""
    try:
        ws = commun.feuille(wb, SHEET_MVS)
        ws_b = commun.feuille(wb, SHEET_BASE)
        normaliser_ligne(ws, L_ISIN)
        normaliser_ligne(ws, L_COM)
        isin = lire_champ(ws, L_ISIN, C_CHAMP, LEN_ISIN).strip().upper()
        isin_aff = cstr(ws_b["H1"].value)

        if isin == "":
            texte = "SAISIR UN CODE ISIN PUIS ENTREE"
        else:
            ecrire_champ(ws, L_ISIN, C_CHAMP, LEN_ISIN, isin)
            r = trouver_isin(ws_b, isin)
            if r == 0:
                effacer_sorties(ws)
                ws_b["H1"] = ""
                texte = "ISIN INCONNU DANS FSIC : " + isin
            elif isin != isin_aff:
                for l, n, col in ((L_LIB, LEN_LIB, 2), (L_STAT, LEN_STAT, 3), (L_ENC, LEN_ENC, 4),
                                  (L_DATE, LEN_DATE, 5), (L_COM, LEN_COM, 6)):
                    ecrire_champ(ws, l, C_CHAMP, n, cstr(ws_b.cell(row=r, column=col).value))
                ws_b["H1"] = isin
                texte = "FSIC - ENREGISTREMENT TROUVE"
            else:
                com = lire_champ(ws, L_COM, C_CHAMP, LEN_COM).strip()
                if com != cstr(ws_b.cell(row=r, column=6).value).strip():
                    ws_b.cell(row=r, column=6).value = com
                    ws_b.cell(row=r, column=5).value = commun.fmt_date(aujourdhui)
                    ecrire_champ(ws, L_DATE, C_CHAMP, LEN_DATE, cstr(ws_b.cell(row=r, column=5).value))
                    texte = "ENREGISTREMENT MIS A JOUR"
                else:
                    texte = "AUCUNE MODIFICATION"
        message(ws, texte)
        return texte
    except Exception as erreur:  # fin: MsgBox "Erreur MVS_Entree : ..."
        commun.journal.error("Erreur MVS_Entree : %s", getattr(erreur, "description", erreur))
        return None


def mvs_effacer(wb) -> None:
    """VBA ``MVS_Effacer`` (touche CLEAR)."""
    ws = commun.feuille(wb, SHEET_MVS)
    normaliser_ligne(ws, L_ISIN)
    ecrire_champ(ws, L_ISIN, C_CHAMP, LEN_ISIN, "")
    effacer_sorties(ws)
    commun.feuille(wb, SHEET_BASE)["H1"] = ""
    message(ws, "ECRAN EFFACE")


def mvs_pf3(wb) -> None:
    """VBA ``MVS_PF3`` (touche PF3)."""
    mvs_effacer(wb)
    message(commun.feuille(wb, SHEET_MVS), "PF3 - SESSION TERMINEE (SIMULATION)")


# ----------------------------------------------------------------------------- ligne de commande

def main(argv=None) -> int:
    p = argparse.ArgumentParser(description="VBA SimuMainframe : simulation d'un écran mainframe MVS")
    sous = p.add_subparsers(dest="commande", required=True)
    aides = {
        "ouvrir-mainframe": "OuvrirMainframe : crée ou affiche l'écran MVS",
        "mvs-entree": "MVS_Entree : touche ENTREE",
        "mvs-effacer": "MVS_Effacer : touche CLEAR",
        "mvs-pf3": "MVS_PF3 : touche PF3",
    }
    for nom, aide in aides.items():
        sous.add_parser(nom, help=aide).add_argument(
            "classeur", help="classeur (remplace ThisWorkbook ; enregistré en place)")
    a = p.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(message)s")

    if a.commande == "ouvrir-mainframe":
        ouvrir_mainframe(a.classeur)
        return 0
    wb = commun.charger_classeur(a.classeur)
    code = 0
    if a.commande == "mvs-entree":
        texte = mvs_entree(wb, date.today())
        code = 0 if texte is not None else 1
        if texte:
            logging.info(texte)
    elif a.commande == "mvs-effacer":
        mvs_effacer(wb)
    else:
        mvs_pf3(wb)
    wb.save(a.classeur)
    return code


if __name__ == "__main__":
    sys.exit(main())
