"""Traduction du module VBA ``UtilisationCSV`` : importe le CSV du jour (``DAT_aaaammjj.csv``), génère
la carte HTML des villes, la synthèse (Carte_V3), sauvegarde datée, mail (classeur + carte + CSV).
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Font, PatternFill

import commun
from commun import ErreurVBA, Resultat, log_msg
from connecteurs import excel as connecteur_excel
from bibliotheque.excel_com import preparer_feuille

CSV_FOLDER = "csv"  # VBA : "C:\Users\jdavoust\natran\VBATF_CSV\" (chemin personnel remplacé)
CSV_PREFIX = "DAT_"
DATE_FORMAT = "%Y%m%d"  # "yyyymmdd"
CSV_SEP = ";"
SHEET_DATA = "Donnees_V3"
SHEET_VISUEL = "Carte_V3"
TABLE_NAME = "tblCSV_V3"
FILE_TITLE = "VBA_Natran_Test"
FILE_SUFFIX = "VBATF_V3"
MAIL_TO = "tom.ferrandis@hotmail.com"
MAIL_CC = ""


def nom_csv(d: date, csv_folder=CSV_FOLDER) -> Path:
    """VBA ``NomCSV``"""
    return Path(csv_folder) / f"{CSV_PREFIX}{d.strftime(DATE_FORMAT)}.csv"


# ----------------------------------------------------------------------------- 1. import du CSV

def lire_lignes_csv(chemin) -> list[str]:
    """``Open ... For Input`` + ``Line Input`` : lignes en Windows-1252, sans saut de ligne."""
    with open(chemin, encoding="cp1252", newline="") as f:
        texte = f.read()
    lignes = re.split(r"\r\n|\r|\n", texte)
    if lignes and lignes[-1] == "":
        lignes.pop()
    return lignes


def est_nombre(s) -> bool:
    """VBA ``EstNombre`` : le point est remplacé par le séparateur décimal français (``,``)."""
    t = commun.cstr(s).strip()
    if t == "" or "," in t:
        return False
    return commun.is_numeric(t.replace(".", ","))


def ligne_valide(champs: list[str]) -> bool:
    """VBA ``LigneValide`` : 6 champs, date aaaa-mm-jj, ville non vide, 4 nombres, coordonnées dans la France."""
    if len(champs) != 6:
        return False
    d = champs[0]
    if len(d.strip()) != 10:
        return False
    if d[4:5] != "-" or d[7:8] != "-":
        return False
    if not (commun.is_numeric(d[:4]) and commun.is_numeric(d[5:7]) and commun.is_numeric(d[8:10])):
        return False
    if champs[1].strip() == "":
        return False
    if not all(est_nombre(c) for c in champs[2:6]):
        return False
    lat, lon = commun.vba_val(champs[2].strip()), commun.vba_val(champs[3].strip())
    return not (lat < 41 or lat > 52 or lon < -5.5 or lon > 10)


def importer_csv(wb, chemin, maintenant: datetime) -> int:
    """VBA ``ImporterCSV`` : remplit Donnees_V3 + tableau tblCSV_V3, retourne le nombre de lignes valides."""
    ws = commun.get_or_create_sheet(wb, SHEET_DATA)
    for nom in list(ws.tables):  # ListObjects(1).Delete
        del ws.tables[nom]
    preparer_feuille(ws)  # Cells.Clear
    for i, titre in enumerate(("Date", "Ville", "Latitude", "Longitude", "Volume_GWh", "Temperature_C"), 1):
        ws.cell(row=1, column=i, value=titre)

    n = 1
    nb_invalides = 0
    lignes = lire_lignes_csv(chemin)
    for ligne in lignes[1:]:  # la première (en-tête) est ignorée
        if len(ligne.strip()) > 0:
            champs = ligne.split(CSV_SEP)
            if ligne_valide(champs):
                n += 1
                c0 = champs[0]
                ws.cell(row=n, column=1, value=commun.date_serial(
                    commun.c_int(c0[:4]), commun.c_int(c0[5:7]), commun.c_int(c0[8:10])))
                ws.cell(row=n, column=2, value=champs[1].strip())
                for k in range(2, 6):
                    ws.cell(row=n, column=k + 1, value=commun.vba_val(champs[k].strip()))
            else:
                nb_invalides += 1
                log_msg(wb, "AVERTISSEMENT", f"Ligne ignoree : {ligne}", maintenant)

    if n == 1:
        raise ErreurVBA(commun.VB_OBJECT_ERROR + 11, f"Aucune ligne valide dans {chemin}")

    commun.creer_table(ws, TABLE_NAME, f"A1:F{n}")
    commun.formater_colonne(ws, "A", "dd/mm/yyyy")
    log_msg(wb, "INFO", f"CSV importe : {n - 1} valide(s), {nb_invalides} ignoree(s) - {chemin}", maintenant)
    return n - 1


# ----------------------------------------------------------------------------- 2. choix de la ville

def demander_ville(wb, saisie: str, maintenant: datetime | None = None) -> str:
    """VBA ``DemanderVille`` : ``saisie`` remplace l'InputBox (vide = toutes les villes)."""
    ws = commun.feuille(wb, SHEET_DATA)
    villes = [l[1] for l in commun.lignes_table(ws, commun.table_de(ws, TABLE_NAME))]
    saisie = (saisie or "").strip()
    if saisie == "":
        return ""
    for v in villes:
        if commun.cstr(v).casefold() == saisie.casefold():
            log_msg(wb, "INFO", f"Ville selectionnee : {commun.cstr(v)}", maintenant)
            return commun.cstr(v)
    commun.journal.warning("Ville inconnue : %s\nToutes les villes seront affichees.", saisie)  # MsgBox
    return ""


# ----------------------------------------------------------------------------- 3. carte HTML

def _js(s: str) -> str:
    return s.replace("'", "\\'")


def generer_carte_html(wb, ville_pre: str, dossier, maintenant: datetime) -> Path:
    """VBA ``GenererCarteHTML`` : écrit la carte (Windows-1252) dans ``dossier``, retourne son chemin."""
    ws = commun.feuille(wb, SHEET_DATA)
    lignes = commun.lignes_table(ws, commun.table_de(ws, TABLE_NAME))
    if not lignes:
        raise ErreurVBA(commun.VB_OBJECT_ERROR + 1, "Table de donnees vide")

    date_donnees = commun.fmt_date(lignes[0][0])
    donnees = ",".join(
        f"['{_js(commun.cstr(l[1]))}',{commun.str_nb(l[2])},{commun.str_nb(l[3])},{commun.str_nb(l[4])},"
        f"{commun.str_nb(l[5])}]" for l in lignes)
    sel = ville_pre or "*"

    h: list[str] = list(commun.HTML_ENTETE)
    h.append(f"<div id='panel'><b>Distribution gaz - donnees du {date_donnees}</b> Ville <select id='selV'></select>")
    h.append("<span id='tot'></span><small>Taille et couleur = volume | Genere le "
             f"{commun.fmt_date_heure(maintenant)} - {FILE_SUFFIX}</small></div>")
    h += commun.HTML_SVG
    h.append("<div id='info'>Cliquer sur une ville pour le detail.</div><script>")
    h.append(f"var DATA=[{donnees}];")
    h += commun.HTML_CONTOUR
    h += [
        "var selV=document.getElementById('selV');",
        "var names=DATA.map(function(r){return r[0];}).sort();",
        "selV.add(new Option('Toutes','*'));names.forEach(function(n){selV.add(new Option(n,n));});",
        f"selV.value='{_js(sel)}';",
        "var MAXV=Math.max.apply(null,DATA.map(function(r){return r[3];}));",
        "function color(v){var x=v/MAXV;return x>0.85?'#b30000':x>0.65?'#e34a33':x>0.45?'#fc8d59':'#fdcc8a';}",
        "function el(n,a){var e=document.createElementNS(NS,n);for(var k in a)e.setAttribute(k,a[k]);return e;}",
        "function draw(){while(g.firstChild)g.removeChild(g.firstChild);var f=selV.value,tot=0;",
        "DATA.forEach(function(r){if(f!=='*'&&r[0]!==f)return;tot+=r[3];",
        "var rad=8+26*r[3]/MAXV,cx=X(r[2]),cy=Y(r[1]);",
        "var ci=el('circle',{cx:cx,cy:cy,r:rad,fill:color(r[3]),'fill-opacity':0.8,stroke:'#333','stroke-width':1});",
        "var tt=el('title',{});tt.textContent=r[0]+' : '+r[3].toFixed(1)+' GWh';ci.appendChild(tt);",
        "ci.onclick=function(){document.getElementById('info').innerText=r[0]+' : '+r[3].toFixed(1)+' GWh, temperature '+r[4].toFixed(1)+' C';};",
        "g.appendChild(ci);",
        "var t=el('text',{x:cx,y:cy+rad+13,'text-anchor':'middle','font-size':11,'font-weight':'bold',fill:'#222',stroke:'#fff','stroke-width':3,'paint-order':'stroke','pointer-events':'none'});",
        "t.textContent=r[0];g.appendChild(t);});",
        "document.getElementById('tot').innerText='Total affiche : '+tot.toFixed(1)+' GWh';}",
        "selV.onchange=draw;draw();</script></body></html>",
    ]
    texte = "".join(ligne + "\n" for ligne in h)

    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    chemin = (dossier / f"{FILE_TITLE}_Carte_{maintenant.strftime('%Y%m%d_%H%M%S')}_{FILE_SUFFIX}.html").resolve()
    commun.ecrire_texte_cp1252(chemin, texte)
    log_msg(wb, "INFO", f"Carte generee : {chemin}", maintenant)
    return chemin


# ----------------------------------------------------------------------------- 4. synthèse

def creer_synthese(wb, chemin_carte, chemin_csv, maintenant: datetime) -> None:
    """VBA ``CreerSynthese`` (partie openpyxl) : feuille Carte_V3. L'AutoFit est fait par Excel."""
    ws = commun.get_or_create_sheet(wb, SHEET_VISUEL)
    lo = commun.table_de(commun.feuille(wb, SHEET_DATA), TABLE_NAME)
    nb = len(commun.lignes_table(wb[SHEET_DATA], lo))
    preparer_feuille(ws)  # ChartObjects.Delete, PivotTables.Clear, Cells.Clear
    for row in ws.iter_rows():
        for c in row:
            if c.hyperlink:
                c.hyperlink = None  # Hyperlinks.Delete
    ws.conditional_formatting = ConditionalFormattingList()

    ws["A1"] = "Suivi distribution gaz (V3 CSV) - mis a jour le " + commun.fmt_date_heure(maintenant)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Volume total (GWh)"
    ws["B2"] = f"=SUM({TABLE_NAME}[Volume_GWh])"
    ws["D2"] = "Moyenne / ville (GWh)"
    ws["E2"] = f"=AVERAGE({TABLE_NAME}[Volume_GWh])"
    ws["G2"] = "Ville du pic"
    ws["H2"] = f"=INDEX({TABLE_NAME}[Ville],MATCH(MAX({TABLE_NAME}[Volume_GWh]),{TABLE_NAME}[Volume_GWh],0))"
    for ref in ("A2", "D2", "G2"):
        ws[ref].font = Font(bold=True)
    for ref in ("B2", "E2"):
        ws[ref].number_format = "#,##0.0"
    ws["A3"] = "Date des donnees"
    ws["B3"] = f"=MAX({TABLE_NAME}[Date])"
    ws["B3"].number_format = "dd/mm/yyyy"
    for i, titre in enumerate(("Ville", "Volume (GWh)", "Temperature (C)"), 1):
        c = ws.cell(row=4, column=i, value=titre)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")

    ligne = 5
    for i in range(1, nb + 1):
        ws.cell(row=ligne, column=1, value=f"=INDEX({TABLE_NAME}[Ville],{i})")
        ws.cell(row=ligne, column=2, value=f"=INDEX({TABLE_NAME}[Volume_GWh],{i})")
        ws.cell(row=ligne, column=3, value=f"=INDEX({TABLE_NAME}[Temperature_C],{i})")
        ligne += 1
    for l in range(5, ligne):
        ws.cell(row=l, column=2).number_format = "#,##0.0"
        ws.cell(row=l, column=3).number_format = "#,##0.0"
    ws.conditional_formatting.add(f"B5:B{ligne - 1}", DataBarRule(start_type="min", end_type="max", color="638EC6"))

    lien = ws.cell(row=ligne + 1, column=1)
    lien.hyperlink = str(chemin_carte)
    lien.value = "Ouvrir la carte interactive"
    lien.style = "Hyperlink"
    ws.cell(row=ligne + 2, column=1, value=f"Source CSV : {chemin_csv}").font = Font(italic=True)
    ws.cell(row=ligne + 3, column=1, value=f"Carte : {chemin_carte}").font = Font(italic=True)

    commun.activer_feuille(wb, ws)
    log_msg(wb, "INFO", "Synthese regeneree", maintenant)


# ----------------------------------------------------------------------------- 5. sauvegarde / 6. mail

def sauvegarder_version(wb, dossier, extension: str, maintenant: datetime) -> Path:
    """VBA ``SauvegarderVersion`` : copie ``titre_date_VBATF_V3`` (dossier et parents créés)."""
    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    chemin = (dossier / commun.nom_version(FILE_TITLE, maintenant, FILE_SUFFIX, extension)).resolve()
    wb.save(chemin)
    return chemin


def envoyer_mail_pipeline(wb, chemin_xlsm, chemin_carte, chemin_csv, mail_to, mail_cc, mode_mail, maintenant,
                          envoyer_mail=None) -> bool:
    """VBA ``EnvoyerMail`` (UtilisationCSV) : classeur, carte et CSV source en pièces jointes."""
    sujet = "[PHASE DE TEST V3] Mise a jour suivi distribution - " + commun.fmt_date(maintenant)
    corps = ("Bonjour,\r\n\r\n"
             "Ceci est un envoi automatique en phase de test (version CSV + carte interactive).\r\n"
             f"Les donnees ont ete importees depuis : {Path(chemin_csv).name}\r\n"
             f"Classeur : {Path(chemin_xlsm).name}\r\n"
             f"Carte : {Path(chemin_carte).name} (a ouvrir dans un navigateur, fonctionne hors ligne)\r\n\r\n"
             "Cordialement,\r\nRobot VBATF")
    return commun.expedier_mail(wb, envoyer_mail, mode_mail, mail_to, mail_cc, sujet, corps,
                                [chemin_xlsm, chemin_carte, chemin_csv], maintenant)


# ----------------------------------------------------------------------------- point d'entrée

def lancer_pipeline_csv_v3(classeur, dossier_sortie="sortie", csv_folder=CSV_FOLDER, mail_to=MAIL_TO,
                           mail_cc=MAIL_CC, mode_mail="envoyer", open_map=True, saisie="", *, envoyer_mail=None,
                           excel=None, ouvrir_fichier=None, maintenant=datetime.now) -> Resultat:
    """VBA ``LancerPipelineCSVV3``. ``classeur`` = ``ThisWorkbook`` (non modifié : le résultat est la copie
    enregistrée) ; ``saisie`` = InputBox de ``DemanderVille`` (vide = toutes les villes)."""
    excel = excel or connecteur_excel.finaliser_csv_v3
    wb = commun.charger_classeur(classeur)
    try:
        chemin_csv = nom_csv(maintenant().date(), csv_folder)
        if not chemin_csv.is_file():
            raise ErreurVBA(commun.VB_OBJECT_ERROR + 10,
                            f"CSV du jour introuvable : {chemin_csv}\r\n"
                            "Genere d'abord le CSV avec le module de creation.")
        chemin_csv = chemin_csv.resolve()
        nb_lignes = importer_csv(wb, chemin_csv, maintenant())
        ville_pre = demander_ville(wb, saisie, maintenant())
        chemin_carte = generer_carte_html(wb, ville_pre, dossier_sortie, maintenant())
        creer_synthese(wb, chemin_carte, chemin_csv, maintenant())
        chemin_xlsm = sauvegarder_version(wb, dossier_sortie, Path(classeur).suffix, maintenant())
        excel(chemin_xlsm)
        log_msg(wb, "INFO", f"Sauvegarde : {chemin_xlsm}", maintenant())
        envoyer_mail_pipeline(wb, chemin_xlsm, chemin_carte, chemin_csv, mail_to, mail_cc, mode_mail,
                              maintenant(), envoyer_mail)
        if open_map:
            commun.ouvrir_fichier_vba(wb, ouvrir_fichier, chemin_carte, maintenant())
        log_msg(wb, "INFO", f"Pipeline V3 termine : {nb_lignes} ligne(s) valide(s), {chemin_xlsm}", maintenant())
        message = f"Pipeline V3 termine.\n{nb_lignes} ville(s) importee(s)\n{chemin_xlsm}\n{chemin_carte}"
        commun.journal.info(message)
        return Resultat(True, [chemin_xlsm, chemin_carte], message)
    except Exception as erreur:
        description = getattr(erreur, "description", str(erreur))
        log_msg(wb, "ERREUR", f"{commun.numero_erreur(erreur)} - {description}", maintenant())
        commun.journal.error("Erreur : %s", description)
        return Resultat(False, [], f"Erreur : {description}")


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description="VBA UtilisationCSV.LancerPipelineCSVV3")
    p.add_argument("classeur", help="classeur de départ (remplace ThisWorkbook)")
    p.add_argument("--dossier-sortie", default="sortie", help="remplace SaveFolder()")
    p.add_argument("--csv-folder", default=CSV_FOLDER, help="dossier où lire le CSV du jour")
    p.add_argument("--mail-to", default=MAIL_TO)
    p.add_argument("--mail-cc", default=MAIL_CC)
    p.add_argument("--mode-mail", choices=["envoyer", "afficher", "aucun"], default="envoyer",
                   help="remplace SEND_MAIL (True = envoyer, False = afficher)")
    p.add_argument("--open-map", choices=["oui", "non"], default="oui", help="OPEN_MAP : ouvrir la carte")
    p.add_argument("--saisie", default="", help="ville à mettre en avant (InputBox ; vide = toutes)")
    a = p.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    res = lancer_pipeline_csv_v3(a.classeur, a.dossier_sortie, a.csv_folder, a.mail_to, a.mail_cc, a.mode_mail,
                                 a.open_map == "oui", a.saisie)
    return 0 if res.ok else 1


if __name__ == "__main__":
    sys.exit(main())
