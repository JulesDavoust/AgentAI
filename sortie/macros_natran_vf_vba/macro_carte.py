"""Traduction du module VBA ``MacroCarte`` : pipeline gaz avec carte HTML interactive (SVG, hors
ligne), synthèse dans Feuille1, sauvegarde datée, mail (classeur + carte) et ouverture de la carte.
"""

from __future__ import annotations

import argparse
import logging
import random
import sys
from datetime import datetime, time, timedelta
from pathlib import Path

from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Font, PatternFill

import commun
from commun import ErreurVBA, Resultat, log_msg
from connecteurs import excel as connecteur_excel
from bibliotheque.excel_com import preparer_feuille

SHEET_VISUEL = "Feuille1"
SHEET_DATA = "Feuille2"
TABLE_NAME = "tblDonnees"
FILE_TITLE = "VBA_Natran_Test"
FILE_SUFFIX = "VBATF_V2"
MAIL_TO = "jdavoust201@headmind.com"
MAIL_CC = ""
COORDONNEES = {
    "Nord": "50.63,3.06",      # Lille
    "Sud": "43.30,5.37",       # Marseille
    "Est": "48.57,7.75",       # Strasbourg
    "Ouest": "47.22,-1.55",    # Nantes
    "Centre": "47.08,2.40",    # Bourges
}


REGIONS = ("Nord", "Sud", "Est", "Ouest", "Centre")


def coord_region(reg: str) -> str:
    """VBA ``CoordRegion`` : « latitude,longitude » du point d'ancrage de la région."""
    return COORDONNEES.get(reg, "")


def simuler_volume(reg: str, rng) -> float:
    """VBA ``SimulerVolume``"""
    base = {"Nord": 120, "Sud": 95, "Est": 80, "Ouest": 70}.get(reg, 60)
    return round(base * (0.85 + rng.random() * 0.3), 1)


def simuler_temperature(rng) -> float:
    """VBA ``SimulerTemperature``"""
    return round(2 + rng.random() * 14, 1)


def _minuit(d) -> datetime:
    return datetime.combine(d, time())


# ----------------------------------------------------------------------------- 1. structure / données

def ensure_structure(wb, maintenant: datetime, rng) -> bool:
    """VBA ``EnsureStructure`` : crée Feuille2 + tableau tblDonnees (14 jours d'historique) s'il
    n'existe pas. Retourne True si la structure a été créée (AutoFit des colonnes à faire)."""
    ws_d = commun.get_or_create_sheet(wb, SHEET_DATA)
    commun.get_or_create_sheet(wb, SHEET_VISUEL)
    if TABLE_NAME in ws_d.tables:
        return False
    preparer_feuille(ws_d)  # wsD.Cells.Clear
    for i, titre in enumerate(("Date", "Region", "Volume_GWh", "Temperature_C", "Statut"), 1):
        ws_d.cell(row=1, column=i, value=titre)
    ligne = 2
    for i in range(14, 0, -1):
        for r in REGIONS:
            ws_d.cell(row=ligne, column=1, value=_minuit(maintenant.date() - timedelta(days=i)))
            ws_d.cell(row=ligne, column=2, value=r)
            ws_d.cell(row=ligne, column=3, value=simuler_volume(r, rng))
            ws_d.cell(row=ligne, column=4, value=simuler_temperature(rng))
            ws_d.cell(row=ligne, column=5, value="Historique")
            ligne += 1
    commun.creer_table(ws_d, TABLE_NAME, f"A1:E{ligne - 1}")
    commun.formater_colonne(ws_d, "A", "dd/mm/yyyy")
    log_msg(wb, "INFO", "Structure initiale creee", maintenant)
    return True


def ajouter_donnees(wb, maintenant: datetime, rng) -> int:
    """VBA ``AjouterDonnees`` : une ligne (statut « Phase de test V2 ») par région absente du jour."""
    ws = commun.feuille(wb, SHEET_DATA)
    lo = commun.table_de(ws, TABLE_NAME)
    nb = 0
    for r in REGIONS:
        if not ligne_existe(ws, lo, maintenant.date(), r):
            ligne = commun.ajouter_ligne_table(ws, lo, [
                _minuit(maintenant.date()), r, simuler_volume(r, rng), simuler_temperature(rng),
                "Phase de test V2"])
            ws.cell(row=ligne, column=1).number_format = "dd/mm/yyyy"
            nb += 1
    log_msg(wb, "INFO", f"{nb} ligne(s) ajoutee(s) pour le {commun.fmt_date(maintenant)}", maintenant)
    return nb


def ligne_existe(ws, lo, d, reg: str) -> bool:
    """VBA ``LigneExiste``"""
    return any(commun.en_jour(v[0]) == d and v[1] == reg for v in commun.lignes_table(ws, lo))


def est_region_valide(s: str) -> bool:
    """VBA ``EstRegionValide`` (comparaison de texte : insensible à la casse)."""
    return any(r.casefold() == s.casefold() for r in REGIONS)


# ----------------------------------------------------------------------------- 2. filtre

def appliquer_filtre_v2(wb, saisie: str, maintenant: datetime | None = None) -> str | None:
    """VBA ``AppliquerFiltreV2`` (partie classeur) ; le filtre est appliqué par Excel.
    Retourne la région saisie (= ``RegionFiltree`` : ``Replace(Criteria1, "=", "")``) ou None."""
    commun.table_de(commun.feuille(wb, SHEET_DATA), TABLE_NAME)
    saisie = (saisie or "").strip()
    if saisie == "":
        return None
    if not est_region_valide(saisie):
        commun.journal.warning("Region inconnue : %s", saisie)
        return None
    log_msg(wb, "INFO", f"Filtre applique : {saisie}", maintenant)
    return saisie.replace("=", "")


# ----------------------------------------------------------------------------- 3. carte HTML

def generer_carte_html(wb, region_filtree: str | None, dossier, maintenant: datetime) -> Path:
    """VBA ``GenererCarteHTML`` : écrit la carte (Windows-1252, comme ``Print #f``) et retourne son chemin."""
    ws = commun.feuille(wb, SHEET_DATA)
    lignes = commun.lignes_table(ws, commun.table_de(ws, TABLE_NAME))
    if not lignes:
        raise ErreurVBA(commun.VB_OBJECT_ERROR + 1, "Table de donnees vide")

    donnees = ",".join(
        f"['{l[0].strftime('%Y-%m-%d')}','{l[1]}',{commun.str_nb(l[2])},{commun.str_nb(l[3])}]" for l in lignes)
    coords = ",".join(f"'{r}':[{coord_region(r)}]" for r in REGIONS)
    reg_pre = region_filtree or "*"

    h: list[str] = list(commun.HTML_ENTETE)
    h.append("<div id='panel'><b>Distribution gaz</b> Date <select id='selDate'></select>")
    h.append("Region <select id='selReg'></select><button onclick='play()'>Lecture / Pause</button>")
    h.append(f"<span id='tot'></span><small>Genere le {commun.fmt_date_heure(maintenant)} - {FILE_SUFFIX}</small></div>")
    h += commun.HTML_SVG
    h.append("<div id='info'>Cliquer sur une region pour le detail.</div><script>")
    h.append(f"var DATA=[{donnees}];")
    h.append(f"var COORD={{{coords}}};")
    h += commun.HTML_CONTOUR
    h += [
        "var dates=[],regs=[];",
        "DATA.forEach(function(r){if(dates.indexOf(r[0])<0)dates.push(r[0]);if(regs.indexOf(r[1])<0)regs.push(r[1]);});",
        "dates.sort();",
        "var selD=document.getElementById('selDate'),selR=document.getElementById('selReg');",
        "dates.forEach(function(d){selD.add(new Option(d,d));});",
        "selR.add(new Option('Toutes','*'));regs.forEach(function(r){selR.add(new Option(r,r));});",
        f"selD.value=dates[dates.length-1];selR.value='{reg_pre}';",
        "var MAXV=Math.max.apply(null,DATA.map(function(r){return r[2];}));",
        "function color(v){var x=v/MAXV;return x>0.85?'#b30000':x>0.65?'#e34a33':x>0.45?'#fc8d59':'#fdcc8a';}",
        "function el(n,a){var e=document.createElementNS(NS,n);for(var k in a)e.setAttribute(k,a[k]);return e;}",
        "function draw(){while(g.firstChild)g.removeChild(g.firstChild);var d=selD.value,f=selR.value,tot=0;",
        "DATA.forEach(function(r){if(r[0]!==d)return;if(f!=='*'&&r[1]!==f)return;",
        "var c=COORD[r[1]];if(!c)return;tot+=r[2];",
        "var rad=12+34*r[2]/MAXV,cx=X(c[1]),cy=Y(c[0]);",
        "var ci=el('circle',{cx:cx,cy:cy,r:rad,fill:color(r[2]),'fill-opacity':0.8,stroke:'#333','stroke-width':1});",
        "var tt=el('title',{});tt.textContent=r[1]+' : '+r[2].toFixed(1)+' GWh';ci.appendChild(tt);",
        "ci.onclick=function(){document.getElementById('info').innerText=r[1]+' - '+d+' : '+r[2].toFixed(1)+' GWh, temperature '+r[3].toFixed(1)+' C';};",
        "g.appendChild(ci);",
        "var t=el('text',{x:cx,y:cy+4,'text-anchor':'middle','font-size':12,'font-weight':'bold',fill:'#222','pointer-events':'none'});",
        "t.textContent=r[1];g.appendChild(t);});",
        "document.getElementById('tot').innerText='Total affiche : '+tot.toFixed(1)+' GWh';}",
        "selD.onchange=draw;selR.onchange=draw;var timer=null;",
        "function play(){if(timer){clearInterval(timer);timer=null;return;}",
        "var i=dates.indexOf(selD.value);",
        "timer=setInterval(function(){i=(i+1)%dates.length;selD.value=dates[i];draw();},800);}",
        "draw();</script></body></html>",
    ]
    texte = "".join(ligne + "\n" for ligne in h)

    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    chemin = (dossier / f"{FILE_TITLE}_Carte_{maintenant.strftime('%Y%m%d_%H%M%S')}_{FILE_SUFFIX}.html").resolve()
    commun.ecrire_texte_cp1252(chemin, texte)
    log_msg(wb, "INFO", f"Carte generee : {chemin}", maintenant)
    return chemin


# ----------------------------------------------------------------------------- 4. synthèse

def creer_synthese(wb, chemin_carte, maintenant: datetime) -> None:
    """VBA ``CreerSynthese`` (partie openpyxl) : KPI, tableau du dernier jour, lien vers la carte.
    L'AutoFit et la largeur de la colonne A sont faits par Excel."""
    ws = commun.feuille(wb, SHEET_VISUEL)
    preparer_feuille(ws)  # ChartObjects.Delete, PivotTables.Clear, Cells.Clear
    for row in ws.iter_rows():
        for c in row:
            if c.hyperlink:
                c.hyperlink = None  # Hyperlinks.Delete
    ws.conditional_formatting = ConditionalFormattingList()

    ws["A1"] = "Suivi distribution gaz (V2 carte) - mis a jour le " + commun.fmt_date_heure(maintenant)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Volume total (GWh)"
    ws["B2"] = f"=SUM({TABLE_NAME}[Volume_GWh])"
    ws["D2"] = "Moyenne / jour (GWh)"
    ws["E2"] = f"=SUM({TABLE_NAME}[Volume_GWh])/SUMPRODUCT(1/COUNTIF({TABLE_NAME}[Date],{TABLE_NAME}[Date]))"
    ws["G2"] = "Region du pic"
    ws["H2"] = f"=INDEX({TABLE_NAME}[Region],MATCH(MAX({TABLE_NAME}[Volume_GWh]),{TABLE_NAME}[Volume_GWh],0))"
    for ref in ("A2", "D2", "G2"):
        ws[ref].font = Font(bold=True)
    for ref in ("B2", "E2"):
        ws[ref].number_format = "#,##0.0"

    ws["A3"] = "Dernier jour"
    ws["B3"] = f"=MAX({TABLE_NAME}[Date])"
    ws["B3"].number_format = "dd/mm/yyyy"
    for i, titre in enumerate(("Region", "Volume du jour (GWh)", "Temp. moyenne (C)"), 1):
        c = ws.cell(row=4, column=i, value=titre)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")

    ligne = 5
    for r in REGIONS:
        ws.cell(row=ligne, column=1, value=r)
        ws.cell(row=ligne, column=2, value=(
            f"=SUMIFS({TABLE_NAME}[Volume_GWh],{TABLE_NAME}[Region],A{ligne},{TABLE_NAME}[Date],$B$3)"))
        ws.cell(row=ligne, column=3, value=(
            f'=IFERROR(AVERAGEIFS({TABLE_NAME}[Temperature_C],{TABLE_NAME}[Region],A{ligne},'
            f'{TABLE_NAME}[Date],$B$3),"")'))
        ligne += 1
    for l in range(5, ligne):
        ws.cell(row=l, column=2).number_format = "#,##0.0"
        ws.cell(row=l, column=3).number_format = "#,##0.0"
    ws.conditional_formatting.add(f"B5:B{ligne - 1}", DataBarRule(start_type="min", end_type="max", color="638EC6"))

    lien = ws.cell(row=ligne + 1, column=1)
    lien.hyperlink = str(chemin_carte)
    lien.value = "Ouvrir la carte interactive"
    lien.style = "Hyperlink"
    ws.cell(row=ligne + 2, column=1, value=f"Fichier : {chemin_carte}").font = Font(italic=True)

    commun.activer_feuille(wb, ws)
    log_msg(wb, "INFO", "Synthese regeneree", maintenant)


# ----------------------------------------------------------------------------- 5. sauvegarde / 6. mail

def sauvegarder_version(wb, dossier, extension: str, maintenant: datetime) -> Path:
    """VBA ``SauvegarderVersion`` : copie ``titre_date_VBATF_V2`` dans le dossier (créé si besoin)."""
    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    chemin = (dossier / commun.nom_version(FILE_TITLE, maintenant, FILE_SUFFIX, extension)).resolve()
    wb.save(chemin)
    return chemin


def envoyer_mail_pipeline(wb, chemin_xlsm, chemin_carte, mail_to, mail_cc, mode_mail, maintenant,
                          envoyer_mail=None) -> bool:
    """VBA ``EnvoyerMail`` (MacroCarte) : classeur et carte en pièces jointes."""
    sujet = "[PHASE DE TEST V2] Mise a jour suivi distribution - " + commun.fmt_date(maintenant)
    corps = ("Bonjour,\r\n\r\n"
             "Ceci est un envoi automatique en phase de test (version carte interactive).\r\n"
             "Les donnees et la carte ont ete mises a jour.\r\n"
             f"Classeur : {Path(chemin_xlsm).name}\r\n"
             f"Carte : {Path(chemin_carte).name} (a ouvrir dans un navigateur, fonctionne hors ligne)\r\n\r\n"
             "Cordialement,\r\nRobot VBATF")
    return commun.expedier_mail(wb, envoyer_mail, mode_mail, mail_to, mail_cc, sujet, corps,
                                [chemin_xlsm, chemin_carte], maintenant)


# ----------------------------------------------------------------------------- point d'entrée

def lancer_pipeline_carte_v2(classeur, dossier_sortie="sortie", mail_to=MAIL_TO, mail_cc=MAIL_CC,
                             mode_mail="envoyer", open_map=True, saisie="", *, envoyer_mail=None, excel=None,
                             ouvrir_fichier=None, rng=None, maintenant=datetime.now) -> Resultat:
    """VBA ``LancerPipelineCarteV2``. ``classeur`` = ``ThisWorkbook`` (non modifié : le résultat est la
    copie enregistrée) ; ``open_map`` = ``OPEN_MAP`` ; ``saisie`` = InputBox de ``AppliquerFiltreV2``."""
    rng = rng or random.Random()
    excel = excel or connecteur_excel.finaliser_carte
    wb = commun.charger_classeur(classeur)
    try:
        creee = ensure_structure(wb, maintenant(), rng)
        ajouter_donnees(wb, maintenant(), rng)
        region = appliquer_filtre_v2(wb, saisie, maintenant())
        chemin_carte = generer_carte_html(wb, region, dossier_sortie, maintenant())
        creer_synthese(wb, chemin_carte, maintenant())
        chemin_xlsm = sauvegarder_version(wb, dossier_sortie, Path(classeur).suffix, maintenant())
        excel(chemin_xlsm, creee, region)
        log_msg(wb, "INFO", f"Sauvegarde : {chemin_xlsm}", maintenant())
        envoyer_mail_pipeline(wb, chemin_xlsm, chemin_carte, mail_to, mail_cc, mode_mail, maintenant(), envoyer_mail)
        if open_map:
            commun.ouvrir_fichier_vba(wb, ouvrir_fichier, chemin_carte, maintenant())
        log_msg(wb, "INFO", f"Pipeline V2 termine : {chemin_xlsm}", maintenant())
        message = f"Pipeline V2 termine.\n{chemin_xlsm}\n{chemin_carte}"
        commun.journal.info(message)
        return Resultat(True, [chemin_xlsm, chemin_carte], message)
    except Exception as erreur:
        description = getattr(erreur, "description", str(erreur))
        log_msg(wb, "ERREUR", f"{commun.numero_erreur(erreur)} - {description}", maintenant())
        commun.journal.error("Erreur : %s", description)
        return Resultat(False, [], f"Erreur : {description}")


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description="VBA MacroCarte.LancerPipelineCarteV2")
    p.add_argument("classeur", help="classeur de départ (remplace ThisWorkbook)")
    p.add_argument("--dossier-sortie", default="sortie", help="remplace SaveFolder()")
    p.add_argument("--mail-to", default=MAIL_TO)
    p.add_argument("--mail-cc", default=MAIL_CC)
    p.add_argument("--mode-mail", choices=["envoyer", "afficher", "aucun"], default="envoyer",
                   help="remplace SEND_MAIL (True = envoyer, False = afficher)")
    p.add_argument("--open-map", choices=["oui", "non"], default="oui", help="OPEN_MAP : ouvrir la carte")
    p.add_argument("--saisie", default="", help="région à mettre en avant (InputBox ; vide = toutes)")
    a = p.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    res = lancer_pipeline_carte_v2(a.classeur, a.dossier_sortie, a.mail_to, a.mail_cc, a.mode_mail,
                                   a.open_map == "oui", a.saisie)
    return 0 if res.ok else 1


if __name__ == "__main__":
    sys.exit(main())
