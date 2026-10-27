"""Outils partagés par les modules traduits (aucun équivalent VBA direct).

Regroupe ce que le VBA recopie à l'identique dans plusieurs modules (``LogMsg``,
``GetOrCreateSheet``, ``Num``, blocs HTML communs des cartes, envoi du mail, ``OuvrirFichier``)
et les fonctions qui reproduisent les comportements de VBA (``Val``, ``IsNumeric``, ``CDate``,
``Str$``, ``DateSerial``, ``End(xlUp)``).
"""

from __future__ import annotations

import logging
import re
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableStyleInfo

from connecteurs import outlook, systeme

journal = logging.getLogger("vba2py")

VB_OBJECT_ERROR = -2147221504  # vbObjectError


class ErreurVBA(Exception):
    """Équivalent de ``Err.Raise`` : ``numero`` = Err.Number, ``description`` = Err.Description."""

    def __init__(self, numero: int, description: str):
        super().__init__(description)
        self.numero = numero
        self.description = description


@dataclass
class Resultat:
    """Ce que la macro annonce par ``MsgBox`` : succès, fichiers produits, message."""

    ok: bool = True
    chemins: list = field(default_factory=list)
    message: str = ""


def numero_erreur(erreur: Exception) -> int:
    """``Err.Number`` : celui de l'ErreurVBA, 0 pour une erreur Python quelconque."""
    return erreur.numero if isinstance(erreur, ErreurVBA) else 0


# ----------------------------------------------------------------------------- classeur

def charger_classeur(chemin) -> openpyxl.Workbook:
    """``ThisWorkbook`` : le classeur passé en paramètre (macros conservées pour un .xlsm)."""
    chemin = Path(chemin)
    return openpyxl.load_workbook(chemin, keep_vba=chemin.suffix.lower() == ".xlsm")


def get_or_create_sheet(wb, nom: str):
    """VBA ``GetOrCreateSheet`` : la feuille, créée après la dernière si elle n'existe pas."""
    if nom in wb.sheetnames:
        return wb[nom]
    return wb.create_sheet(nom)


def feuille(wb, nom: str):
    """``ThisWorkbook.Worksheets(nom)`` : erreur 9 si la feuille n'existe pas."""
    if nom not in wb.sheetnames:
        raise ErreurVBA(9, "L'indice n'appartient pas à la sélection")
    return wb[nom]


def table_de(ws, nom: str):
    """``ws.ListObjects(nom)`` : erreur 9 si le tableau n'existe pas."""
    if nom not in ws.tables:
        raise ErreurVBA(9, "L'indice n'appartient pas à la sélection")
    return ws.tables[nom]


def derniere_ligne(ws, colonne: int = 1) -> int:
    """``Cells(Rows.Count, col).End(xlUp).Row`` : dernière ligne non vide (1 si la colonne est vide)."""
    for r in range(ws.max_row, 0, -1):
        if ws.cell(row=r, column=colonne).value is not None:
            return r
    return 1


def log_msg(wb, niveau: str, msg: str, maintenant: datetime | None = None) -> None:
    """VBA ``LogMsg`` (MacroInitiale, MacroCarte, UtilisationCSV) : une ligne dans la feuille Log."""
    ws = get_or_create_sheet(wb, "Log")
    if ws["A1"].value in (None, ""):
        for i, titre in enumerate(("Horodatage", "Niveau", "Message"), 1):
            ws.cell(row=1, column=i, value=titre)
    n = derniere_ligne(ws, 1) + 1
    cellule = ws.cell(row=n, column=1, value=maintenant or datetime.now())
    cellule.number_format = "dd/mm/yyyy hh:mm:ss"
    ws.cell(row=n, column=2, value=niveau)
    ws.cell(row=n, column=3, value=msg)
    journal.log(logging.ERROR if niveau == "ERREUR" else logging.INFO, "[%s] %s", niveau, msg)


def activer_feuille(wb, ws) -> None:
    """``ws.Activate`` : feuille active à l'enregistrement."""
    wb.active = ws
    for s in wb.worksheets:
        s.sheet_view.tabSelected = s is ws


def selectionner(ws, plage: str) -> None:
    """``Range(...).Select`` : sélection enregistrée dans la feuille."""
    sel = ws.sheet_view.selection[0]
    sel.activeCell = plage.split(":")[0]
    sel.sqref = plage


# ----------------------------------------------------------------------------- tableaux

def creer_table(ws, nom: str, ref: str, style: str = "TableStyleMedium2"):
    """``ListObjects.Add`` + ``.Name`` + ``.TableStyle``."""
    t = Table(displayName=nom, name=nom, ref=ref)
    t.autoFilter = AutoFilter(ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)
    return t


def lignes_table(ws, table) -> list[list]:
    """``ListObject.DataBodyRange.Value`` : lignes de données (liste vide = DataBodyRange Nothing)."""
    c0, r0, c1, r1 = range_boundaries(table.ref)
    return [[ws.cell(row=r, column=c).value for c in range(c0, c1 + 1)] for r in range(r0 + 1, r1 + 1)]


def ajouter_ligne_table(ws, table, valeurs) -> int:
    """``ListRows.Add`` + écriture des cellules : agrandit le tableau d'une ligne."""
    c0, r0, c1, r1 = range_boundaries(table.ref)
    r = r1 + 1
    for i, v in enumerate(valeurs):
        cellule = ws.cell(row=r, column=c0 + i)
        if r - 1 > r0:
            cellule._style = copy(ws.cell(row=r - 1, column=c0 + i)._style)
        cellule.value = v
    ref = f"{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{r}"
    table.ref = ref
    if table.autoFilter is not None:
        table.autoFilter.ref = ref
    return r


def formater_colonne(ws, lettre: str, format_nombre: str) -> None:
    """``Columns("A").NumberFormat = ...``"""
    ws.column_dimensions[lettre].number_format = format_nombre
    for row in ws.iter_rows(min_col=openpyxl.utils.column_index_from_string(lettre),
                            max_col=openpyxl.utils.column_index_from_string(lettre)):
        for c in row:
            c.number_format = format_nombre


def en_jour(v) -> date:
    """``CLng(cellule)`` sur une date : le jour (erreur 13 si ce n'est pas une date)."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return date(1899, 12, 30) + timedelta(days=round(v))
    raise ErreurVBA(13, "Incompatibilité de type")


# ----------------------------------------------------------------------------- fonctions VBA

_VAL = re.compile(r"[+-]?(\d+\.?\d*|\.\d+)([eEdD][+-]?\d+)?")
_NUM = re.compile(r"[+-]?(\d+(,\d*)?|,\d+)([eEdD][+-]?\d+)?")
_BASE = re.compile(r"&[hH][0-9A-Fa-f]+&?|&[oO][0-7]+&?")


def vba_val(texte) -> float:
    """VBA ``Val`` : point décimal, espaces ignorés, arrêt au premier caractère invalide."""
    s = re.sub(r"\s", "", str(texte))
    m = _VAL.match(s)
    return float(m.group().replace("d", "e").replace("D", "e")) if m else 0.0


def is_numeric(v) -> bool:
    """VBA ``IsNumeric`` en réglages français (séparateur décimal ``,``)."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return True
    s = str(v).strip()
    return bool(_NUM.fullmatch(s) or _BASE.fullmatch(s))


def str_nb(v) -> str:
    """VBA ``Trim$(Str$(CDbl(v)))`` : nombre avec point décimal, sans espace."""
    t = format(float(v), ".15g")
    if "e" in t:
        mantisse, exposant = t.split("e")
        t = f"{mantisse}E{exposant[0]}{abs(int(exposant)):02d}"
    return t


def cstr(v) -> str:
    """VBA ``CStr`` d'une valeur de cellule (réglages français)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Vrai" if v else "Faux"
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v).replace(".", ",")
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y") if v.time() == datetime.min.time() else v.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return str(v)


def c_int(texte) -> int:
    """VBA ``CInt`` d'un texte (arrondi bancaire, virgule décimale)."""
    try:
        return int(round(float(str(texte).strip().replace(",", "."))))
    except ValueError:
        raise ErreurVBA(13, "Incompatibilité de type") from None


def date_serial(annee: int, mois: int, jour: int) -> datetime:
    """VBA ``DateSerial`` (les mois et jours hors limites débordent sur l'année / le mois suivant)."""
    if 0 <= annee <= 99:
        annee += 2000 if annee < 30 else 1900
    total = annee * 12 + (mois - 1)
    base = date(total // 12, total % 12 + 1, 1)
    d = base + timedelta(days=jour - 1)
    return datetime(d.year, d.month, d.day)


def c_date(texte: str) -> date:
    """VBA ``CDate`` en réglages français : ``05/10/2026`` est le 5 octobre (jj/mm/aaaa, aaaa-mm-jj)."""
    s = str(texte).strip()
    m = re.fullmatch(r"(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2}|\d{4})", s)
    if m:
        j, mo, a = (int(x) for x in m.groups())
    else:
        m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
        if not m:
            raise ErreurVBA(13, "Incompatibilité de type")
        a, mo, j = (int(x) for x in m.groups())
    if a < 100:
        a += 2000 if a < 30 else 1900
    try:
        return date(a, mo, j)
    except ValueError:
        raise ErreurVBA(13, "Incompatibilité de type") from None


def is_date(texte: str) -> bool:
    """VBA ``IsDate`` (mêmes formats que ``c_date``)."""
    try:
        c_date(texte)
        return True
    except ErreurVBA:
        return False


def fmt_date(d) -> str:
    """``Format(d, "dd/mm/yyyy")``"""
    return d.strftime("%d/%m/%Y")


def fmt_date_heure(d) -> str:
    """``Format(d, "dd/mm/yyyy hh:nn")``"""
    return d.strftime("%d/%m/%Y %H:%M")


def nom_version(titre: str, maintenant: datetime, suffixe: str, extension: str) -> str:
    """``titre_aaaammjj_hhnnss_suffixe.ext`` (nom des sauvegardes)."""
    return f"{titre}_{maintenant.strftime('%Y%m%d_%H%M%S')}_{suffixe}{extension}"


# ----------------------------------------------------------------------------- fichiers

def ecrire_texte_cp1252(chemin, texte: str) -> None:
    """``Open ... For Output`` + ``Print #f, texte`` : Windows-1252, le Print final ajoute CRLF
    (les ``vbLf`` internes restent des LF)."""
    with open(chemin, "w", encoding="cp1252", errors="replace", newline="") as f:
        f.write(texte + "\r\n")


# ----------------------------------------------------------------------------- mail / ouverture

def expedier_mail(wb, envoyer_mail, mode: str, destinataire: str, cc: str, sujet: str, corps: str,
                  pieces_jointes, maintenant: datetime | None = None) -> bool:
    """Bloc ``EnvoyerMail`` commun : ``On Error GoTo MailErr`` consigne l'erreur dans le Log sans
    interrompre la macro. ``mode`` : envoyer (.Send), afficher (.Display), aucun (rien)."""
    if mode == "aucun":
        log_msg(wb, "INFO", "Mail non envoye (mode-mail aucun)", maintenant)
        return False
    try:
        fonction = envoyer_mail or outlook.envoyer_mail
        fonction(destinataire, cc, sujet, corps, [str(Path(p).resolve()) for p in pieces_jointes], mode)
        log_msg(wb, "INFO", ("Mail envoye a " if mode == "envoyer" else "Brouillon cree pour ") + destinataire,
                maintenant)
        return True
    except Exception as erreur:  # MailErr : la macro d'origine consigne et continue
        log_msg(wb, "ERREUR", f"Mail impossible : {erreur}", maintenant)
        return False


def ouvrir_fichier_vba(wb, ouvrir, chemin, maintenant: datetime | None = None) -> None:
    """VBA ``OuvrirFichier`` (``On Error Resume Next`` + ``FollowHyperlink``)."""
    try:
        (ouvrir or systeme.ouvrir_fichier)(str(chemin))
    except Exception as erreur:
        log_msg(wb, "ERREUR", f"Ouverture impossible : {erreur}", maintenant)


# ----------------------------------------------------------------------------- HTML des cartes

HTML_ENTETE = [
    "<!DOCTYPE html><html><head><meta charset='utf-8'>",
    "<title>Suivi distribution gaz - carte</title>",
    "<style>html,body{height:100%;margin:0;font-family:Arial,sans-serif;background:#fff}",
    "#panel{padding:8px;background:#f2f2f2;border-bottom:1px solid #ccc}",
    "#panel *{margin-right:8px}#svg{width:100%;height:calc(100% - 78px);display:block}",
    "#info{padding:6px 8px;height:26px;font-size:14px;border-top:1px solid #ccc;background:#fafafa}",
    "circle{cursor:pointer}circle:hover{stroke-width:3}</style></head><body>",
]

HTML_SVG = [
    "<svg id='svg' viewBox='0 0 640 640' preserveAspectRatio='xMidYMid meet'>",
    "<path id='fr' fill='#eaf0e0' stroke='#6b7a5e' stroke-width='1.5' stroke-linejoin='round'/>",
    "<path id='co' fill='#eaf0e0' stroke='#6b7a5e' stroke-width='1.5' stroke-linejoin='round'/>",
    "<g id='g'></g></svg>",
]

# Contour simplifié de la France métropolitaine [longitude, latitude] puis fonctions de tracé
HTML_CONTOUR = [
    "var FR=[[2.55,51.09],[1.85,50.95],[1.60,50.73],[1.55,50.35],[1.38,50.06],[1.08,49.93],[0.37,49.76],[0.10,49.50],",
    "[0.25,49.43],[0.07,49.35],[-0.25,49.30],[-0.62,49.34],[-1.10,49.40],[-1.27,49.69],[-1.95,49.72],[-1.80,49.38],",
    "[-1.60,48.83],[-1.50,48.64],[-1.85,48.68],[-2.02,48.65],[-2.32,48.69],[-2.75,48.55],[-3.05,48.78],[-3.45,48.83],",
    "[-3.98,48.73],[-4.77,48.33],[-4.73,48.04],[-4.37,47.80],[-3.92,47.87],[-3.37,47.73],[-3.12,47.48],[-2.75,47.55],",
    "[-2.45,47.45],[-2.52,47.29],[-2.20,47.27],[-2.10,47.11],[-2.20,46.95],[-1.93,46.68],[-1.78,46.49],[-1.15,46.15],",
    "[-1.25,45.95],[-1.05,45.62],[-1.07,45.57],[-1.20,45.00],[-1.25,44.65],[-1.29,44.20],[-1.45,43.65],[-1.56,43.48],",
    "[-1.78,43.37],[-1.45,43.05],[-0.75,42.95],[-0.30,42.85],[0.20,42.75],[0.67,42.85],[1.00,42.80],[1.45,42.60],",
    "[1.75,42.50],[2.00,42.35],[2.50,42.35],[3.05,42.43],[3.03,42.70],[3.05,43.05],[3.70,43.40],[4.20,43.50],",
    "[4.80,43.35],[5.35,43.30],[5.93,43.10],[6.75,43.40],[7.00,43.55],[7.27,43.70],[7.53,43.78],[7.60,44.15],",
    "[6.95,44.40],[7.05,44.65],[6.75,45.00],[7.05,45.25],[6.80,45.55],[6.85,45.95],[6.80,46.38],[6.25,46.30],",
    "[6.05,46.20],[6.30,46.55],[6.40,46.80],[6.95,47.20],[7.15,47.50],[7.55,47.58],[7.55,48.00],[7.80,48.60],",
    "[8.20,48.98],[7.40,49.15],[6.35,49.45],[5.85,49.55],[5.45,49.50],[5.05,49.65],[4.85,49.95],[4.20,50.25],",
    "[3.60,50.35],[3.25,50.75],[2.60,50.85]];",
    "var CORSE=[[9.15,41.38],[9.40,41.60],[9.55,42.10],[9.45,42.70],[9.36,43.00],[9.30,42.70],[8.75,42.57],",
    "[8.60,42.28],[8.72,41.92],[8.90,41.67]];",
    "var K=60,C=0.688;",
    "function X(lon){return (lon+5.2)*C*K;}function Y(lat){return (51.4-lat)*K;}",
    "function path(a){return a.map(function(p,i){return (i?'L':'M')+X(p[0]).toFixed(1)+' '+Y(p[1]).toFixed(1);}).join('')+'Z';}",
    "document.getElementById('fr').setAttribute('d',path(FR));",
    "document.getElementById('co').setAttribute('d',path(CORSE));",
    "var NS='http://www.w3.org/2000/svg',g=document.getElementById('g');",
]
