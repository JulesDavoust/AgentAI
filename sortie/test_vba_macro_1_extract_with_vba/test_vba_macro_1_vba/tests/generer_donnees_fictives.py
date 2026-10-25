"""Crée les classeurs fictifs de tests/fixtures (idempotent).

- classeur_vide.xlsx   : Feuil1..Feuil4 vides (point de départ : la macro crée tout).
- classeur_donnees.xlsx : Feuille2 avec le tableau tblDonnees (3 jours x 5 régions,
  du 20/10/2025 au 22/10/2025), Feuille1 vide.
Utilisable dans Excel pour comparer avec la macro VBA d'origine.
"""
import argparse
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

REGIONS = ["Nord", "Sud", "Est", "Ouest", "Centre"]
BASES = {"Nord": 120.0, "Sud": 95.0, "Est": 80.0, "Ouest": 70.0, "Centre": 60.0}


def creer_classeur_vide(chemin: Path) -> None:
    wb = Workbook()
    wb.active.title = "Feuil1"
    for i in range(2, 5):
        wb.create_sheet(f"Feuil{i}")
    wb.save(chemin)


def creer_classeur_donnees(chemin: Path) -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Feuille1"
    ws = wb.create_sheet("Feuille2")
    ws.append(["Date", "Region", "Volume_GWh", "Temperature_C", "Statut"])
    for jour in (20, 21, 22):
        for k, reg in enumerate(REGIONS):
            # volume = base + jour - 20 (valeurs lisibles) ; température = 5 + k
            ws.append([datetime(2025, 10, jour), reg, BASES[reg] + (jour - 20), 5.0 + k, "Historique"])
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=1).number_format = "dd/mm/yyyy"
    tbl = Table(displayName="tblDonnees", ref=f"A1:E{ws.max_row}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)
    wb.save(chemin)


def generer(dossier: Path) -> None:
    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    creer_classeur_vide(dossier / "classeur_vide.xlsx")
    creer_classeur_donnees(dossier / "classeur_donnees.xlsx")


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--dossier", type=Path, default=Path(__file__).parent / "fixtures")
    generer(p.parse_args().dossier)
