import random
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

import marco_25_10 as m
from connecteurs import outlook

# Référence capturée à l'import, avant que la fixture autouse ne remplace la fonction.
ENVOYER_MAIL_REEL = outlook.envoyer_mail

JOUR = date(2025, 10, 23)
MAINT = datetime(2025, 10, 23, 14, 5, 9)


class FauxMail:
    def __init__(self, erreur=None):
        self.appels = []
        self.erreur = erreur

    def __call__(self, destinataire, cc, sujet, corps, pieces_jointes, mode):
        if self.erreur:
            raise self.erreur
        self.appels.append(dict(destinataire=destinataire, cc=cc, sujet=sujet, corps=corps,
                                pj=list(pieces_jointes), mode=mode))


def lignes(wb):
    return m.lignes_table(wb)


# --- simulation --------------------------------------------------------------
def test_simuler_volume_bornes():
    rng = random.Random(1)
    for reg, base in [("Nord", 120), ("Sud", 95), ("Est", 80), ("Ouest", 70), ("Centre", 60), ("Autre", 60)]:
        for _ in range(50):
            v = m.simuler_volume(reg, rng)
            assert base * 0.85 - 0.05 <= v <= base * 1.15 + 0.05


def test_simuler_temperature_bornes():
    rng = random.Random(1)
    assert all(2 <= m.simuler_temperature(rng) <= 16 for _ in range(100))


# --- structure ---------------------------------------------------------------
def test_ensure_structure_cree_historique(classeur_vide):
    wb = load_workbook(classeur_vide)
    assert m.ensure_structure(wb, JOUR, random.Random(0)) is True
    assert {"Feuille1", "Feuille2", "Log"} <= set(wb.sheetnames)
    data = lignes(wb)
    assert len(data) == 14 * 5  # 14 jours x 5 régions
    assert {d[4] for d in data} == {"Historique"}
    assert data[0][0] == datetime(2025, 10, 9)  # Date - 14
    assert data[-1][0] == datetime(2025, 10, 22)  # Date - 1
    assert [c.value for c in wb["Feuille2"][1]] == m.EN_TETES


def test_ensure_structure_idempotent(classeur_donnees):
    wb = load_workbook(classeur_donnees)
    assert m.ensure_structure(wb, JOUR) is False
    assert len(lignes(wb)) == 15


# --- données du jour ---------------------------------------------------------
def test_ajouter_donnees_du_jour(classeur_donnees):
    wb = load_workbook(classeur_donnees)
    assert m.ajouter_donnees_du_jour(wb, JOUR, random.Random(0)) == 5
    data = lignes(wb)
    assert len(data) == 20
    assert {d[4] for d in data[-5:]} == {"Phase de test"}
    assert wb["Feuille2"].tables["tblDonnees"].ref == "A1:E21"


def test_ajouter_donnees_pas_de_doublon(classeur_donnees):
    wb = load_workbook(classeur_donnees)
    m.ajouter_donnees_du_jour(wb, JOUR, random.Random(0))
    assert m.ajouter_donnees_du_jour(wb, JOUR, random.Random(0)) == 0
    assert len(lignes(wb)) == 20


def test_ajouter_donnees_jour_deja_present(classeur_donnees):
    wb = load_workbook(classeur_donnees)
    assert m.ajouter_donnees_du_jour(wb, date(2025, 10, 22)) == 0  # déjà 5 lignes ce jour-là


def test_ligne_existe(classeur_donnees):
    wb = load_workbook(classeur_donnees)
    assert m.ligne_existe(wb, date(2025, 10, 20), "Nord")
    assert not m.ligne_existe(wb, date(2025, 10, 20), "nord")  # '=' VBA sensible à la casse
    assert not m.ligne_existe(wb, JOUR, "Nord")


def test_table_absente_erreur(classeur_vide):
    wb = load_workbook(classeur_vide)
    with pytest.raises(KeyError):
        m.ajouter_donnees_du_jour(wb, JOUR)


# --- filtre ------------------------------------------------------------------
def test_est_region_valide():
    assert m.est_region_valide("nord") and m.est_region_valide("OUEST")
    assert not m.est_region_valide("Paris") and not m.est_region_valide("")


def test_filtre_region(classeur_donnees):
    wb = load_workbook(classeur_donnees)
    assert m.appliquer_filtre(wb, "  est ") == "est"
    ws = wb["Feuille2"]
    visibles = [r for r in range(2, 17) if not ws.row_dimensions[r].hidden]
    assert len(visibles) == 3 and all(ws.cell(row=r, column=2).value == "Est" for r in visibles)
    assert ws.tables["tblDonnees"].autoFilter.filterColumn[0].filters.filter == ["est"]


@pytest.mark.parametrize("saisie", ["", None, "   ", "Paris"])
def test_filtre_aucun(classeur_donnees, saisie):
    wb = load_workbook(classeur_donnees)
    m.appliquer_filtre(wb, "Nord")
    assert m.appliquer_filtre(wb, saisie) is None  # ShowAllData d'abord
    ws = wb["Feuille2"]
    assert not any(ws.row_dimensions[r].hidden for r in range(2, 17))


# --- visuel ------------------------------------------------------------------
def test_tcd(classeur_donnees):
    wb = load_workbook(classeur_donnees)
    tcd = m.construire_tcd(wb)
    assert tcd.shape == (3, 5)
    # Nord : 120, 121, 122 ; total du 20/10 = 120+95+80+70+60 = 425
    assert list(tcd["Nord"]) == [120, 121, 122]
    assert tcd.iloc[0].sum() == 425


def test_creer_visuel(classeur_donnees, tmp_path):
    wb = load_workbook(classeur_donnees)
    png = m.creer_visuel(wb, tmp_path / "png", MAINT)
    assert png.exists() and png.stat().st_size > 0
    ws = wb["Feuille1"]
    assert ws["A1"].value == "Suivi distribution gaz - mis a jour le 23/10/2025 14:05"
    assert ws["B2"].value == "=SUM(tblDonnees[Volume_GWh])"
    assert ws["E2"].value.startswith("=SUM(tblDonnees[Volume_GWh])/SUMPRODUCT(1/COUNTIF(")
    assert ws["H2"].value.startswith("=INDEX(tblDonnees[Region],MATCH(MAX(")
    assert ws["A6"].value == "Date" and ws["B6"].value == "Centre"  # tri alphabétique du TCD
    assert ws["A7"].value == datetime(2025, 10, 20)
    assert len(ws._images) == 1
    # Régénération : pas d'accumulation d'images
    m.creer_visuel(wb, tmp_path / "png", MAINT)
    assert len(wb["Feuille1"]._images) == 1


def test_creer_visuel_tcd_hors_filtre(classeur_donnees, tmp_path):
    wb = load_workbook(classeur_donnees)
    m.appliquer_filtre(wb, "Nord")
    m.creer_visuel(wb, tmp_path, MAINT)
    ws = wb["Feuille1"]
    # le TCD porte sur toute la table : 5 colonnes de régions
    assert [ws.cell(row=6, column=c).value for c in range(2, 8)] == [
        "Centre", "Est", "Nord", "Ouest", "Sud", "Total general"]


# --- sauvegarde --------------------------------------------------------------
def test_sauvegarder_version(classeur_donnees, tmp_path):
    wb = load_workbook(classeur_donnees)
    chemin = m.sauvegarder_version(wb, tmp_path / "nouveau" / "dossier", ".xlsx", MAINT)
    assert chemin.name == "VBA_Natran_Test_20251023_140509_VBATF.xlsx"
    assert chemin.exists()


# --- mail --------------------------------------------------------------------
def test_corps_mail():
    c = m.construire_corps_mail(r"C:\x\y\fichier.xlsm")
    assert "Fichier : " in c and c.endswith("Robot VBATF")
    assert c.startswith("Bonjour,\r\n\r\n")


def test_mail_envoyer(classeur_donnees, tmp_path):
    wb = load_workbook(classeur_donnees)
    faux = FauxMail()
    assert m.envoyer_mail_pipeline(wb, tmp_path / "f.xlsm", envoyer_mail=faux, aujourdhui=JOUR)
    a = faux.appels[0]
    assert a["destinataire"] == "jdavoust201@headmind.com" and a["cc"] == "" and a["mode"] == "envoyer"
    assert a["sujet"] == "[PHASE DE TEST] Mise a jour suivi distribution - 23/10/2025"
    assert "Fichier : f.xlsm\r\n" in a["corps"]
    assert a["pj"] == [tmp_path / "f.xlsm"]


def test_mail_afficher_et_destinataire(classeur_donnees, tmp_path):
    wb = load_workbook(classeur_donnees)
    faux = FauxMail()
    m.envoyer_mail_pipeline(wb, tmp_path / "f.xlsm", "a@b.fr", "c@d.fr", "afficher", faux, JOUR)
    assert faux.appels[0]["mode"] == "afficher" and faux.appels[0]["cc"] == "c@d.fr"
    assert wb["Log"].cell(row=wb["Log"].max_row, column=3).value == "Brouillon cree pour a@b.fr"


def test_mail_aucun(classeur_donnees, tmp_path):
    wb = load_workbook(classeur_donnees)
    faux = FauxMail()
    assert not m.envoyer_mail_pipeline(wb, tmp_path / "f", mode_mail="aucun", envoyer_mail=faux)
    assert faux.appels == []


def test_mail_erreur_journalisee(classeur_donnees, tmp_path):
    wb = load_workbook(classeur_donnees)
    assert not m.envoyer_mail_pipeline(wb, tmp_path / "f", envoyer_mail=FauxMail(RuntimeError("Outlook absent")))
    ws = wb["Log"]
    assert ws.cell(row=ws.max_row, column=2).value == "ERREUR"
    assert "Outlook absent" in ws.cell(row=ws.max_row, column=3).value


def test_mode_mail_invalide(classeur_donnees, tmp_path):
    with pytest.raises(ValueError):
        m.envoyer_mail_pipeline(load_workbook(classeur_donnees), tmp_path / "f", mode_mail="x")


def test_connecteur_bloque_par_defaut(classeur_donnees, tmp_path):
    """Sans fausse fonction, le connecteur par défaut est appelé : la fixture le bloque,
    l'erreur est journalisée (comme MailErr)."""
    wb = load_workbook(classeur_donnees)
    assert not m.envoyer_mail_pipeline(wb, tmp_path / "f")
    assert "Appel réel" in wb["Log"].cell(row=wb["Log"].max_row, column=3).value


def test_connecteur_outlook_mode_invalide():
    with pytest.raises(ValueError):
        ENVOYER_MAIL_REEL("a", "", "s", "c", [], "bizarre")


# --- journal -----------------------------------------------------------------
def test_log_msg(classeur_vide):
    wb = load_workbook(classeur_vide)
    m.log_msg(wb, "INFO", "un", MAINT)
    m.log_msg(wb, "ERREUR", "deux", MAINT)
    ws = wb["Log"]
    assert [c.value for c in ws[1]] == ["Horodatage", "Niveau", "Message"]
    assert [ws.cell(row=r, column=3).value for r in (2, 3)] == ["un", "deux"]
    assert ws["A2"].value == MAINT


# --- pipeline ----------------------------------------------------------------
def test_pipeline_complet(classeur_vide, tmp_path):
    faux = FauxMail()
    chemin = m.lancer_pipeline_complet(
        classeur_vide, tmp_path / "sortie", region="Nord", envoyer_mail=faux,
        aujourdhui=JOUR, maintenant=MAINT, rng=random.Random(0), mode_mail="afficher")
    assert chemin.name == "VBA_Natran_Test_20251023_140509_VBATF.xlsx" and chemin.exists()
    assert (tmp_path / "sortie" / "graphiques" / "graphique_volumes.png").stat().st_size > 0
    assert faux.appels[0]["pj"] == [chemin] and faux.appels[0]["mode"] == "afficher"
    wb = load_workbook(chemin)
    assert len(lignes(wb)) == 75  # 70 historiques + 5 du jour
    assert len(wb["Feuille1"]._images) == 1
    # le classeur source n'est pas modifié (SaveCopyAs)
    assert "Feuille2" not in load_workbook(classeur_vide).sheetnames


def test_pipeline_enregistrer_source(classeur_vide, tmp_path):
    m.lancer_pipeline_complet(classeur_vide, tmp_path / "s", envoyer_mail=FauxMail(), aujourdhui=JOUR,
                              maintenant=MAINT, enregistrer_source=True)
    assert "Feuille2" in load_workbook(classeur_vide).sheetnames


def test_pipeline_erreur_retourne_none(tmp_path, classeur_vide):
    # dossier de sauvegarde impossible : un fichier existe à la place
    bloc = tmp_path / "bloc"
    bloc.write_text("x")
    assert m.lancer_pipeline_complet(classeur_vide, bloc / "sous", envoyer_mail=FauxMail(),
                                     aujourdhui=JOUR) is None


def test_pipeline_classeur_inexistant(tmp_path):
    chemin = m.lancer_pipeline_complet(tmp_path / "neuf.xlsx", tmp_path / "s", envoyer_mail=FauxMail(),
                                       aujourdhui=JOUR, maintenant=MAINT)
    assert chemin is not None and chemin.exists()


def test_main_cli(classeur_donnees, tmp_path):
    code = m.main([str(classeur_donnees), "--dossier", str(tmp_path / "o"), "--mode-mail", "aucun"])
    assert code == 0
    assert any((tmp_path / "o").glob("*_VBATF.xlsx"))
