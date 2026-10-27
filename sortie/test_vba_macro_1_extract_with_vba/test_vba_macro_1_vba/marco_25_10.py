"""Traduction du module VBA ``marco_25_10``.

Pipeline : Maj données -> Filtre -> Visuel -> Sauvegarde -> Mail.
Feuille2 : données (tableau ``tblDonnees``) ; Feuille1 : visuel (KPI + TCD + graphique).
"""
import argparse
import logging
import random
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.drawing.image import Image as XlImage  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from openpyxl.utils import get_column_letter, range_boundaries  # noqa: E402
from openpyxl.worksheet.filters import AutoFilter, FilterColumn, Filters  # noqa: E402
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: E402

logger = logging.getLogger(__name__)

# --- CONFIGURATION ---------------------------------------------------------
SHEET_VISUEL = "Feuille1"
SHEET_DATA = "Feuille2"
SHEET_LOG = "Log"
TABLE_NAME = "tblDonnees"
FILE_TITLE = "VBA_Natran_Test"
FILE_SUFFIX = "VBATF"
MAIL_TO = "jdavoust201@headmind.com"
MAIL_CC = ""
SEND_MAIL = True  # False = brouillon seulement
SAVE_FOLDER = "C:\\Users\\Jules\\Documents\\Jules\\Projects\\AgentAI\\examples\\sortiesVBA\\"
REGIONS = ("Nord", "Sud", "Est", "Ouest", "Centre")
EN_TETES = ["Date", "Region", "Volume_GWh", "Temperature_C", "Statut"]

EnvoyerMailFn = Callable[..., None]


def regions() -> tuple:
    """VBA : fonction privée ``Regions``."""
    return REGIONS


# --- OUTILS ------------------------------------------------------------------
def get_or_create_sheet(wb: Workbook, nom: str):
    """VBA : ``GetOrCreateSheet`` (ajoute la feuille après la dernière si absente)."""
    if nom in wb.sheetnames:
        return wb[nom]
    return wb.create_sheet(title=nom)


def log_msg(wb: Workbook, niveau: str, msg: str, maintenant: datetime | None = None) -> None:
    """VBA : ``LogMsg`` - ajoute une ligne (horodatage, niveau, message) dans la feuille Log."""
    ws = get_or_create_sheet(wb, SHEET_LOG)
    if ws["A1"].value in (None, ""):
        for i, titre in enumerate(("Horodatage", "Niveau", "Message"), start=1):
            ws.cell(row=1, column=i, value=titre)
    n = 1
    for row in range(ws.max_row, 0, -1):  # équivalent de End(xlUp)
        if ws.cell(row=row, column=1).value not in (None, ""):
            n = row
            break
    n += 1
    cell = ws.cell(row=n, column=1, value=maintenant or datetime.now())
    cell.number_format = "dd/mm/yyyy hh:mm:ss"
    ws.cell(row=n, column=2, value=niveau)
    ws.cell(row=n, column=3, value=msg)
    logger.info("[%s] %s", niveau, msg)


def _ajuster_colonnes(ws, premiere: int, derniere: int) -> None:
    """Approximation de ``Columns(...).AutoFit``."""
    for c in range(premiere, derniere + 1):
        largeur = max(
            (len(_texte(cell.value)) for cell in ws[get_column_letter(c)] if cell.value is not None),
            default=8,
        )
        ws.column_dimensions[get_column_letter(c)].width = min(largeur + 2, 60)


def _texte(v) -> str:
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    return str(v)


def charger_ou_creer_classeur(chemin: Path) -> Workbook:
    """VBA : ``ThisWorkbook`` - ouvre le classeur (le crée avec Feuil1..4 s'il n'existe pas)."""
    chemin = Path(chemin)
    if chemin.exists():
        return load_workbook(chemin, keep_vba=chemin.suffix.lower() == ".xlsm")
    wb = Workbook()
    wb.active.title = "Feuil1"
    for i in range(2, 5):
        wb.create_sheet(f"Feuil{i}")
    return wb


# --- 1. STRUCTURE / DONNEES -------------------------------------------------
def simuler_volume(reg: str, rng: random.Random) -> float:
    """VBA : ``SimulerVolume`` - base par région * (0,85 + Rnd*0,3), arrondi à 1 décimale."""
    base = {"Nord": 120, "Sud": 95, "Est": 80, "Ouest": 70}.get(reg, 60)
    return round(base * (0.85 + rng.random() * 0.3), 1)


def simuler_temperature(rng: random.Random) -> float:
    """VBA : ``SimulerTemperature`` - 2 + Rnd*14, arrondi à 1 décimale."""
    return round(2 + rng.random() * 14, 1)


def _table(wb: Workbook):
    ws = wb[SHEET_DATA]
    if TABLE_NAME not in ws.tables:
        raise KeyError(f"Tableau {TABLE_NAME} introuvable dans {SHEET_DATA}")
    return ws, ws.tables[TABLE_NAME]


def lignes_table(wb: Workbook) -> list[tuple]:
    """Lignes de données (sans l'en-tête) de ``tblDonnees`` sous forme de tuples."""
    ws, tbl = _table(wb)
    c1, r1, c2, r2 = range_boundaries(tbl.ref)
    return [
        tuple(ws.cell(row=r, column=c).value for c in range(c1, c2 + 1))
        for r in range(r1 + 1, r2 + 1)
    ]


def _ajouter_ligne_table(ws, tbl, valeurs: list) -> None:
    """VBA : ``lo.ListRows.Add`` - ajoute une ligne et étend le tableau."""
    c1, r1, c2, r2 = range_boundaries(tbl.ref)
    r = r2 + 1
    for i, v in enumerate(valeurs):
        cell = ws.cell(row=r, column=c1 + i, value=v)
        if i == 0:
            cell.number_format = "dd/mm/yyyy"
    tbl.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r}"
    if tbl.autoFilter is not None:
        tbl.autoFilter.ref = tbl.ref


def ensure_structure(
    wb: Workbook, aujourdhui: date | None = None, rng: random.Random | None = None
) -> bool:
    """VBA : ``EnsureStructure`` - crée Feuille1/Feuille2 et, si le tableau manque,
    14 jours d'historique. Retourne True si la structure a été créée."""
    aujourdhui = aujourdhui or date.today()
    rng = rng or random.Random()
    ws = get_or_create_sheet(wb, SHEET_DATA)
    get_or_create_sheet(wb, SHEET_VISUEL)
    if TABLE_NAME in ws.tables:
        return False

    ws.delete_rows(1, ws.max_row)  # Cells.Clear
    for c, titre in enumerate(EN_TETES, start=1):
        ws.cell(row=1, column=c, value=titre)
    ligne = 2
    for i in range(14, 0, -1):
        jour = datetime.combine(aujourdhui, datetime.min.time()) - timedelta(days=i)
        for r in regions():
            ws.cell(row=ligne, column=1, value=jour)
            ws.cell(row=ligne, column=2, value=r)
            ws.cell(row=ligne, column=3, value=simuler_volume(r, rng))
            ws.cell(row=ligne, column=4, value=simuler_temperature(rng))
            ws.cell(row=ligne, column=5, value="Historique")
            ligne += 1
    tbl = Table(displayName=TABLE_NAME, ref=f"A1:E{ligne - 1}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    tbl.autoFilter = AutoFilter(ref=tbl.ref)
    ws.add_table(tbl)
    for row in range(2, ligne):
        ws.cell(row=row, column=1).number_format = "dd/mm/yyyy"
    _ajuster_colonnes(ws, 1, 5)
    log_msg(wb, "INFO", "Structure initiale creee")
    return True


def ligne_existe(wb: Workbook, d: date, reg: str) -> bool:
    """VBA : ``LigneExiste`` - une ligne existe-t-elle pour ce jour et cette région ?
    (comparaison de la partie date uniquement, région sensible à la casse comme ``=`` VBA)."""
    for valeurs in lignes_table(wb):
        v = valeurs[0]
        if isinstance(v, datetime):
            v = v.date()
        if v == d and valeurs[1] == reg:
            return True
    return False


def ajouter_donnees_du_jour(
    wb: Workbook, aujourdhui: date | None = None, rng: random.Random | None = None
) -> int:
    """VBA : ``AjouterDonneesDuJour`` - ajoute une ligne par région absente pour aujourd'hui.
    Retourne le nombre de lignes ajoutées."""
    aujourdhui = aujourdhui or date.today()
    rng = rng or random.Random()
    ws, tbl = _table(wb)
    nb = 0
    for r in regions():
        if not ligne_existe(wb, aujourdhui, r):
            _ajouter_ligne_table(
                ws,
                tbl,
                [
                    datetime.combine(aujourdhui, datetime.min.time()),
                    r,
                    simuler_volume(r, rng),
                    simuler_temperature(rng),
                    "Phase de test",
                ],
            )
            nb += 1
    log_msg(wb, "INFO", f"{nb} ligne(s) ajoutee(s) pour le {aujourdhui:%d/%m/%Y}")
    return nb


# --- 2. FILTRE --------------------------------------------------------------
def est_region_valide(s: str) -> bool:
    """VBA : ``EstRegionValide`` - comparaison insensible à la casse (vbTextCompare)."""
    return any(r.casefold() == s.casefold() for r in regions())


def _afficher_tout(ws, tbl) -> None:
    if tbl.autoFilter is not None:
        tbl.autoFilter.filterColumn = []
    _, r1, _, r2 = range_boundaries(tbl.ref)
    for r in range(r1 + 1, r2 + 1):
        ws.row_dimensions[r].hidden = False


def appliquer_filtre(wb: Workbook, saisie: str | None) -> str | None:
    """VBA : ``AppliquerFiltre`` - ``saisie`` remplace l'``InputBox`` (vide/None = aucun filtre).
    Retourne la région filtrée, ou None si aucun filtre n'est appliqué (saisie vide ou inconnue)."""
    ws, tbl = _table(wb)
    _afficher_tout(ws, tbl)  # ShowAllData
    saisie = (saisie or "").strip()
    if saisie == "":
        return None
    if not est_region_valide(saisie):
        logger.warning("Region inconnue : %s", saisie)  # MsgBox vbExclamation
        return None
    if tbl.autoFilter is None:
        tbl.autoFilter = AutoFilter(ref=tbl.ref)
    tbl.autoFilter.filterColumn = [FilterColumn(colId=1, filters=Filters(filter=[saisie]))]
    c1, r1, _, r2 = range_boundaries(tbl.ref)
    for r in range(r1 + 1, r2 + 1):
        v = ws.cell(row=r, column=c1 + 1).value
        ws.row_dimensions[r].hidden = not (isinstance(v, str) and v.casefold() == saisie.casefold())
    log_msg(wb, "INFO", f"Filtre applique : {saisie}")
    return saisie


# --- 3. VISUEL --------------------------------------------------------------
def construire_tcd(wb: Workbook) -> pd.DataFrame:
    """Calcule le contenu du TCD ``TCD_Volume`` (lignes = Date, colonnes = Region,
    valeurs = somme de Volume_GWh) sur toute la table, hors filtre."""
    df = pd.DataFrame(lignes_table(wb), columns=EN_TETES)
    if df.empty:
        return pd.DataFrame()
    df["Date"] = pd.to_datetime(df["Date"]).dt.normalize()
    return df.pivot_table(index="Date", columns="Region", values="Volume_GWh", aggfunc="sum")


def exporter_graphique(tcd: pd.DataFrame, dossier_sortie: Path, nom: str = "graphique_volumes.png") -> Path:
    """VBA : ``ChartObjects.Add`` (courbes xlLine) - export PNG matplotlib."""
    dossier_sortie = Path(dossier_sortie)
    dossier_sortie.mkdir(parents=True, exist_ok=True)
    chemin = dossier_sortie / nom
    fig, ax = plt.subplots(figsize=(520 / 72, 300 / 72), dpi=100)
    for col in tcd.columns:
        ax.plot([d.strftime("%d/%m") for d in tcd.index], tcd[col].values, label=col)
    ax.set_title("Volumes distribues par region (GWh)")
    if len(tcd.columns):
        ax.legend()
    ax.tick_params(axis="x", rotation=45)
    fig.tight_layout()
    fig.savefig(chemin)
    plt.close(fig)
    return chemin


def creer_visuel(
    wb: Workbook, dossier_graphiques: Path, maintenant: datetime | None = None
) -> Path:
    """VBA : ``CreerVisuel`` - KPI (formules), TCD (valeurs statiques) et graphique (PNG intégré
    en J5) sur Feuille1. Retourne le chemin du PNG."""
    maintenant = maintenant or datetime.now()
    wsV = wb[SHEET_VISUEL]
    _table(wb)  # vérifie l'existence

    # Nettoyage (Cells.Clear)
    wsV._images = []
    if wsV.max_row:
        wsV.delete_rows(1, wsV.max_row)

    t = TABLE_NAME
    wsV["A1"] = f"Suivi distribution gaz - mis a jour le {maintenant:%d/%m/%Y %H:%M}"
    wsV["A1"].font = Font(bold=True, size=14)
    wsV["A2"] = "Volume total (GWh)"
    wsV["B2"] = f"=SUM({t}[Volume_GWh])"
    wsV["D2"] = "Moyenne / jour (GWh)"
    wsV["E2"] = f"=SUM({t}[Volume_GWh])/SUMPRODUCT(1/COUNTIF({t}[Date],{t}[Date]))"
    wsV["G2"] = "Region du pic"
    for c in ("A2", "D2", "G2"):
        wsV[c].font = Font(bold=True)
    for c in ("B2", "E2"):
        wsV[c].number_format = "#,##0.0"
    wsV["H2"] = f"=INDEX({t}[Region],MATCH(MAX({t}[Volume_GWh]),{t}[Volume_GWh],0))"

    # TCD (openpyxl ne sait pas créer de tableau croisé dynamique) : version statique en A5
    # TODO(vba2py): TCD "TCD_Volume" écrit en valeurs, non actualisable dans Excel.
    tcd = construire_tcd(wb)
    wsV["A5"] = "Volume (GWh)"
    wsV["B5"] = "Region"
    wsV["A6"] = "Date"
    colonnes = sorted(tcd.columns) if not tcd.empty else []
    for j, reg in enumerate(colonnes, start=2):
        wsV.cell(row=6, column=j, value=reg)
    wsV.cell(row=6, column=len(colonnes) + 2, value="Total general")
    ligne = 7
    for jour, ligne_tcd in tcd.iterrows() if not tcd.empty else []:
        c = wsV.cell(row=ligne, column=1, value=jour.to_pydatetime())
        c.number_format = "dd/mm"
        total = 0.0
        for j, reg in enumerate(colonnes, start=2):
            v = ligne_tcd[reg]
            if pd.notna(v):
                wsV.cell(row=ligne, column=j, value=float(v)).number_format = "#,##0.0"
                total += float(v)
        wsV.cell(row=ligne, column=len(colonnes) + 2, value=total).number_format = "#,##0.0"
        ligne += 1
    if not tcd.empty:
        wsV.cell(row=ligne, column=1, value="Total general")
        for j, reg in enumerate(colonnes, start=2):
            wsV.cell(row=ligne, column=j, value=float(tcd[reg].sum())).number_format = "#,##0.0"
        wsV.cell(row=ligne, column=len(colonnes) + 2, value=float(tcd.sum().sum())).number_format = "#,##0.0"

    # Graphique à côté du TCD
    png = exporter_graphique(tcd, dossier_graphiques)
    img = XlImage(str(png))
    img.width, img.height = 520, 300
    wsV.add_image(img, "J5")

    _ajuster_colonnes(wsV, 1, 8)
    log_msg(wb, "INFO", "Visuel regenere")
    return png


# --- 4. SAUVEGARDE ----------------------------------------------------------
def sauvegarder_version(
    wb: Workbook, dossier: Path, extension: str = ".xlsm", maintenant: datetime | None = None
) -> Path:
    """VBA : ``SauvegarderVersion`` - copie ``titre_aaaammjj_hhmmss_VBATF`` (SaveCopyAs)."""
    maintenant = maintenant or datetime.now()
    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    nom = f"{FILE_TITLE}_{maintenant:%Y%m%d_%H%M%S}_{FILE_SUFFIX}{extension}"
    chemin = dossier / nom
    log_msg(wb, "INFO", f"Sauvegarde : {chemin}", maintenant)  # logué avant la copie, comme en VBA
    wb.save(chemin)
    return chemin


# --- 5. MAIL ----------------------------------------------------------------
def construire_corps_mail(chemin_pj: Path) -> str:
    """Corps du mail de ``EnvoyerMail`` (retours à la ligne vbCrLf)."""
    return (
        "Bonjour,\r\n\r\n"
        "Ceci est un envoi automatique en phase de test.\r\n"
        "Le graphique et les donnees ont ete mis a jour.\r\n"
        f"Fichier : {Path(chemin_pj).name}\r\n\r\n"
        "Cordialement,\r\nRobot VBATF"
    )


def envoyer_mail_pipeline(
    wb: Workbook,
    chemin_pj: Path,
    destinataire: str = MAIL_TO,
    cc: str = MAIL_CC,
    mode_mail: str = "envoyer",
    envoyer_mail: EnvoyerMailFn | None = None,
    aujourdhui: date | None = None,
) -> bool:
    """VBA : ``EnvoyerMail`` - envoie ou affiche le mail avec la PJ via le connecteur Outlook.
    Les erreurs sont journalisées (``MailErr``), pas propagées. Retourne True si le mail a été traité."""
    aujourdhui = aujourdhui or date.today()
    if mode_mail == "aucun":
        log_msg(wb, "INFO", "Mail desactive (mode aucun)")
        return False
    if mode_mail not in ("envoyer", "afficher"):
        raise ValueError(f"Mode de mail inconnu : {mode_mail!r}")
    if envoyer_mail is None:
        from connecteurs import outlook  # résolu à l'appel (substituable dans les tests)

        envoyer_mail = outlook.envoyer_mail
    try:
        envoyer_mail(
            destinataire,
            cc,
            f"[PHASE DE TEST] Mise a jour suivi distribution - {aujourdhui:%d/%m/%Y}",
            construire_corps_mail(chemin_pj),
            [chemin_pj],
            mode_mail,
        )
    except Exception as e:  # équivalent de On Error GoTo MailErr
        log_msg(wb, "ERREUR", f"Mail impossible : {e}")
        return False
    log_msg(wb, "INFO", ("Mail envoye a " if mode_mail == "envoyer" else "Brouillon cree pour ") + destinataire)
    return True


# --- POINT D'ENTREE ---------------------------------------------------------
def lancer_pipeline_complet(
    classeur: Path,
    dossier_sauvegarde: Path | str = SAVE_FOLDER,
    dossier_graphiques: Path | str | None = None,
    region: str | None = None,
    destinataire: str = MAIL_TO,
    cc: str = MAIL_CC,
    mode_mail: str = "envoyer" if SEND_MAIL else "afficher",
    envoyer_mail: EnvoyerMailFn | None = None,
    aujourdhui: date | None = None,
    maintenant: datetime | None = None,
    rng: random.Random | None = None,
    enregistrer_source: bool = False,
) -> Path | None:
    """VBA : ``LancerPipelineComplet`` - structure, données du jour, filtre, visuel, sauvegarde, mail.

    ``classeur`` remplace ``ThisWorkbook`` ; ``region`` remplace l'``InputBox`` du filtre.
    Retourne le chemin de la copie sauvegardée, ou None en cas d'erreur (journalisée,
    comme le ``ErrHandler`` VBA). ``enregistrer_source`` réécrit aussi le classeur d'origine.
    """
    classeur = Path(classeur)
    wb = charger_ou_creer_classeur(classeur)
    try:
        ensure_structure(wb, aujourdhui, rng)
        ajouter_donnees_du_jour(wb, aujourdhui, rng)
        appliquer_filtre(wb, region)
        dossier_graphiques = Path(dossier_graphiques) if dossier_graphiques else Path(dossier_sauvegarde) / "graphiques"
        creer_visuel(wb, dossier_graphiques, maintenant)
        chemin = sauvegarder_version(wb, Path(dossier_sauvegarde), classeur.suffix or ".xlsm", maintenant)
        envoyer_mail_pipeline(wb, chemin, destinataire, cc, mode_mail, envoyer_mail, aujourdhui)
        log_msg(wb, "INFO", f"Pipeline termine : {chemin}")
        logger.info("Pipeline termine.\n%s", chemin)  # MsgBox vbInformation
        if enregistrer_source:
            wb.save(classeur)
        return chemin
    except Exception as e:  # équivalent de ErrHandler
        log_msg(wb, "ERREUR", f"{type(e).__name__} - {e}")
        logger.error("Erreur : %s", e)
        if enregistrer_source:
            wb.save(classeur)
        return None


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Pipeline distribution gaz (traduction de LancerPipelineComplet)")
    p.add_argument("classeur", type=Path, help="classeur .xlsx/.xlsm (créé s'il n'existe pas)")
    p.add_argument("--dossier", type=Path, default=Path(SAVE_FOLDER), help="dossier de sauvegarde des versions")
    p.add_argument("--dossier-graphiques", type=Path, default=None, help="dossier du PNG (défaut : <dossier>/graphiques)")
    p.add_argument("--region", default=None, help=f"filtre ({', '.join(REGIONS)}) ; vide = aucun")
    p.add_argument("--destinataire", default=MAIL_TO)
    p.add_argument("--cc", default=MAIL_CC)
    p.add_argument("--mode-mail", choices=["envoyer", "afficher", "aucun"], default="envoyer" if SEND_MAIL else "afficher")
    p.add_argument("--enregistrer-source", action="store_true", help="réécrit aussi le classeur d'origine")
    a = p.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    res = lancer_pipeline_complet(
        a.classeur, a.dossier, a.dossier_graphiques, a.region, a.destinataire, a.cc,
        a.mode_mail, enregistrer_source=a.enregistrer_source,
    )
    return 0 if res else 1


if __name__ == "__main__":
    raise SystemExit(main())
